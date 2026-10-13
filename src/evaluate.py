#!/usr/bin/env python3
"""Read a graded check sheet, store the grades as labels, and measure the
search and the classifier against them.

A person grades paragraphs in data/review/check_<set>_s<seed>.xlsx
(src/check_sheet.py). This script:

1. stores each grade in data/labels/<set>.csv, keyed by the SHA-256 of the
   paragraph text, the label set, the label and the scale it was graded on,
   never by position (AGENTS.md rule 9). A later grade of the same text and
   label replaces an earlier one;
2. for each label, measures precision and recall at a range of cut-offs:
   the classifier's probability that the grade is at least k, and the
   search's best cosine score. The sheet was drawn by stratum, so every graded
   row is weighted by its stratum's size over the number drawn, which
   estimates the figures for the whole corpus; the report gives the number of
   graded rows behind each figure.

Nothing else reads the labels. No cut-off is chosen here: the report shows
the trade-off; a person picks.

Reads  data/review/check_<set>_s<seed>.{xlsx,key.csv}, data/classify/<run>/scores.csv
Writes data/labels/<set>.csv, data/reports/evaluate_<set>.md

  python3 src/evaluate.py assets --seed 1
"""
import argparse, csv, datetime, hashlib, os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from paths import data_root, where  # noqa: E402
import classify  # noqa: E402
from search import TAXONOMY, _rows  # noqa: E402

LABEL_COLS = ["label_set", "label_id", "text_sha256", "scale_sha256", "grade",
              "paragraph_id", "source", "graded_on"]
THRESHOLDS = [0.1, 0.2, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9]


def scale_sha(scale):
    return hashlib.sha256("\n".join(scale).encode("utf-8")).hexdigest()


def read_grades(xlsx):
    """{row_id: grade} for the rows a person graded."""
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, read_only=True)["grade these"]
    rows = ws.iter_rows(values_only=True)
    head = list(next(rows))
    i, g = head.index("row_id"), head.index("grade")
    out = {}
    for r in rows:
        if r[g] is None or str(r[g]).strip() == "":
            continue
        out[str(r[i])] = int(float(str(r[g]).strip()))
    return out


def store_labels(name, data, key, grades, scale, source):
    path = where(data, "labels", f"{name}.csv")
    current = {}
    if os.path.exists(path):
        for r in _rows(path):
            current[(r["label_id"], r["text_sha256"], r["scale_sha256"])] = r
    today = datetime.date.today().isoformat()
    sha = scale_sha(scale)
    for k in key:
        if k["row_id"] in grades:
            current[(k["label_id"], k["text_sha256"], sha)] = {
                "label_set": name, "label_id": k["label_id"], "text_sha256": k["text_sha256"],
                "scale_sha256": sha, "grade": grades[k["row_id"]],
                "paragraph_id": k["paragraph_id"], "source": source, "graded_on": today}
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path + ".part", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=LABEL_COLS, lineterminator="\n")
        w.writeheader()
        w.writerows(sorted(current.values(), key=lambda r: (r["label_id"], r["text_sha256"])))
    os.replace(path + ".part", path)
    return path


def weighted_pr(rows, predicted, positive):
    """(precision, recall, graded rows predicted, graded rows positive) with
    stratum weights; None where undefined."""
    tp = sum(r["w"] for r in rows if predicted(r) and positive(r))
    pp = sum(r["w"] for r in rows if predicted(r))
    ap = sum(r["w"] for r in rows if positive(r))
    return (tp / pp if pp else None, tp / ap if ap else None,
            sum(1 for r in rows if predicted(r)), sum(1 for r in rows if positive(r)))


def _f(x):
    return "-" if x is None else f"{x:.2f}"


def evaluate(name, data, seed=1, taxonomy=TAXONOMY):
    base = where(data, "review", f"check_{name}_s{seed}")
    key = _rows(base + ".key.csv")
    grades = read_grades(base + ".xlsx")
    labels, scale = classify.load_labels(name, taxonomy)
    lpath = store_labels(name, data, key, grades, scale, os.path.basename(base))
    cdir = where(data, "classify", key[0]["classify_dir"])
    probs = {(r["label_id"], r["paragraph_id"]): r for r in _rows(os.path.join(cdir,
                                                                             "scores.csv"))}
    graded = []
    for k in key:
        if k["row_id"] not in grades:
            continue
        p = probs[(k["label_id"], k["paragraph_id"])]
        graded.append({"label_id": k["label_id"], "grade": grades[k["row_id"]],
                       "w": int(k["stratum_size"]) / int(k["stratum_drawn"]),
                       "p": [float(p[f"p{i}"]) for i in range(len(scale))],
                       "search": float(k["search_best"]) if k["search_best"] else None})
    top = len(scale) - 1
    lines = [f"# Evaluation of {name} against graded paragraphs", "",
             f"{len(graded)} of {len(key)} rows graded in {os.path.basename(base)}.xlsx; "
             f"labels stored in {os.path.relpath(lpath, data)}. Figures are weighted to the "
             f"corpus by sampling stratum; n is the number of graded rows behind each.", "",
             "Positive means the person graded at least k. The classifier predicts positive "
             "when its probability of a grade of at least k reaches the cut-off; the search "
             "when the paragraph's best cosine score reaches it.", ""]
    for lab in labels:
        rows = [r for r in graded if r["label_id"] == lab["id"]]
        if not rows:
            continue
        lines += [f"## {lab['name']} ({lab['id']}): {len(rows)} graded", ""]
        for k in sorted({top - 1, top}):
            pos = (lambda r, k=k: r["grade"] >= k)
            lines += [f"Positive: grade {k} or above.", "",
                      "| cut-off | classifier precision | recall | n predicted | n positive |",
                      "|---|---|---|---|---|"]
            for t in THRESHOLDS:
                pr = weighted_pr(rows, lambda r, t=t, k=k: sum(r["p"][k:]) >= t, pos)
                lines.append(f"| {t:.1f} | {_f(pr[0])} | {_f(pr[1])} | {pr[2]} | {pr[3]} |")
            scored = sorted({round(r["search"], 2) for r in rows if r["search"] is not None})
            if scored:
                lines += ["", "| search score at least | precision | recall | n predicted |",
                          "|---|---|---|---|"]
                for s in scored[::max(1, len(scored) // 8)]:
                    pr = weighted_pr([r for r in rows if r["search"] is not None],
                                     lambda r, s=s: r["search"] >= s, pos)
                    lines.append(f"| {s:.2f} | {_f(pr[0])} | {_f(pr[1])} | {pr[2]} |")
            lines.append("")
    out = where(data, "reports", f"evaluate_{name}.md")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    return out, len(graded)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("label_set")
    ap.add_argument("--data", default=data_root())
    ap.add_argument("--seed", type=int, default=1)
    args = ap.parse_args()
    out, n = evaluate(args.label_set, args.data, args.seed)
    print(f"{n} graded rows; wrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
