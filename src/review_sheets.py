#!/usr/bin/env python3
"""Write two plain spreadsheets for reading the cleaning by eye.

Reads  data/appraisal/paragraphs.csv, data/appraisal/text/{doc_id}.txt,
       data/raw/pdf_text/{doc_id}.txt, data/procurement/packages.csv
Writes data/review/paragraphs_s{seed}.xlsx    100 paragraphs: cleaned, split into
                                              sentences, and the raw text beside them
       data/review/packages_s{seed}.xlsx      100 packages, raw next to cleaned

One sheet each, one row per unit, and two empty columns - ok, note - for the
person reading. The files live in data/, so they are never committed.

Paragraphs are drawn only from those given to the model (for_model): front
matter and tables are kept for reading but are not reviewed.

A file that already exists is left alone: it may hold someone's notes. Pass a
new --seed for a fresh draw, or --force to overwrite.

  python3 src/review_sheets.py
  python3 src/review_sheets.py --n 200 --seed 2

The raw column is the stretch of the PDF's own text the paragraph was built
from (raw_start:raw_end in data/raw/pdf_text/), shown exactly as read: the
lines, any running header or footnote number the cleaner removed from between
them, and the page break if the paragraph crossed one.
"""
import argparse, csv, os, sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from paths import data_root, where                    # noqa: E402
from review import draw                        # noqa: E402
import sentences                               # noqa: E402

CELL_MAX = 32_000          # Excel refuses a cell over 32,767 characters

PARA_COLS = [("paragraph_id", 18), ("project_ids", 12), ("page_from", 6),
             ("section_path", 12), ("block", 12), ("n_tokens", 8),
             ("cleaned", 70), ("sentences", 70), ("raw", 70), ("ok", 6), ("note", 30)]
PKG_COLS = [("package_id", 18), ("project_id", 10), ("plan_version", 10),
            ("category", 12), ("borrower_ref", 18), ("description", 45),
            ("description_clean", 45), ("lot_or_phase", 10), ("component_number", 6),
            ("component", 24), ("component_source", 10), ("is_rebid", 7),
            ("is_placeholder", 8), ("method", 8), ("method_name", 18),
            ("method_source", 10), ("market_approach", 14), ("status_raw", 14),
            ("status", 14), ("status_source", 10), ("status_as_of", 11),
            ("estimated_amount", 12), ("currency", 7), ("amount_source", 10), ("planned_date", 11), ("ok", 6), ("note", 30)]


def cell_text(value):
    """Excel refuses control characters. A form feed is a page break in the
    raw file, which the reader wants to see, so it gets a visible marker."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    text = str(value if value is not None else "").replace("\f", "\n[page break]\n")
    return ILLEGAL_CHARACTERS_RE.sub("", text)[:CELL_MAX]


def write_sheet(path, cols, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append([c for c, _ in cols])
    for r in rows:
        ws.append([cell_text(r.get(c)) for c, _ in cols])
    for i, (_, width) in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "A2"
    wb.save(path)


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def paragraph_rows(data, n, seed):
    paras = [r for r in read_csv(where(data, "appraisal", "paragraphs.csv"))
             if r.get("for_model") == "true"]
    picked = draw(paras, n, seed)
    clean_cache, raw_cache = {}, {}
    out, missed = [], 0
    for r in sorted(picked, key=lambda r: r["paragraph_id"]):
        did = r["doc_id"]
        if did not in clean_cache:
            with open(where(data, "text", f"{did}.txt"), encoding="utf-8") as fh:
                clean_cache[did] = fh.read()
            path = where(data, "pdf_text", f"{did}.txt")
            raw_cache[did] = open(path, encoding="utf-8").read() \
                if os.path.exists(path) else None
        body = clean_cache[did][int(r["char_start"]):int(r["char_end"])]
        raw = raw_cache[did]
        if raw is None or r.get("raw_start", "") == "":
            found, missed = "(no PDF text for this paragraph)", missed + 1
        else:
            found = raw[int(r["raw_start"]):int(r["raw_end"])]
        cuts = sentences.split(body)
        out.append(dict(r, cleaned=body,
                        sentences="\n".join(f"[{k}] {body[a:b]}" for k, (a, b)
                                             in enumerate(cuts, 1)),
                        raw=found))
    return out, missed


def package_rows(data, n, seed):
    path = where(data, "procurement", "packages.csv")
    if not os.path.exists(path):
        return None
    picked = draw(read_csv(path), n, seed, key="project_id")
    return sorted(picked, key=lambda r: (r["project_id"], r["package_id"]))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default=data_root())
    ap.add_argument("--n", type=int, default=100)
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--force", action="store_true", help="overwrite existing files")
    args = ap.parse_args()

    out_dir = where(args.data, "review")
    os.makedirs(out_dir, exist_ok=True)
    jobs = [("paragraphs", PARA_COLS), ("packages", PKG_COLS)]
    for name, cols in jobs:
        dest = os.path.join(out_dir, f"{name}_s{args.seed}.xlsx")
        if os.path.exists(dest) and not args.force:
            print(f"kept {dest} (exists; may hold notes - use --seed or --force)")
            continue
        if name == "paragraphs":
            rows, missed = paragraph_rows(args.data, args.n, args.seed)
            extra = f", {missed} with no PDF text"
        else:
            rows, extra = package_rows(args.data, args.n, args.seed), ""
            if rows is None:
                print("skipped packages: run the procurement stages first")
                continue
        write_sheet(dest, cols, rows)
        print(f"wrote {dest}: {len(rows)} rows{extra}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
