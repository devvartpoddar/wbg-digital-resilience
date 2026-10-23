#!/usr/bin/env python3
"""A sheet of a measure's units for a person to grade, and the measurement
of the decision model against those grades: the precision and recall that
can be quoted, with their confidence intervals.

Draws from the current measure run (src/measure.py) by stratum of the
pipeline's decision, so every figure can be weighted back to the graded pool:
a training sheet (--design train, the default) near the floor, where grades
teach the probability layer and the skip rule most; an evaluation sheet
(--design eval), a random draw of what counts and what does not, whose
grades are never used to fit anything. The person sees the sentence with its lead-in, the paragraph around it, its
section and component, and the measure with its definition; never the
model's answer. They enter a grade on the set's scale and, if they wish, the
asset the sentence is about.

  ./wbg check cra                       draw: data/review/check_units_cra_s1.xlsx (+ .key.csv)
  ./wbg check cra --evaluate            read the grades back, store them as labels
                                        (data/labels/cra_units.csv, keyed by text and
                                        question hash), write data/reports/evaluate_units_cra.md
                                        and fit the probability layer (src/posterior.py)
  ./wbg check cra --elusion             draw units the funnel never read, from just below
                                        the pool threshold and from far below it:
                                        data/review/check_units_cra_elusion_s1.xlsx
  ./wbg check cra --elusion --evaluate  what the pool missed, with intervals

What the report says, and does not: the weighted precision and recall of the
pipeline's decision (the posterior at the floor) against the person's grade
2 or more, over the graded pool (candidates and feedback units), with 95 per
cent Wilson intervals on the effective sample size. The probability layer
(src/posterior.py) and the skip rule's gate (src/skip.py) are fitted on the
training sheets' grades under the current question; the next run reads
them. Recall outside the pool is
the elusion sheet's job: a random draw of the units the funnel never read,
in two strata (near: within ELUSION_BAND below the pool threshold; far: the
rest), weighted back to their strata, which is how a technology-assisted
review is signed off.
"""
import argparse, csv, datetime, hashlib, json, math, os, random, re, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from paths import data_root, where  # noqa: E402
import classify  # noqa: E402
import measure  # noqa: E402
import panel  # noqa: E402

CELL_MAX = 32_000   # Excel's limit on a cell's characters

# What a string may keep to be written into the workbook: XML 1.0's Char
# production, because openpyxl writes through lxml and lxml refuses anything
# outside it. openpyxl's own ILLEGAL_CHARACTERS_RE is NOT enough - it covers
# the C0 controls but not the noncharacters U+FFFE and U+FFFF, and one U+FFFF
# inside a Chad PAD paragraph killed a whole `./wbg check all` draw at write
# time with "All strings must be XML compatible". Nothing is stripped at parse
# time: a paragraph is kept as the document says it. This is the display
# boundary, so the loss is here, and the tab has to open for a person to grade
# it at all.
XML_TEXT_RE = re.compile("[^\t\n\r\x20-\ud7ff\ue000-\ufffd\U00010000-\U0010ffff]")


def cell_text(value):
    """Excel refuses control characters. A form feed is a page break in the
    raw file, which the reader wants to see, so it gets a visible marker."""
    text = str(value if value is not None else "").replace("\f", "\n[page break]\n")
    return XML_TEXT_RE.sub("", text)[:CELL_MAX]

# The two sheets the method grades on (docs/method/measures.md, "Training and
# evaluation sheets"). A training sheet (--design train) is drawn where a
# grade teaches most: near the floor, plus a few far from it. An evaluation
# sheet (--design eval) is a random draw of the pool by the decision, so its
# weighted precision and recall are the dashboards'; its grades are never
# used to fit a layer or a gate, nor to edit a definition.
TRAIN_PER = {"train:near": 30, "train:far": 10}
TRAIN_NEAR = 0.2         # "near": the posterior within this of the floor
EVAL_PER = {"eval:committed": 60, "eval:not_committed": 30}
EVAL = "eval:"           # the stratum prefix of an evaluation sheet's rows
ELUSION_PER = {"near": 25, "far": 25}
ELUSION_BAND = 0.10      # "near": best cosine within this much below the pool threshold
KEY_COLS = ["row_id", "label_id", "unit_id", "paragraph_id", "text_sha256", "question_sha256",
            "stratum", "stratum_size", "stratum_drawn", "measure_dir"]
SHEET_COLS = [("row_id", 8), ("measure", 20), ("measure means", 40), ("project", 10),
              ("section", 20), ("component", 18), ("sentence", 70), ("paragraph", 70),
              ("grade", 8), ("asset", 14), ("type", 18), ("note", 24)]
LABEL_COLS = ["label_set", "label_id", "unit_id", "text_sha256", "question_sha256", "grade",
              "asset", "type", "source", "graded_on", "stratum", "weight", "unsure", "sheet"]
THRESHOLDS = [0.15, 0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9]
# The band design (--design bands): drawn by the posterior, densest where a
# cut-off would be chosen, so the floor can be read off the graded sheet. Each band is (stratum, low, high, rows).
BANDS = [("p<0.15", 0.0, 0.15, 5), ("p0.15-0.4", 0.15, 0.4, 8), ("p0.4-0.6", 0.4, 0.6, 12),
         ("p0.6-0.8", 0.6, 0.8, 10), ("p0.8+", 0.8, 1.01, 5)]
# The floor design (--design floor): about 20 units a measure in bands around
# its floor (control panel floor.<set>), to check a chosen floor; a measure
# without one is drawn in BANDS at about half the rows.
FLOOR_STEPS = [(-1.0, -0.2, 2), (-0.2, -0.1, 3), (-0.1, 0.0, 5), (0.0, 0.1, 5), (0.1, 0.25, 3),
               (0.25, 1.0, 2)]


def floor_bands(floor, scale=1.0):
    """[(stratum, low, high, rows)] around a floor, or BANDS halved;
    `scale` multiplies every band's rows (the tuning draw: 1.4, about 28)."""
    rows = lambda n: int(n * scale + 0.5)
    if floor is None:
        return [(s, lo, hi, rows((n + 1) // 2)) for s, lo, hi, n in BANDS]
    out = []
    for a, b, n in FLOOR_STEPS:
        lo, hi = max(0.0, floor + a), min(1.01, floor + b)
        if hi > lo:
            out.append((f"p{lo:.2f}-{min(hi, 1.0):.2f}", lo, hi, rows(n)))
    return out


TUNE_SCALE = 1.4   # the tuning draw: the floor bands at 1.4 times (about 28), links over the whole range


LINK_FLOOR_BANDS = [("p0.5-0.6", 0.5, 0.6, 8), ("p0.6-0.7", 0.6, 0.7, 8), ("p0.7-0.8", 0.7, 0.8, 8),
                    ("p0.8+", 0.8, 1.01, 6)]
LINK_BANDS = [("p<0.3", 0.0, 0.3, 10), ("p0.3-0.5", 0.3, 0.5, 15), ("p0.5-0.7", 0.5, 0.7, 15),
              ("p0.7-0.9", 0.7, 0.9, 15), ("p0.9+", 0.9, 1.01, 15)]
HOW_TO = [
    "Each measure tab: read the sentence (with its paragraph, section and component) and enter a "
    "grade on the measure's scale (tab 'grades'). Asset and type are optional (tab 'types' "
    "defines the types). Asset: 'other' when the commitment is about something that is not a "
    "digital asset, 'none' when the sentence names no asset; for several, type them joined by "
    "| (fiber|data_hosting). Type: the same, | for several. Leave a row blank to skip it.",
    "Rows were drawn by the model's probability that the sentence is a commitment, more of them "
    "near where a cut-off would fall; the model's answer is not shown.",
    "Tab 'package links': does carrying out the package deliver or apply the commitment? "
    "Enter yes, no or unsure.",
    "Tab 'pair checks': two measures both count the sentence; which does it commit the project "
    "to? only A, only B, both, neither, or unsure.",
    "Tab 'package assets': does the package buy, build, install, connect, operate or design "
    "the asset named? yes, no or unsure.",
    "Do not reorder, insert or delete rows or columns: row_id joins each row to its key file.",
]
ALL = "all"              # ./wbg check all: every measure, one workbook, plus the package links
LINK_KEY_COLS = ["row_id", "label_set", "unit_id", "package_id", "text_sha256", "package_sha256",
                 "question_sha256", "probability", "stratum", "stratum_size", "stratum_drawn",
                 "link_run"]
LINK_SHEET_COLS = [("row_id", 8), ("project", 10), ("measure", 20), ("commitment", 70),
                   ("commitment component", 14), ("package", 60), ("package component", 18),
                   ("applies", 10), ("note", 24)]
LINK_LABEL_COLS = ["label_set", "unit_id", "package_id", "text_sha256", "package_sha256",
                   "question_sha256", "applies", "weight", "source", "graded_on", "stratum",
                   "sheet"]


def key_weight(k):
    """A key row's design weight: its own (one over its chance in a combined
    draw, src/combined.py), else its stratum's size over the rows drawn."""
    if k.get("weight"):
        return float(k["weight"])
    return int(k["stratum_size"]) / int(k["stratum_drawn"])


def wilson(k, n, z=1.96):
    if not n:
        return (None, None)
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return (max(0.0, c - h), min(1.0, c + h))


def _posterior(u):
    p = u.get("p_committed_calibrated", "")
    return float(p) if p != "" else measure.score_of(u)


def draw(units, floor, seed, design="train"):
    """A training or evaluation draw of the units the model read, by the
    decision: (unit, stratum, stratum size, rows drawn)."""
    rng = random.Random(seed)
    strata = {}
    for u in units:
        if u["sent"] != "true":
            continue
        p = _posterior(u)
        if design == "eval":
            s = "eval:committed" if p >= floor else "eval:not_committed"
        else:
            s = "train:near" if abs(p - floor) <= TRAIN_NEAR else "train:far"
        strata.setdefault(s, []).append(u)
    per = EVAL_PER if design == "eval" else TRAIN_PER
    picked = []
    for s in sorted(per):
        pool = sorted(strata.get(s, []), key=lambda u: u["unit_id"])
        n = min(per[s], len(pool))
        for u in rng.sample(pool, n):
            picked.append((u, s, len(pool), n))
    return picked


def draw_bands(units, top, seed, bands=BANDS):
    """The units the model read, by band of the probability the floor reads
    (the probability layer's, else the model's score for the grades the
    control panel counts): (unit, stratum, stratum size, rows drawn)."""
    rng = random.Random(seed)
    strata = {}
    for u in units:
        if u["sent"] != "true" or u.get(f"p{top}", "") == "":
            continue
        p = float(u["p_committed_calibrated"]) if u.get("p_committed_calibrated") \
            else measure.score_of(u)
        s = next((b[0] for b in bands if b[1] <= p < b[2]), None)
        if s is not None:
            strata.setdefault(s, []).append(u)
    picked = []
    for s, _, _, per in bands:
        pool = sorted(strata.get(s, []), key=lambda u: u["unit_id"])
        n = min(per, len(pool))
        for u in rng.sample(pool, n):
            picked.append((u, s, len(pool), n))
    return picked


def draw_elusion(run, units, per, seed):
    """Units the funnel never read, by stratum of their best cosine over the
    measure's phrases: near (within ELUSION_BAND below the pool threshold)
    and far (the rest). Returns rows shaped like the run's own."""
    import search
    rng = random.Random(seed)
    seen = {u["unit_id"] for u in units}
    folder, _ = search.run_folder(run.name, run.data)
    best = {h["unit_id"]: float(h["score"])
            for h in search.hits(folder, run.data, labels=[run.labels[0]["id"]])}
    floor = run.settings["min_cosine"] or 1.0
    strata = {"near": [], "far": []}
    for r in run.index:
        uid = r["unit_id"]
        if uid in seen or uid not in best:
            continue
        strata["near" if best[uid] >= floor - ELUSION_BAND else "far"].append(r)
    picked = []
    for s in sorted(strata):
        pool = sorted(strata[s], key=lambda r: r["unit_id"])
        n = min(per.get(s, 0), len(pool))
        for r in rng.sample(pool, n):
            row = run.base_row(r, "elusion", "", "", f"{best[r['unit_id']]:.6f}")
            picked.append((row, s, len(pool), n))
    return picked


def set_types(name, data):
    """[(type_id, name, definition)] a measure's committed units can be: its
    list and the OTH types the gap step proposed, as the dashboards show them."""
    import load_pg
    return [(t["type_id"], t["type_name"], t["definition"]) for t in load_pg.type_rows(data)
            if t["label_set"] == name]


def sheet_rows(name, data, seed=1, elusion=False, design="train", **kw):
    """(key rows, sheet rows, scale, asset ids) of one measure's draw."""
    folder = measure.current_run(name, data, **kw)
    if not folder:
        raise SystemExit(f"check: no measure run for {name} as it stands; run ./wbg measure {name}")
    run = measure.Run(name, data, build=False, log=lambda *_: None, **kw)
    run.load_corpus()
    qsha = classify.question_shas(run.qs)
    with open(os.path.join(folder, "units.csv"), newline="", encoding="utf-8") as fh:
        units = list(csv.DictReader(fh))
    if elusion:
        picked = draw_elusion(run, units, ELUSION_PER, seed)
    elif design == "bands":
        picked = draw_bands(units, len(run.scale) - 1, seed)
    elif design in ("floor", "tune"):
        import labels
        picked = draw_bands(units, len(run.scale) - 1, seed,
                            floor_bands(panel.floor(name),
                                        TUNE_SCALE if design == "tune" else 1.0))
    elif design in ("train", "eval"):
        picked = draw(units, run.floor, seed, design)
    else:
        raise SystemExit(f"check: unknown design {design}")
    rng = random.Random(seed)
    rng.shuffle(picked)
    idx = {r["unit_id"]: r for r in run.index}
    key, sheet, texts = [], [], {}
    lab = run.labels[0]
    for i, (u, stratum, size, drawn) in enumerate(picked, 1):
        para = run.paras[u["paragraph_id"]]
        doc = measure._text(texts, data, u["doc_id"])
        state = measure.state_for(idx[u["unit_id"]], para, u["text"], run.names)
        paragraph = doc[int(para["char_start"]):int(para["char_end"])]
        key.append({"row_id": i, "label_id": u["label_id"], "unit_id": u["unit_id"],
                    "paragraph_id": u["paragraph_id"],
                    "text_sha256": text_key(u["text"]),
                    "question_sha256": qsha[u["label_id"]], "stratum": stratum,
                    "stratum_size": size, "stratum_drawn": drawn,
                    "measure_dir": os.path.basename(folder)})
        sheet.append({"row_id": i, "measure": lab["name"], "measure means": lab["definition"],
                      "project": u["project_ids"], "section": state.get("section", ""),
                      "component": state.get("component", ""), "sentence": u["text"],
                      "paragraph": paragraph})
    return key, sheet, run.scale, [a["id"] for a in run.assets]


def _write_key(path, cols, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def build(name, data, seed=1, elusion=False, design="train", **kw):
    near_gate = kw.pop("near_gate", False)
    if kw.get("dry_run") and not (name == ALL and design in ("train", "eval")):
        raise SystemExit("check: --dry-run is for ./wbg check all with --design train or eval")
    if name == ALL:
        return build_all(data, seed, design=design, near_gate=near_gate, **kw)
    key, sheet, scale, assets = sheet_rows(name, data, seed, elusion, design, **kw)
    base = where(data, "review", f"check_units_{name}{'_elusion' if elusion else ''}_s{seed}")
    if os.path.exists(base + ".xlsx"):
        raise SystemExit(f"check: {base}.xlsx exists and may hold grades; use another --seed")
    os.makedirs(os.path.dirname(base), exist_ok=True)
    _write_key(base + ".key.csv", KEY_COLS, key)
    _write_xlsx(base + ".xlsx", [{"tab": "grade these", "rows": sheet, "scale": scale,
                                  "assets": assets, "types": set_types(name, data),
                                  "measure": name}])
    return base + ".xlsx", len(sheet)


def all_path(data, seed):
    return where(data, "review", f"check_units_{ALL}_s{seed}.xlsx")


def build_all(data, seed, design="train", near_gate=False, **kw):
    """One workbook for a person to grade: a tab per measure (drawn by
    `design`), and a tab of package links. Each tab has its own key file,
    check_units_<set>_s<seed>.key.csv and check_links_s<seed>.key.csv, so
    `--evaluate` reads each measure as it reads a sheet of its own. The
    training and evaluation designs draw the combined sheets instead (a tab
    of sentences and one of packages, every decision read off each item
    once; src/combined.py), into the same key files."""
    import labels
    if design in ("train", "eval"):
        import combined
        xlsx, n, k = combined.build(data, seed, design, near_gate, kw.get("projects"),
                                    kw.get("dry_run", False))
        return xlsx, n + k
    xlsx = all_path(data, seed)
    if os.path.exists(xlsx):
        raise SystemExit(f"check: {xlsx} exists and may hold grades; use another --seed")
    tabs, keys = [], []
    for name in labels.label_sets():
        if name == measure.ASSET_SET:
            continue
        if not measure.current_run(name, data, **kw):
            print(f"check: {name}: no current measure run; left out")
            continue
        key, sheet, scale, assets = sheet_rows(name, data, seed, design=design, **kw)
        base = where(data, "review", f"check_units_{name}_s{seed}")
        if os.path.exists(base + ".xlsx") or os.path.exists(base + ".key.csv"):
            raise SystemExit(f"check: {base} exists; use another --seed")
        tabs.append({"tab": name[:31], "rows": sheet, "scale": scale, "assets": assets,
                     "types": set_types(name, data), "measure": name})
        keys.append((base + ".key.csv", KEY_COLS, key))
    lkey, lsheet = link_rows(data, seed, kw.get("projects"),
                             LINK_FLOOR_BANDS if design == "floor" else LINK_BANDS)
    if lsheet:
        tabs.append({"tab": "package links", "rows": lsheet})
        keys.append((where(data, "review", f"check_links_s{seed}.key.csv"), LINK_KEY_COLS, lkey))
    # The decisions after the measures, each with its own tab (src/judged.py).
    import judged
    for tab, fname, rows_of, kcols, cols, answers in (
            ("pair checks", "check_pairs", judged.pair_rows, judged.PAIR_KEY_COLS,
             judged.PAIR_SHEET, judged.PAIR_ANSWERS),
            ("package assets", "check_package_assets", judged.package_asset_rows,
             judged.ASSET_KEY_COLS, judged.ASSET_SHEET, ["yes", "no", "unsure"])):
        k, sheet = rows_of(data, seed, kw.get("projects"))
        if sheet:
            tabs.append({"tab": tab, "rows": sheet, "cols": cols, "drops": [("answer", answers)]})
            keys.append((where(data, "review", f"{fname}_s{seed}.key.csv"), kcols, k))
    os.makedirs(os.path.dirname(xlsx), exist_ok=True)
    for path, cols, rows in keys:
        _write_key(path, cols, rows)
    _write_xlsx(xlsx, tabs)
    return xlsx, sum(len(t["rows"]) for t in tabs)


def link_question_sha():
    import link_packages
    return hashlib.sha256(json.dumps(link_packages.PAIR_QUESTION, sort_keys=True)
                          .encode("utf-8")).hexdigest()


def link_rows(data, seed, projects=None, bands=LINK_BANDS):
    """(key rows, sheet rows) of the package links drawn by band of the
    link's probability, from the current link run; ([], []) without one."""
    import link_packages
    folder = link_packages.current_run(data, projects=projects)
    if not folder:
        print("check: no current link run; the package links tab is left out")
        return [], []
    with open(os.path.join(folder, "commitment_packages.csv"), newline="", encoding="utf-8") as fh:
        pairs = list(csv.DictReader(fh))
    units = {(name, u["unit_id"]): u
             for name, (_, us) in link_packages.committed_units(data, projects).items() for u in us}
    packages = {p["package_id"]: p for p in link_packages.load_packages(data, projects)}
    names = {}
    import labels
    for name in {p["label_set"] for p in pairs}:
        try:
            names[name] = labels.load_set(name)[0]["name"]
        except (SystemExit, IndexError):
            names[name] = name
    pairs = [p for p in pairs if (p["label_set"], p["unit_id"]) in units
             and p["package_id"] in packages and p.get("probability", "") != ""]
    for p in pairs:                  # the decision reads the posterior (src/judged.py)
        p["probability"] = p.get("p_calibrated") or p["probability"]
    picked = draw_links(pairs, seed, bands)
    random.Random(seed).shuffle(picked)
    qsha = link_question_sha()
    key, sheet = [], []
    for i, (p, stratum, size, drawn) in enumerate(picked, 1):
        u, pk = units[(p["label_set"], p["unit_id"])], packages[p["package_id"]]
        state = link_packages.pair_state(u, pk)
        key.append({"row_id": i, "label_set": p["label_set"], "unit_id": p["unit_id"],
                    "package_id": p["package_id"],
                    "text_sha256": text_key(u["text"]),
                    "package_sha256": pk["sha"], "question_sha256": qsha,
                    "probability": p["probability"], "stratum": stratum, "stratum_size": size,
                    "stratum_drawn": drawn, "link_run": os.path.basename(folder)})
        sheet.append({"row_id": i, "project": p["project_id"],
                      "measure": names.get(p["label_set"], p["label_set"]),
                      "commitment": u["text"],
                      "commitment component": state.get("commitment_component", ""),
                      "package": pk["text"], "package component": pk["component"]})
    return key, sheet


def draw_links(pairs, seed, bands=LINK_BANDS):
    """Commitment and package pairs by band of the link probability:
    (pair, stratum, stratum size, rows drawn)."""
    rng = random.Random(seed)
    strata = {}
    for p in pairs:
        x = float(p["probability"])
        s = next((b[0] for b in bands if b[1] <= x < b[2]), None)
        if s is not None:
            strata.setdefault(s, []).append(p)
    picked = []
    for s, _, _, per in bands:
        pool = sorted(strata.get(s, []), key=lambda p: (p["label_set"], p["unit_id"], p["package_id"]))
        n = min(per, len(pool))
        for p in rng.sample(pool, n):
            picked.append((p, s, len(pool), n))
    return picked


def _write_xlsx(path, tabs):
    """tabs: [{tab, rows, and for a measure: scale, assets, types, measure}].
    A measure tab has dropdowns for grade, asset and type; the package links
    tab one for applies. The dropdown lists sit on a hidden sheet (a list
    written into the validation itself is capped at 255 characters)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    wb.active.title = "how to"
    for line in HOW_TO:
        wb.active.append([line])
    wb.active.column_dimensions["A"].width = 120
    lists = wb.create_sheet("lists")
    lists.sheet_state = "hidden"
    ncol = [0]

    def listed(values):
        ncol[0] += 1
        c = lists.cell(1, ncol[0]).column_letter
        for i, v in enumerate(values, 1):
            lists.cell(i, ncol[0], cell_text(v))
        return f"lists!${c}$1:${c}${len(values)}"

    wrap = Alignment(wrap_text=True, vertical="top")
    for t in tabs:
        spec = t.get("cols") or (LINK_SHEET_COLS if "scale" not in t else SHEET_COLS)
        ws = wb.create_sheet(t["tab"])
        cols = [c for c, _ in spec]
        ws.append(cols)
        for r in t["rows"]:
            ws.append([cell_text(r.get(c)) for c in cols])
        for i, (_, width) in enumerate(spec, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        n = len(t["rows"]) + 1
        if t.get("drops"):
            drops = t["drops"]
        elif "scale" in t:
            drops = [("grade", [str(i) for i in range(len(t["scale"]))]),
                     ("asset", t["assets"] + ["none", "other"]),
                     ("type", [x[0] for x in t["types"]] + ["none", "other"])]
        else:
            drops = [("applies", ["yes", "no", "unsure"])]
        for col, values in drops:
            dv = DataValidation(type="list", formula1=listed(values), allow_blank=True)
            if col in ("asset", "type"):
                dv.showErrorMessage = False     # several may be typed, joined by |
            ws.add_data_validation(dv)
            c = ws.cell(1, cols.index(col) + 1).column_letter
            dv.add(f"{c}2:{c}{n}")
        ws.freeze_panes = "A2"
    measures = [t for t in tabs if "scale" in t]
    s = wb.create_sheet("grades")
    s.append(["measure", "grade", "meaning ({name} is the measure)"])
    for t in measures:
        for i, line in enumerate(t["scale"]):
            s.append([t["measure"], i, cell_text(line)])
    s.column_dimensions["A"].width = 26
    s.column_dimensions["C"].width = 120
    if any(t["types"] for t in measures):
        s = wb.create_sheet("types")
        s.append(["measure", "type", "name", "definition"])
        for t in measures:
            for tid, nm, d in t["types"]:
                s.append([t["measure"], cell_text(tid), cell_text(nm), cell_text(d)])
        for c, w in zip("ABCD", (26, 30, 30, 100)):
            s.column_dimensions[c].width = w
    wb.save(path)


def _sheet_values(xlsx, tab):
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, read_only=True)[tab]
    rows = ws.iter_rows(values_only=True)
    head = list(next(rows))
    return head, rows


def _clean(v):
    return str(v).strip() if v is not None else ""


def read_sheet(xlsx, tab="grade these"):
    """{row_id: (grade, asset, type)} for the rows a person graded."""
    head, rows = _sheet_values(xlsx, tab)
    gi, ai, ri = head.index("grade"), head.index("asset"), head.index("row_id")
    ti = head.index("type") if "type" in head else None
    out = {}
    for r in rows:
        if _clean(r[gi]) == "":
            continue
        out[str(r[ri])] = (int(float(_clean(r[gi]))), _clean(r[ai]),
                           _clean(r[ti]) if ti is not None else "")
    return out


def read_links(xlsx, tab="package links"):
    """{row_id: yes|no|unsure} for the links a person answered."""
    head, rows = _sheet_values(xlsx, tab)
    ai, ri = head.index("applies"), head.index("row_id")
    return {str(r[ri]): _clean(r[ai]).lower() for r in rows if _clean(r[ai])}


def graded_sheet(name, data, seed, elusion=False):
    """(key path, {row_id: (grade, asset, type)}, sheet name) of a measure's
    sheet: its own workbook, else its tab in the combined one."""
    base = where(data, "review", f"check_units_{name}{'_elusion' if elusion else ''}_s{seed}")
    if os.path.exists(base + ".xlsx"):
        return base + ".key.csv", read_sheet(base + ".xlsx"), os.path.basename(base)
    import combined
    if not elusion and combined.is_combined(data, seed) and os.path.exists(base + ".key.csv"):
        return (base + ".key.csv", combined.Graded(data, seed).measure_grades(name),
                os.path.basename(base))
    xlsx = all_path(data, seed)
    if not elusion and os.path.exists(xlsx) and os.path.exists(base + ".key.csv"):
        return base + ".key.csv", read_sheet(xlsx, name[:31]), os.path.basename(base)
    raise SystemExit(f"check: no sheet for {name} at seed {seed} ({base}.xlsx or {xlsx})")


from schema import QUALITY_COLS  # noqa: E402  (one list, loaded in this shape)


def write_quality(data, name, key, rows, floor):
    """The headline numbers of an evaluation, one row per measure and exact
    question in data/labels/measure_quality.csv (loaded as
    wbg.measure_quality): the weighted precision and recall of the
    pipeline's decision (posterior at the floor) against the person's grade
    (2 or more), and which kind of sheet they come from."""
    import posterior
    ys, ws = [r["yes"] for r in rows], [r["w"] for r in rows]
    prec, rec, _ = posterior.weighted_pr([r["p"] for r in rows], ys, ws, floor)
    path = where(data, "labels", "measure_quality.csv")
    current = {}
    if os.path.exists(path):
        with open(path, newline="", encoding="utf-8") as fh:
            current = {(r["label_set"], r["question_sha256"], r.get("sheet", "")): r
                       for r in csv.DictReader(fh)}
    import labels
    version = next((lab.get("version") or "" for lab in labels.load_set(name)), "")
    sheet = "evaluation" if all(r["stratum"].startswith(EVAL) for r in rows) else "training"
    current[(name, key["question_sha256"], sheet)] = {
        "label_set": name, "version": version, "question_sha256": key["question_sha256"],
        "run_id": key["measure_dir"], "sheet": sheet, "graded_units": len(rows),
        "precision": f"{prec:.3f}" if prec is not None else "",
        "recall": f"{rec:.3f}" if rec is not None else "",
        "evaluated_at": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "report_path": f"reports/evaluate_units_{name}.md"}
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path + ".part", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=QUALITY_COLS, lineterminator="\n", restval="")
        w.writeheader()
        w.writerows(sorted(current.values(), key=lambda r: (r["label_set"], r["evaluated_at"])))
    os.replace(path + ".part", path)


def posterior_path(name, data):
    return where(data, "labels", f"{name}_posterior.json")


def text_key(text):
    """The SHA-256 a label is keyed by: of the unit's text as the sheet showed
    it (with its lead-in)."""
    return hashlib.sha256((text or "").encode("utf-8")).hexdigest()


def graded_labels(name, data, question_sha256=None, fitting=False, units_by_text=None):
    """{unit_id: (grade, asset, weight)} of every sentence a person graded for
    the set under the exact question given (every question when None; the
    latest grade wins). A grade given under another wording of the label is
    a grade of another question (AGENTS.md rule 9), so the layer and the
    skip rule read only the current one. weight is the row's design weight
    (stratum size over rows drawn; 1 where the sheet did not record it).
    fitting=True leaves out the rows of evaluation sheets: what fits a layer
    or a gate never grades it. units_by_text ({text SHA-256: unit_id} of the
    corpus as it stands) finds each grade's sentence by its text, never by
    the stored position: a re-parse moves unit ids (AGENTS.md rule 9), and a
    grade whose text is no longer in the corpus is left out."""
    lpath = where(data, "labels", f"{name}_units.csv")
    out = {}
    if os.path.exists(lpath):
        with open(lpath, newline="", encoding="utf-8") as fh:
            for r in sorted(csv.DictReader(fh), key=lambda r: r.get("graded_on", "")):
                if r.get("grade", "") == "":
                    continue
                if question_sha256 is not None and r.get("question_sha256") != question_sha256:
                    continue
                if fitting and (r.get("stratum") or "").startswith(EVAL):
                    continue
                uid = r["unit_id"] if units_by_text is None else \
                    units_by_text.get(r.get("text_sha256", ""))
                if uid is None:
                    continue
                out[uid] = (int(r["grade"]), r.get("asset", ""), float(r.get("weight") or 1.0))
    return out


def latest_run(name, data):
    """(units.csv rows by unit_id, question SHA-256) of the newest measure run
    asking the set's question as it stands now, or (None, sha)."""
    import glob
    import measure
    run = measure.Run(name, data, log=lambda *_: None, build=False)
    lid = run.labels[0]["id"]
    sha = classify.question_shas(run.qs)[lid]
    for m in sorted(glob.glob(where(data, "measures", f"{glob.escape(name)}-*", "manifest.json")),
                    key=os.path.getmtime, reverse=True):
        with open(m, encoding="utf-8") as fh:
            if json.load(fh).get("question_sha256", {}).get(lid) != sha:
                continue
        with open(os.path.join(os.path.dirname(m), "units.csv"), newline="", encoding="utf-8") as fh:
            return {u["unit_id"]: u for u in csv.DictReader(fh)}, sha
    return None, sha


def layer_block(name, data):
    """What the measure's probability layer is fitted on: (graded rows
    [(unit_id, 0/1, design weight, projects)], {unit_id: model score},
    {unit_id: vector}, question SHA-256, {unit_id: [P(0)..P(3)]}) from the
    sentences a person graded under the question the set asks now and the
    newest run asking it; None with no such run."""
    import panel, posterior
    lo = min(panel.grades())
    units, sha = latest_run(name, data)
    if not units:
        return None
    graded = graded_labels(name, data, sha, fitting=True,
                           units_by_text={text_key(u.get("text")): uid for uid, u in units.items()})
    vector_of, projects_of = posterior.lookups(data)
    asked = [u for u in graded if u in units and units[u].get("sent") == "true"]
    score = {u: sum(float(units[u].get(f"p{g}") or 0) for g in panel.grades()) for u in asked}
    probs = {u: [float(units[u].get(f"p{g}") or 0) for g in range(4)] for u in asked}
    vecs = {u: vector_of(u) for u in score}
    vecs = {u: v for u, v in vecs.items() if v is not None}
    rows = [(u, 1 if graded[u][0] >= lo else 0, graded[u][2], projects_of(u)) for u in score]
    return rows, score, vecs, sha, probs


def layer_lines(name, rows, score, probs, head):
    """The report on one measure's layer: precision and recall out of sample
    against the model's score alone, and the floor the panel's target
    precision suggests."""
    import panel, posterior
    target = float(panel.get("floor.target_precision", 0.8))
    floor = panel.floor(name)
    ys, ws = [r[1] for r in rows], [r[2] for r in rows]
    raw = [score[u] for u, *_ in rows]
    probs = [probs[u] for u, *_ in rows]
    neff = posterior.effective_n(ws)
    f = lambda x: "-" if x is None else f"{x:.2f}"
    lines = head + ["", f"Out of sample (folds by project), weighted; effective sample size "
                    f"{neff:.0f} of {len(rows)} rows.", "",
                    "| floor | precision, model's score | recall | precision, layer | recall |",
                    "|---|---|---|---|---|"]
    for t in sorted({0.3, 0.4, 0.5, 0.6, 0.7, 0.8, floor or 0.5}):
        p1, r1, _ = posterior.weighted_pr(raw, ys, ws, t)
        p2, r2, _ = posterior.weighted_pr(probs, ys, ws, t)
        lines.append(f"| {t:.2f} | {f(p1)} | {f(r1)} | {f(p2)} | {f(r2)} |")
    t = posterior.suggest_floor(probs, ys, target, ws)
    if t:
        lines += ["", f"- suggested floor (floor.{name}): {t[0]:.2f}, weighted precision "
                  f"{t[1]:.2f}, recall {t[2]:.2f} out of sample (target precision {target}); "
                  f"the panel now reads {floor}"]
    return lines + [""]


def fit_layer(name, data):
    """Fit the measure's probability layer (src/posterior.py) on the
    sentences a person graded under the question the set asks now, write it
    for the next run, and report how it does out of sample (folds by
    project, every row at its design weight) and the floor the panel's
    target precision suggests on those out-of-fold probabilities. With
    posterior.pool the layer is fitted with every measure's grades
    (fit_pooled_layers). Returns report lines."""
    import panel, posterior
    if posterior.settings()["pool"]:
        return fit_pooled_layers(data).get(name, ["## Probability layer and suggested floor", "",
                                                  "- no layer in the pool: the model's score "
                                                  "stands", ""])
    lo = min(panel.grades())
    lines = ["## Probability layer and suggested floor", ""]
    block = layer_block(name, data)
    if not block:
        return lines + ["- no run asks the current question yet: run the measure, then evaluate again", ""]
    rows, score, vecs, sha, _ = block
    layer = posterior.fit(rows, score, vecs, sha, name)
    if layer is None:
        if os.path.exists(posterior_path(name, data)):
            os.remove(posterior_path(name, data))
        probs = score
        lines.append(f"- too few sentences graded under the current question ({len(rows)}, "
                     f"{sum(r[1] for r in rows)} yes) for a layer: the model's score stands")
    else:
        posterior.write(posterior_path(name, data), layer)
        probs = posterior.out_of_fold(rows, score, vecs)
        lines.append(f"- layer fitted on {layer['n_yes']} graded yes and {layer['n_no']} no "
                     f"(grade {lo} or more = yes; current question only), written to "
                     f"{os.path.basename(posterior_path(name, data))}")
    return layer_lines(name, rows, score, probs, lines)


def layer_blocks(data):
    """{set: layer_block} for every measure with a run asking its current
    question."""
    import labels
    out = {}
    for name in labels.label_sets():
        if name == measure.ASSET_SET:
            continue
        try:
            b = layer_block(name, data)
        except SystemExit:            # a set the measure stage cannot run
            continue
        if b:
            out[name] = b
    return out


def fit_pooled_layers(data, blocks=None):
    """The pooled layer (posterior.pool): every measure's grades fitted at
    once, one intercept per measure and the weights shared; each measure's
    layer written to its file (a measure that cannot join loses its file)
    and checked out of sample with the same project folds. Returns {set:
    report lines}."""
    import panel, posterior
    lo = min(panel.grades())
    blocks = layer_blocks(data) if blocks is None else blocks
    four = {n: b[:4] for n, b in blocks.items()}
    specs = posterior.fit_pooled(four)
    oof = posterior.out_of_fold_all(four, dict(posterior.settings(), pool=True)) if specs else {}
    out = {}
    for name, (rows, score, vecs, sha, _) in sorted(blocks.items()):
        path = posterior_path(name, data)
        lines = ["## Probability layer and suggested floor", ""]
        if name in specs:
            posterior.write(path, specs[name])
            lines.append(f"- pooled layer: this measure's {specs[name]['n_yes']} graded yes and "
                         f"{specs[name]['n_no']} no (grade {lo} or more = yes; current question "
                         f"only) with {len(specs) - 1} other measures, its own intercept and "
                         f"the shared weights; written to {os.path.basename(path)}")
            probs = oof[name]
        else:
            if os.path.exists(path):
                os.remove(path)
            lines.append(f"- too few sentences graded under the current question ({len(rows)}, "
                         f"{sum(r[1] for r in rows)} yes) to join the pooled layer: the model's "
                         f"score stands")
            probs = score
        out[name] = layer_lines(name, rows, score, probs, lines)
    return out


def fit_gate(name, data):
    """The skip rule's gate for the measure (src/skip.py): from the sentences
    a person graded yes under the current question (pool, feedback and
    elusion sheets alike), each at its design weight, the highest best-cosine
    gate that loses at most skip.max_loss of them; written for the next run.
    Returns report lines."""
    import panel, skip
    lo = min(panel.grades())
    run = measure.Run(name, data, build=False, log=lambda *_: None)
    run.load_corpus()
    graded = graded_labels(name, data, run.question_sha, fitting=True,
                           units_by_text={text_key(t): u for u, t in run.texts.items()})
    lines = ["## Skip rule (what the pool leaves out)", ""]
    yes = {u: w for u, (g, _, w) in graded.items() if g >= lo}
    if not yes:
        return lines + ["- no sentence graded yes under the current question: the panel's "
                        f"gate ({measure.MIN_COSINE}) stands", ""]
    lid = run.labels[0]["id"]
    mine = [p for p in run.patterns if p["label_id"] == lid]
    pos = []
    for u, w in sorted(yes.items()):
        if u not in run.best_cosine:
            continue
        text = run.texts.get(u, "")
        pos.append((run.best_cosine[u], any(p["regex"].search(text) for p in mine), w))
    got = skip.fit(pos)
    path = skip.path(data, name)
    if got is None:
        if os.path.exists(path):
            os.remove(path)
        return lines + [f"- {len(pos)} graded yes, too few or all matched by a pattern: the "
                        f"panel's gate ({measure.MIN_COSINE}) stands", ""]
    gate, lost, n = got
    bound = skip.loss_bound(lost, [p[2] for p in pos])
    skip.write(data, name, {lid: {"gate": gate, "lost": lost, "lost_upper_95": bound,
                                  "positives": n, "question_sha256": run.question_sha,
                                  "search_run": os.path.basename(run.search_run)}})
    return lines + [f"- gate {gate:.4f} on the best cosine to the phrases: of {n} graded yes "
                    f"({sum(1 for p in pos if p[1])} matched by a pattern, always asked), "
                    f"{lost:.1%} of their weight falls below it (rule: at most "
                    f"{skip.settings()['max_loss']:.0%}); the panel's starting gate was "
                    f"{measure.MIN_COSINE}. The next run's pool uses it.",
                    f"- what the grades can vouch for: the loss is at most {bound:.0%} "
                    f"(95 per cent upper bound on {n} yes); the elusion sheet checks below the "
                    f"gate.", ""]


def fit_asset_layers(data):
    """One probability layer per asset (src/posterior.py), fitted on the
    assets people gave the sentences they graded for any measure (yes = the
    person named this asset), against the current sentence asset map, and
    written to data/labels/assets_posterior.json for the measure stage.
    Returns report lines."""
    import labels
    import panel, posterior
    folder = classify.current_run("assets", data) or classify.last_run("assets", data)
    if not folder:
        return ["- asset layers: no sentence asset map yet"]
    with open(os.path.join(folder, "manifest.json"), encoding="utf-8") as fh:
        man = json.load(fh)
    import csv as _csv
    with open(where(data, "embeddings", "sentence_index.csv"), newline="", encoding="utf-8") as fh:
        texts = measure.unit_texts(data, list(_csv.DictReader(fh)))
    by_text = {text_key(t): u for u, t in texts.items()}
    # The assets a person gave a sentence on a measure's own sheet: every
    # asset asked of that sentence is graded by it, given or not ...
    named, weight = {}, {}
    for name in labels.label_sets():
        if name == measure.ASSET_SET:
            continue
        for u, (_, asset, w) in graded_labels(name, data, fitting=True,
                                              units_by_text=by_text).items():
            if asset:
                named[u] = set(asset.split("|")) - {""}   # the assets a person gave it
                weight[u] = w
    unsure = unsure_assets(data, by_text)
    # ... and the combined sheets' asset lines, each graded on its own (a
    # line not drawn is no grade, never a no)
    own = asset_line_grades(data, by_text, man["question_sha256"])
    graded, scores = {}, {}
    with open(os.path.join(folder, "scores.csv"), newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            u, a = r.get("unit_id"), r["label_id"]
            if r.get("asked") != "true":
                continue
            if (u, a) in own:
                graded[(u, a)] = own[(u, a)]
            elif u in named and a not in unsure.get(u, ()):
                graded[(u, a)] = (int(a in named[u]), weight[u])
            else:
                continue
            scores.setdefault(a, {})[u] = sum(float(r.get(f"p{g}") or 0) for g in panel.grades())
    vector_of, projects_of = posterior.lookups(data)
    vecs = {u: vector_of(u) for u, _ in graded}
    vecs = {u: v for u, v in vecs.items() if v is not None}
    out, lines = {}, ["## Asset probability layers", ""]
    blocks = {asset: ([(u, graded[(u, asset)][0], graded[(u, asset)][1], projects_of(u))
                       for u in sc], sc, vecs, man["question_sha256"].get(asset, ""))
              for asset, sc in sorted(scores.items())}
    # pooled across assets when posterior.pool is on, as the measures are
    pooled = posterior.fit_pooled(blocks) if posterior.settings()["pool"] else None
    for asset, (ys, sc, _, qsha) in blocks.items():
        layer = pooled.get(asset) if pooled is not None else \
            posterior.fit(ys, sc, vecs, qsha, asset)
        if layer:
            out[asset] = layer
            lines.append(f"- {asset}: fitted on {layer['n_yes']} yes, {layer['n_no']} no")
        else:
            lines.append(f"- {asset}: too few graded sentences ({len(ys)}, "
                         f"{sum(y for _, y, *_ in ys)} yes); the map's score stands")
    path = classify.layers_path(data)
    if out:
        posterior.write(path, out)
    elif os.path.exists(path):
        os.remove(path)
    # The skip rule for each asset (src/skip.py), as for a measure: the gate
    # on a sentence's own cosine; a yes a word pattern of the asset matches
    # is always asked, so it sets no bound.
    import search, skip
    cos = classify.best_cosines("assets", data)
    folder_s, _ = search.run_folder("assets", data)
    pats = labels.load_patterns("assets", taxonomy=classify.TAXONOMY)
    gates = {}
    for asset in sorted(scores):
        mine = [p["regex"] for p in pats if p["label_id"] == asset]
        pos = [(cos[(asset, u)], any(rx.search(texts.get(u, "")) for rx in mine), w)
               for (u, a), (y, w) in sorted(graded.items())
               if a == asset and y and (asset, u) in cos]
        got = skip.fit(pos)
        if got:
            bound = skip.loss_bound(got[1], [p[2] for p in pos])
            gates[asset] = {"gate": got[0], "lost": got[1], "lost_upper_95": bound,
                            "positives": got[2],
                            "question_sha256": man["question_sha256"].get(asset, ""),
                            "search_run": os.path.basename(folder_s)}
            lines.append(f"- {asset}: skip-rule gate {got[0]:.4f} ({got[2]} graded yes, "
                         f"{got[1]:.1%} below, at most {bound:.0%} at 95 per cent)")
    gpath = skip.path(data, "assets")
    if gates:
        skip.write(data, "assets", gates)
    elif os.path.exists(gpath):
        os.remove(gpath)
    return lines + [""]


def is_eval(r):
    """Whether a key or label row is an evaluation sheet's."""
    return (r.get("stratum") or "").startswith(EVAL)


def with_sheet(rows):
    """Each label row with its kind of sheet (part of its key in Postgres)."""
    for r in rows:
        r["sheet"] = "evaluation" if is_eval(r) else "training"
    return rows


def asset_line_grades(data, by_text, qshas, fitting=True):
    """{(unit_id, asset): (1 if graded 2 or more else 0, weight)} of the
    combined sheets' asset lines (data/labels/assets_units.csv, label_id
    the asset) under each asset's current question; with fitting, an
    evaluation sheet's are left out."""
    import panel
    lo = min(panel.grades())
    path = where(data, "labels", f"{measure.ASSET_SET}_units.csv")
    out = {}
    if os.path.exists(path):
        with open(path, newline="", encoding="utf-8") as fh:
            for r in sorted(csv.DictReader(fh), key=lambda r: r.get("graded_on", "")):
                a, u = r.get("label_id", ""), by_text.get(r.get("text_sha256", ""))
                if u is None or a not in qshas or r.get("question_sha256") != qshas[a] \
                        or r.get("grade", "") == "" or (fitting and is_eval(r)):
                    continue
                out[(u, a)] = (int(int(r["grade"]) >= lo), float(r.get("weight") or 1.0))
    return out


def unsure_assets(data, by_text):
    """{unit_id: assets a person left unsure} (data/labels/assets_units.csv):
    those questions are left out, not read as no."""
    path = where(data, "labels", f"{measure.ASSET_SET}_units.csv")
    out = {}
    if os.path.exists(path):
        with open(path, newline="", encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                u = by_text.get(r.get("text_sha256", ""))
                if u and r.get("unsure") and not is_eval(r):
                    out[u] = set(r["unsure"].split("|")) - {""}
    return out


def store_labels(name, data, key, graded, source):
    """The person's grades as labels in data/labels/<set>_units.csv, keyed by
    text, question and kind of sheet, never by position; a later grade of the
    same text under the same question on the same kind of sheet replaces an
    earlier one. A training and an evaluation grade of one sentence are kept
    apart, so neither changes what the other measures or fits."""
    lpath = where(data, "labels", f"{name}_units.csv")
    current = {}
    if os.path.exists(lpath):
        with open(lpath, newline="", encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                current[(r["label_id"], r["text_sha256"], r["question_sha256"],
                         is_eval(r))] = r
    today = datetime.date.today().isoformat()
    for k in key:
        if k["row_id"] in graded:
            g, a, t, *rest = graded[k["row_id"]]
            current[(k["label_id"], k["text_sha256"], k["question_sha256"], is_eval(k))] = {
                "label_set": name, "label_id": k["label_id"], "unit_id": k["unit_id"],
                "text_sha256": k["text_sha256"], "question_sha256": k["question_sha256"],
                "grade": g, "asset": a, "type": t, "source": source, "graded_on": today,
                "stratum": k.get("stratum", ""),
                "weight": f"{key_weight(k):.6g}"
                if k.get("weight") or (k.get("stratum_size") and k.get("stratum_drawn")) else "",
                "unsure": rest[0] if rest else ""}
    os.makedirs(os.path.dirname(lpath), exist_ok=True)
    with open(lpath + ".part", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=LABEL_COLS, lineterminator="\n", restval="")
        w.writeheader()
        w.writerows(sorted(with_sheet(list(current.values())),
                           key=lambda r: (r["label_id"], r["text_sha256"], r["sheet"])))
    os.replace(lpath + ".part", lpath)


def assigned_types(name, data):
    """{unit_id: set of types that count} of the set's current type
    assignment ({'none'} for a unit with none), or {}."""
    import kinds
    folder = kinds.current_assign(name, data)
    if not folder:
        return {}
    with open(os.path.join(folder, "assignments.csv"), newline="", encoding="utf-8") as fh:
        return {r["unit_id"]: set((r.get("types_yes") or "").split("|")) - {""} or {"none"}
                for r in csv.DictReader(fh)}


def evaluate(name, data, seed=1, elusion=False, layer=True, **kw):
    if name == ALL:
        return evaluate_all(data, seed, **kw)
    if elusion:
        return evaluate_elusion(name, data, seed, **kw)
    kpath, graded, source = graded_sheet(name, data, seed)
    with open(kpath, newline="", encoding="utf-8") as fh:
        key = list(csv.DictReader(fh))
    if not graded:
        raise SystemExit(f"check: no grade entered in {source}")
    store_labels(name, data, key, graded, source)
    # measure the run the sheet was drawn from
    mdir = where(data, "measures", key[0]["measure_dir"])
    with open(os.path.join(mdir, "units.csv"), newline="", encoding="utf-8") as fh:
        units = {u["unit_id"]: u for u in csv.DictReader(fh)}
    import panel, posterior
    lo = min(panel.grades())
    floor = panel.floor(name)
    rows = []
    for k in key:
        # a combined sheet's 'outside' rows are stored (the skip rule reads
        # them) but are no decision of this run's
        if k["row_id"] not in graded or k.get("frame") == "outside":
            continue
        u = units[k["unit_id"]]
        g, a, t, *_ = graded[k["row_id"]]
        rows.append({"grade": g, "yes": g >= lo, "asset": a, "type": t, "unit_id": k["unit_id"],
                     "w": key_weight(k),
                     "p": _posterior(u), "model": measure.score_of(u),
                     "committed": u.get("committed") == "true",
                     "link": u.get("asset_link", ""),
                     "links": set((u.get("asset_links") or u.get("asset_link") or "").split("|")) - {""},
                     "stratum": k["stratum"]})
    if not rows:
        raise SystemExit(f"check: no graded row of {source} is a sentence {name} read")
    is_eval = all(r["stratum"].startswith(EVAL) for r in rows)
    lines = [f"# {name}: the pipeline's decision against {len(rows)} graded units", "",
             f"Drawn from {key[0]['measure_dir']} by stratum "
             f"({', '.join(sorted({k['stratum'] for k in key if k.get('frame') != 'outside'}))})"
             f". Yes = the person graded "
             f"{lo} or more; the pipeline says yes when the posterior reaches the floor "
             f"({floor}). Weighted estimates use each row's design weight; "
             f"intervals are 95 per cent Wilson on the effective sample size. Recall is "
             f"within the graded pool (candidates and feedback units), not the corpus.",
             "", ("An evaluation sheet: a random draw by the decision, never used to fit; "
                  "its weighted figures are the dashboards' precision and recall within the pool."
                  if is_eval else
                  "A training sheet: drawn where grades teach most, so its figures describe the "
                  "hard cases, not the dashboards (an evaluation sheet measures those)."), "",
             "| pipeline says yes at | weighted precision (95% CI) | weighted recall (95% CI) "
             "| precision, model score alone | recall | rows kept |", "|---|---|---|---|---|---|"]
    ys, ws = [r["yes"] for r in rows], [r["w"] for r in rows]
    f = lambda x: "-" if x is None else f"{x:.2f}"
    for t in sorted(set(THRESHOLDS) | {floor}):
        p1, r1, kept = posterior.weighted_pr([r["p"] for r in rows], ys, ws, t)
        p2, r2, _ = posterior.weighted_pr([r["model"] for r in rows], ys, ws, t)
        nk = posterior.effective_n([r["w"] for r in rows if r["p"] >= t])
        npos = posterior.effective_n([r["w"] for r in rows if r["yes"]])
        pl, ph = wilson(round(p1 * nk), round(nk)) if p1 is not None and nk else (None, None)
        rl, rh = wilson(round(r1 * npos), round(npos)) if r1 is not None and npos else (None, None)
        mark = " (the floor)" if t == floor else ""
        lines.append(f"| posterior >= {t}{mark} | {f(p1)} ({f(pl)}-{f(ph)}) | {f(r1)} "
                     f"({f(rl)}-{f(rh)}) | {f(p2)} | {f(r2)} | "
                     f"{sum(1 for r in rows if r['p'] >= t)} |")
    lines.append("")
    withasset = [r for r in rows if r["asset"] and r["yes"]]
    if withasset:
        given = lambda r: set(r["asset"].split("|")) - {""}
        agree = sum(1 for r in withasset if given(r) & r["links"])
        extra = sum(len(r["links"] - given(r)) for r in withasset)
        lines += [f"## Asset on units the person graded yes and gave an asset: "
                  f"{agree} of {len(withasset)} have it among the pipeline's assets "
                  f"({extra} further assets the pipeline gave them)", ""]
    withtype = [r for r in rows if r["type"] and r["yes"]]
    if withtype:
        best = assigned_types(name, data)
        typed = [r for r in withtype if r["unit_id"] in best]
        given = lambda r: set(r["type"].split("|")) - {""}
        agree = sum(1 for r in typed if given(r) & best[r["unit_id"]])
        extra = sum(len(best[r["unit_id"]] - given(r) - {"none"}) for r in typed)
        lines += [f"## Type on units the person graded yes and gave a type: {agree} of "
                  f"{len(typed)} share a type with the pipeline's ({extra} further types the "
                  f"pipeline gave them; {len(withtype) - len(typed)} not committed by the "
                  f"pipeline, so not typed)", ""]
    by = {}
    for r in rows:
        by.setdefault(r["stratum"], []).append(r)
    lines += ["## By stratum: rows, graded yes", ""]
    for st, rs in sorted(by.items()):
        lines.append(f"- {st}: {len(rs)} rows, graded yes {sum(1 for r in rs if r['yes'])}")
    lines.append("")
    write_quality(data, name, key[0], rows, floor)
    lines += fit_layer(name, data) if layer else []
    lines += fit_gate(name, data)
    out = where(data, "reports", f"evaluate_units_{name}.md")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return out, len(rows)


def evaluate_all(data, seed, **kw):
    """Every measure tab of the combined workbook a person graded, then its
    package links. Returns (reports written, graded rows)."""
    import labels
    xlsx = all_path(data, seed)
    if not os.path.exists(xlsx):
        raise SystemExit(f"check: {xlsx} not found; draw it with ./wbg check all --seed {seed}")
    import combined
    import posterior, tune
    pool = posterior.settings()["pool"]
    outs, n, report = [], 0, {}
    both = combined.Graded(data, seed) if combined.is_combined(data, seed) else None
    if both:
        print(f"  assets: {combined.store_assets(data, seed, both)} sentences stored")
    for name in labels.label_sets():
        if not os.path.exists(where(data, "review", f"check_units_{name}_s{seed}.key.csv")):
            continue
        try:
            out, k = evaluate(name, data, seed, layer=not pool, **kw)
        except SystemExit as exc:
            print(f"  {name}: {exc}")
            continue
        outs.append(out)
        report[name] = out
        n += k
    # Pooled, the layer is fitted once every sheet's grades are stored, and
    # its section is added to each measure's report.
    if pool:
        for name, lines in fit_pooled_layers(data).items():
            if name in report:
                with open(report[name], "a", encoding="utf-8") as fh:
                    fh.write("\n".join(lines) + "\n")
    try:
        lines = fit_asset_layers(data)
        print("\n".join(lines))
    except SystemExit as exc:
        print(f"  asset layers: {exc}")
    if os.path.exists(where(data, "review", f"check_links_s{seed}.key.csv")):
        try:
            out, k = evaluate_links(data, seed,
                                    both.package_answers("package_link") if both else None)
            outs.append(out)
            n += k
        except SystemExit as exc:
            print(f"  package links: {exc}")
    # The decisions after the measures: their tabs' answers, layers and
    # floors (src/judged.py).
    import judged
    try:
        outs.append(judged.evaluate(data, seed, {
            "pairs": both.pair_answers(),
            "package_assets": both.package_answers("package_asset")} if both else None))
    except SystemExit as exc:
        print(f"  after the measures: {exc}")
    if both:
        outs.append(combined.report(data, seed, both))
    # What the grades say every tunable setting should be (src/tune.py).
    outs.append(tune.report(data))
    return ", ".join(outs) or "nothing", n


def evaluate_links(data, seed, answered=None):
    """The person's answers on the package links tab: stored as labels in
    data/labels/package_links.csv (keyed by both texts and the question), and
    the precision of the link at each probability threshold, weighted back
    to the bands, in data/reports/evaluate_links.md."""
    kpath = where(data, "review", f"check_links_s{seed}.key.csv")
    with open(kpath, newline="", encoding="utf-8") as fh:
        key = list(csv.DictReader(fh))
    if answered is None:
        answered = read_links(all_path(data, seed))
    if not answered:
        raise SystemExit("no link answered")
    lpath = where(data, "labels", "package_links.csv")
    current = {}
    if os.path.exists(lpath):
        with open(lpath, newline="", encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                current[(r["text_sha256"], r["package_sha256"], r["question_sha256"],
                         is_eval(r))] = r
    today = datetime.date.today().isoformat()
    for k in key:
        if k["row_id"] in answered:
            current[(k["text_sha256"], k["package_sha256"], k["question_sha256"],
                     is_eval(k))] = {
                **{c: k[c] for c in ("label_set", "unit_id", "package_id", "text_sha256",
                                      "package_sha256", "question_sha256")},
                "applies": answered[k["row_id"]],
                "weight": f"{key_weight(k):.6g}", "stratum": k.get("stratum", ""),
                "source": os.path.basename(kpath)[:-8], "graded_on": today}
    os.makedirs(os.path.dirname(lpath), exist_ok=True)
    with open(lpath + ".part", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=LINK_LABEL_COLS, lineterminator="\n", restval="")
        w.writeheader()
        w.writerows(sorted(with_sheet(list(current.values())),
                           key=lambda r: (r["text_sha256"], r["package_sha256"], r["sheet"])))
    os.replace(lpath + ".part", lpath)
    rows = [{"p": float(k["probability"]), "yes": answered[k["row_id"]] == "yes",
             "w": key_weight(k), "stratum": k["stratum"]}
            for k in key if answered.get(k["row_id"]) in ("yes", "no")]
    unsure = sum(1 for v in answered.values() if v not in ("yes", "no"))
    f = lambda x: "-" if x is None else f"{x:.2f}"
    lines = [f"# Package links: the decision model against {len(rows)} answered pairs", "",
             f"Drawn from {key[0]['link_run']} by band of the link probability; {unsure} "
             f"answered unsure are left out. Weighted figures use band size over rows drawn; "
             f"intervals are 95 per cent Wilson on the raw counts. Recall is among the pairs "
             f"the link asked, not every package.", "",
             "| link kept at P >= | precision (95% CI) | recall (95% CI) | weighted precision | "
             "weighted recall | n kept |", "|---|---|---|---|---|---|"]
    for t in [0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9]:
        tp = sum(1 for r in rows if r["p"] >= t and r["yes"])
        pp = sum(1 for r in rows if r["p"] >= t)
        ap = sum(1 for r in rows if r["yes"])
        wtp = sum(r["w"] for r in rows if r["p"] >= t and r["yes"])
        wpp = sum(r["w"] for r in rows if r["p"] >= t)
        wap = sum(r["w"] for r in rows if r["yes"])
        pl, ph = wilson(tp, pp)
        rl, rh = wilson(tp, ap)
        lines.append(f"| {t} | {f(tp / pp if pp else None)} ({f(pl)}-{f(ph)}) | "
                     f"{f(tp / ap if ap else None)} ({f(rl)}-{f(rh)}) | "
                     f"{f(wtp / wpp if wpp else None)} | {f(wtp / wap if wap else None)} | {pp} |")
    lines += ["", "## By band: pairs answered, share yes", ""]
    for s in sorted({r["stratum"] for r in rows}):
        rs = [r for r in rows if r["stratum"] == s]
        lines.append(f"- {s}: {len(rs)} answered, {sum(r['yes'] for r in rs)} yes")
    out = where(data, "reports", "evaluate_links.md")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return out, len(answered)


def evaluate_elusion(name, data, seed=1, **kw):
    """What the pool missed: the elusion sheet's grades, weighted back to
    their strata."""
    kpath, graded, source = graded_sheet(name, data, seed, elusion=True)
    with open(kpath, newline="", encoding="utf-8") as fh:
        key = list(csv.DictReader(fh))
    if not graded:
        raise SystemExit(f"check: no grade entered in {source}")
    store_labels(name, data, key, graded, source)
    mdir = where(data, "measures", key[0]["measure_dir"])
    with open(os.path.join(mdir, "manifest.json"), encoding="utf-8") as fh:
        man = json.load(fh)
    import panel
    yes_from = min(panel.grades())
    with open(os.path.join(mdir, "units.csv"), newline="", encoding="utf-8") as fh:
        in_pool = sum(1 for u in csv.DictReader(fh) if u["sent"] == "true"
                      and u.get("committed") == "true")
    lines = [f"# {name}: what the pool missed, from {len(graded)} graded units it never read", "",
             f"Drawn from the units outside {key[0]['measure_dir']} (pool: cosine >= "
             f"{man['min_cosine']}), in two strata: near (within {ELUSION_BAND} below the "
             f"threshold) and far (the rest). Each stratum's count is weighted by its size "
             f"over rows drawn; intervals are 95 per cent Wilson on the raw counts.", "",
             "| stratum | units | graded | graded yes | yes share (95% CI) "
             "| estimated yes units in stratum |", "|---|---|---|---|---|---|"]
    est_total = 0.0
    for s in sorted({k["stratum"] for k in key}):
        ks = [k for k in key if k["stratum"] == s and k["row_id"] in graded]
        if not ks:
            continue
        size, n = int(ks[0]["stratum_size"]), len(ks)
        c = sum(1 for k in ks if graded[k["row_id"]][0] >= yes_from)
        lo, hi = wilson(c, n)
        est = size * c / n
        est_total += est
        lines.append(f"| {s} | {size:,} | {n} | {c} | {c / n:.2f} ({lo:.2f}-{hi:.2f}) | "
                     f"{est:,.0f} (up to {size * hi:,.0f}) |")
    lines += ["", f"The pipeline found {in_pool} committed units inside the pool (its decision, "
                  f"not a person's). The strata estimate {est_total:,.0f} committed units outside it, "
                  f"so the pool's recall of the corpus is about "
                  f"{in_pool / (in_pool + est_total):.0%} if the model's count inside is right; "
                  f"a zero count in a stratum bounds it only by the interval's upper end.", ""]
    lines += fit_gate(name, data)
    out = where(data, "reports", f"evaluate_elusion_{name}.md")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return out, len(graded)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("label_set")
    ap.add_argument("--evaluate", action="store_true")
    ap.add_argument("--elusion", action="store_true",
                    help="units the funnel never read, from just below and far below the pool")
    ap.add_argument("--seed", type=int, default=1)
    ap.add_argument("--design", default="train",
                    choices=("train", "eval", "bands", "floor", "tune"),
                    help="train (the default): 30 near the floor, 10 far from it; eval: a "
                         "random draw by the decision (60 committed, 30 not), kept out of "
                         "every fit; bands: by band of the posterior (5/8/12/10/5); floor or "
                         "tune: around each measure's floor (about 20 or 28; links around "
                         "0.7 or over the whole range)")
    ap.add_argument("--near-gate", action="store_true",
                    help="check all, train or eval: also draw sentences just below each "
                         "measure's gate (the last training round)")
    ap.add_argument("--dry-run", action="store_true",
                    help="check all, train or eval: how many sentences and packages the draw "
                         "would bring, per decision; nothing is written")
    ap.add_argument("--projects", default="", help="comma-separated project ids: the run's scope")
    ap.add_argument("--data", default=data_root())
    args = ap.parse_args()
    kw = dict(projects=classify.scope(args.projects))
    if args.evaluate:
        out, n = evaluate(args.label_set, args.data, args.seed, elusion=args.elusion, **kw)
        print(f"{n} graded rows; wrote {out}")
    else:
        if args.near_gate:
            kw["near_gate"] = True
        if args.dry_run:
            kw["dry_run"] = True
        path, n = build(args.label_set, args.data, args.seed, elusion=args.elusion,
                        design=args.design, **kw)
        print(f"{'would write' if args.dry_run else 'wrote'} {path}: {n} rows to grade")
    return 0


if __name__ == "__main__":
    sys.exit(main())
