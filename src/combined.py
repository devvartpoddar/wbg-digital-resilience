"""The two sheets a person grades: one of sentences, one of packages, each
drawn once for every decision (docs/method/checks-and-postgres.md, "The
combined sheets").

Every decision still says how much it wants each item: a measure the
sentences near its floor (training) or a random draw of what it commits
and what it does not (evaluation), an asset the same over its own
posterior, the paragraph rule the sentences whose asset comes from the
paragraph, the pair checks a band of their keep probability, and the
package assets and links bands of theirs; a few random sentences check
what no gate let in. A decision's chance of taking an item is its
stratum's target over the stratum's size; the item's chance is the
largest over every decision it is in, and the item is drawn once
(Poisson sampling, a hash of the seed and the item standing for the
uniform). Every decision then weights every item it sees by one over that
chance, so each decision's figures stay unbiased however many others
read the same item, and a sentence wanted by three decisions is read
once.

A sentence's block shows its passage once, with the sentence marked, and
under it one line per question: each measure the sentence was asked
about (0-3), each type of a measure that committed it (yes, no, unsure),
each asset it was asked about or given (0-3), and a last line naming
anything else that applies, which closes the block. In a closed block a
blank line is 0 or no; a block not closed is not read. A package's block
does the same for the assets it was asked about and the commitments it
was paired with. The pair checks are not asked: two measure lines on the
same sentence answer them.

The answers are written to the label files every decision already reads
(data/labels/<set>_units.csv, assets_units.csv, pairs.csv,
package_assets.csv, package_links.csv), each row with its weight, so the
evaluation, the layers, the gates and the tuning report are unchanged.
The evaluation design's rows carry an eval: stratum and never fit.
"""
import csv
import hashlib
import json
import math
import os

from paths import where

SHEET = "sentences"
PACKAGES = "packages"
NEAR = 0.2                       # "near" a floor: within this of it (check_units.TRAIN_NEAR)
TRAIN = {"measure": {"train:near": 30, "train:far": 10},
         "asset": {"train:near": 15, "train:far": 5},
         "asset_paragraph": 60, "random": 50, "near_gate": 25}
# The evaluation: 900 committed sentences over the measures (about +/-4
# points on the pooled precision whatever it turns out to be, 2026-10-23),
# 30 not committed a measure for its recall, and the assets the same way.
EVAL = {"committed_total": 900, "not_committed": 30,
        "asset": {"eval:asset_yes": 20, "eval:asset_no": 5},
        "asset_paragraph": 40, "random": 50, "near_gate": 25}
GRADES = ["0", "1", "2", "3", "unsure"]
YES_NO = ["yes", "no", "unsure"]


# ------------------------------------------------------------ the draw

def uniform(seed, item):
    """A number in [0, 1) fixed by the seed and the item (AGENTS.md rule 5)."""
    h = hashlib.sha256(f"{seed}\t{item}".encode("utf-8")).hexdigest()
    return int(h[:15], 16) / 16 ** 15


def inclusion(members, targets):
    """{item: chance} from members ({item: [stratum]}) and targets
    ({stratum: rows wanted}): each stratum's target over its size, the
    largest over an item's strata, at most 1."""
    size = {}
    for strata in members.values():
        for s in set(strata):
            size[s] = size.get(s, 0) + 1
    out = {}
    for item, strata in members.items():
        p = max((min(1.0, targets.get(s, 0) / size[s]) for s in set(strata)), default=0.0)
        if p > 0:
            out[item] = p
    return out


def poisson(chance, seed):
    """The items drawn, in the order the sheet shows them (by their
    uniform, so the order is random and repeatable)."""
    got = [(uniform(seed, i), i) for i, p in chance.items() if uniform(seed, i) < p]
    return [i for _, i in sorted(got)]


def band(x, bands):
    return next((b[0] for b in bands if b[1] <= x < b[2]), None)


def eval_targets(n_measures):
    """(committed, not committed): the evaluation's rows wanted a measure."""
    per = math.ceil(EVAL["committed_total"] / max(1, n_measures))
    return per, EVAL["not_committed"]


def sentence_members(frames, design, near_gate=False):
    """({unit_id: [stratum]}, {stratum: target}): every decision's stratum
    for every sentence it reads."""
    members, targets = {}, {}
    add = lambda u, s: members.setdefault(u, []).append(s)
    ev = design == "eval"
    per_yes, per_no = eval_targets(len(frames["measures"]))
    for name, m in frames["measures"].items():
        for uid, u in m["units"].items():
            p = m["posterior"](u)
            if ev:
                s = "eval:committed" if p >= m["floor"] else "eval:not_committed"
            else:
                s = "train:near" if abs(p - m["floor"]) <= NEAR else "train:far"
            add(uid, f"measure:{name}:{s}")
        if ev:
            targets[f"measure:{name}:eval:committed"] = per_yes
            targets[f"measure:{name}:eval:not_committed"] = per_no
        else:
            for s, n in TRAIN["measure"].items():
                targets[f"measure:{name}:{s}"] = n
        for uid in m.get("near_gate", ()) if near_gate else ():
            add(uid, f"near_gate:{name}")
            targets[f"near_gate:{name}"] = (EVAL if ev else TRAIN)["near_gate"]
    a = frames["assets"]
    for uid, own in a["amap"].items():
        win = {x for x, _ in a["winners"].get(uid, ([], ""))[0]}
        unit = a["winners"].get(uid, ([], ""))[1] == "unit"
        for asset, p in own.items():
            if ev:
                s = "eval:asset_yes" if unit and asset in win else "eval:asset_no"
            else:
                s = "train:near" if abs(p - a["floor"]) <= NEAR else "train:far"
            add(uid, f"asset:{asset}:{s}")
    for asset in a["ids"]:
        for s, n in (EVAL if ev else TRAIN)["asset"].items():
            targets[f"asset:{asset}:{s}"] = n
    for uid, (_, level) in a["winners"].items():
        if level == "paragraph":
            add(uid, "asset_paragraph")
    targets["asset_paragraph"] = (EVAL if ev else TRAIN)["asset_paragraph"]
    if not ev:
        import judged
        for r in frames["pairs"]:
            s = band(min(float(r.get("p_keep_a") or 0), float(r.get("p_keep_b") or 0)),
                     judged.BANDS)
            if s:
                add(r["unit_id"], f"pairs:{s}")
        for s, _, _, n in judged.BANDS:
            targets[f"pairs:{s}"] = n
    for uid in frames["corpus_units"]:
        add(uid, "random")
    targets["random"] = (EVAL if ev else TRAIN)["random"]
    return members, targets


def package_members(frames, design):
    """({package_id: [stratum]}, {stratum: target}) for the packages."""
    import check_units
    import judged
    pre = "eval:" if design == "eval" else ""
    members, targets = {}, {}
    p = frames["packages"]
    for r in p["assets"]:
        s = band(r["p"], judged.BANDS)
        if s:
            members.setdefault(r["package_id"], []).append(f"package_assets:{pre}{s}")
    for s, _, _, n in judged.BANDS:
        targets[f"package_assets:{pre}{s}"] = n
    for r in p["links"]:
        s = band(r["p"], check_units.LINK_BANDS)
        if s:
            members.setdefault(r["package_id"], []).append(f"package_links:{pre}{s}")
    for s, _, _, n in check_units.LINK_BANDS:
        targets[f"package_links:{pre}{s}"] = n
    return members, targets


# ------------------------------------------------------------ what the runs say

def load_frames(data, design, near_gate=False, projects=None):
    """Everything the draw reads, from the current runs: each measure's
    units the model read, its types' assignment, the asset map and its
    winners, the pair checks and the package runs, plus the corpus."""
    import check_units
    import classify
    import commit_checks
    import kinds
    import labels
    import link_packages
    import measure
    import panel
    kw = {"projects": projects}
    frames = {"measures": {}}
    corpus = None
    for name in labels.label_sets():
        if name == measure.ASSET_SET:
            continue
        try:
            folder = measure.current_run(name, data, **kw)
        except SystemExit:
            folder = None
        if not folder:
            print(f"check: {name}: no current measure run; left out")
            continue
        run = measure.Run(name, data, build=False, log=lambda *_: None, **kw)
        if corpus is None:
            run.load_corpus()
            corpus = run
        with open(os.path.join(folder, "units.csv"), newline="", encoding="utf-8") as fh:
            units = {u["unit_id"]: u for u in csv.DictReader(fh) if u["sent"] == "true"}
        assigned = {}
        af = kinds.current_assign(name, data, **kw)
        if af:
            with open(os.path.join(af, "assignments.csv"), newline="", encoding="utf-8") as fh:
                assigned = {r["unit_id"]: r for r in csv.DictReader(fh)}
        lab = run.labels[0]
        m = {"units": units, "label_id": lab["id"], "name": lab["name"],
             "qsha": classify.question_shas(run.qs)[lab["id"]], "floor": run.floor,
             "posterior": check_units._posterior, "dir": os.path.basename(folder),
             "assigned": assigned, "types": {t: (n, d) for t, n, d in
                                             check_units.set_types(name, data)}}
        if near_gate:
            m["near_gate"] = near_gate_units(run, units, data)
        frames["measures"][name] = m
    if corpus is None:
        raise SystemExit("check: no measure has a current run; nothing to draw")
    amap = corpus.unit_assets
    winners = {}
    for r in corpus.index:
        uid = r["unit_id"]
        win, level = measure.sentence_assets(amap.get(uid, {}),
                                             [amap.get(x, {}) for x in corpus.before[uid]])
        if win:
            winners[uid] = (win, level)
    assets = [a for a in labels.load_set(measure.ASSET_SET) if a["active"] == "true"]
    afolder = classify.current_run(measure.ASSET_SET, data, **kw) or \
        classify.last_run(measure.ASSET_SET, data)
    qshas = {}
    if afolder:
        with open(os.path.join(afolder, "manifest.json"), encoding="utf-8") as fh:
            qshas = json.load(fh).get("question_sha256", {})
    frames["assets"] = {"amap": amap, "winners": winners, "ids": [a["id"] for a in assets],
                        "names": {a["id"]: a["name"] for a in assets},
                        "floor": float(panel.get("asset.floor", 0.3)),
                        "qsha": hashlib.sha256(json.dumps(qshas, sort_keys=True)
                                               .encode("utf-8")).hexdigest()}
    frames["pairs"], frames["pairs_run"] = [], ""
    pf = commit_checks.current_run(data, projects)
    if pf:
        with open(os.path.join(pf, "pairs.csv"), newline="", encoding="utf-8") as fh:
            frames["pairs"] = list(csv.DictReader(fh))
        frames["pairs_run"] = os.path.basename(pf)
    frames["packages"] = package_frames(data, projects)
    frames["corpus"] = corpus
    frames["corpus_units"] = [r["unit_id"] for r in corpus.index]
    return frames


def near_gate_units(run, units, data):
    """The sentences a measure never read whose best cosine to its phrases
    is within check_units.ELUSION_BAND below its gate."""
    import check_units
    import search
    folder, _ = search.run_folder(run.name, data)
    gate = run.settings["min_cosine"]
    if not gate:
        return []                    # no gate: every sentence is read, none is below it
    return sorted(h["unit_id"] for h in search.hits(folder, data, labels=[run.labels[0]["id"]])
                  if h["unit_id"] not in units
                  and gate - check_units.ELUSION_BAND <= float(h["score"]) < gate)


def package_frames(data, projects=None):
    """The current link run's package-asset and commitment-package pairs,
    each with the probability its floor reads."""
    import link_packages
    out = {"assets": [], "links": [], "packages": {}, "units": {}, "run": ""}
    folder = link_packages.current_run(data, projects=projects)
    if not folder:
        print("check: no current link run; the packages sheet is left out")
        return out
    out["run"] = os.path.basename(folder)
    out["packages"] = {p["package_id"]: p for p in link_packages.load_packages(data, projects)}
    out["units"] = {(n, u["unit_id"]): u for n, (_, us) in
                    link_packages.committed_units(data, projects).items() for u in us}
    import labels
    asked = link_packages.asset_questions([a for a in labels.load_set("assets")
                                           if a["active"] == "true"])
    with open(os.path.join(folder, "package_assets.csv"), newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if r["package_id"] in out["packages"] and r.get("probability", "") != "" \
                    and r["label_id"] in asked:
                out["assets"].append({**r, "p": float(r.get("p_calibrated") or r["probability"])})
    with open(os.path.join(folder, "commitment_packages.csv"), newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            if (r["package_id"] in out["packages"] and r.get("probability", "") != ""
                    and (r["label_set"], r["unit_id"]) in out["units"]):
                out["links"].append({**r, "p": float(r.get("p_calibrated") or r["probability"])})
    return out


# ------------------------------------------------------------ blocks and lines

def sentence_blocks(frames, drawn, chance, design, prior=None):
    """[block] for the drawn sentences: {block, unit_id, weight, lines:
    [{kind, label_set, label_id, type_id, asset_id, label, text, answer}]}.
    prior (prior_grades) fills in what a sheet of the same kind already
    graded on the sentence under the same questions: the measure's grade
    and its types, and the sentence's assets, so keeping them loses
    nothing."""
    prior = prior or {"measures": {}, "assets": {}}
    a = frames["assets"]
    texts = frames["corpus"].texts
    out = []
    for i, uid in enumerate(drawn, 1):
        b = {"block": f"S{i:04d}", "unit_id": uid, "weight": 1.0 / chance[uid],
             "chance": chance[uid], "lines": []}
        for name, m in frames["measures"].items():
            u = m["units"].get(uid)
            if not u:
                continue
            before = prior["measures"].get((name, text_sha(u["text"]), m["qsha"]))
            b["lines"].append({"kind": "measure", "label_set": name, "label_id": m["label_id"],
                               "label": m["name"], "text": "",
                               "answer": before["grade"] if before else ""})
            row = m["assigned"].get(uid)
            if row:
                said = set((before or {}).get("type", "").split("|"))
                unsure = set((before or {}).get("unsure", "").split("|"))
                for t, (tname, tdef) in m["types"].items():
                    if f"p_{t}" in row:
                        b["lines"].append({"kind": "type", "label_set": name, "type_id": t,
                                           "label": f"{m['name']}: {tname}", "text": tdef or "",
                                           "answer": "yes" if t in said else
                                           "unsure" if t in unsure else ""})
        win = {x for x, _ in a["winners"].get(uid, ([], ""))[0]}
        gave = prior["assets"].get(text_sha(texts.get(uid, "")))
        for asset in a["ids"]:
            if asset in a["amap"].get(uid, {}) or asset in win:
                b["lines"].append({"kind": "asset", "asset_id": asset,
                                   "label": a["names"][asset], "text": "",
                                   "answer": "" if gave is None else
                                   "2" if asset in gave[0] else
                                   "unsure" if asset in gave[1] else ""})
        b["lines"].append({"kind": "other", "label": "anything else that applies",
                           "text": "Required, closes the block: none, or the ids of the "
                                   "measures and assets that apply and are not listed, joined "
                                   "by | (tab 'lists')."})
        b["text"] = texts.get(uid, "")
        out.append(b)
    return out


def package_blocks(frames, drawn, chance):
    p = frames["packages"]
    names = frames["assets"]["names"]
    mnames = {n: m["name"] for n, m in frames["measures"].items()}
    by_pkg_a, by_pkg_l = {}, {}
    for r in p["assets"]:
        by_pkg_a.setdefault(r["package_id"], []).append(r)
    for r in p["links"]:
        by_pkg_l.setdefault(r["package_id"], []).append(r)
    out = []
    for i, pid in enumerate(drawn, 1):
        b = {"block": f"K{i:04d}", "package_id": pid, "weight": 1.0 / chance[pid],
             "chance": chance[pid], "lines": []}
        for r in sorted(by_pkg_a.get(pid, []), key=lambda r: r["label_id"]):
            b["lines"].append({"kind": "package_asset", "asset_id": r["label_id"], "row": r,
                               "label": names.get(r["label_id"], r["label_id"]), "text": ""})
        for r in sorted(by_pkg_l.get(pid, []), key=lambda r: (r["label_set"], r["unit_id"])):
            u = p["units"][(r["label_set"], r["unit_id"])]
            b["lines"].append({"kind": "package_link", "label_set": r["label_set"], "row": r,
                               "label": mnames.get(r["label_set"], r["label_set"]),
                               "text": u["text"]})
        b["lines"].append({"kind": "other", "label": "anything else that applies",
                           "text": "Required, closes the block: none, or the ids of assets "
                                   "the package buys or builds that are not listed, joined by |."})
        out.append(b)
    return out


def text_sha(text):
    return hashlib.sha256((text or "").encode("utf-8")).hexdigest()


def number_lines(blocks):
    for b in blocks:
        for j, line in enumerate(b["lines"], 1):
            line["line"] = f"{b['block']}.{j:02d}"


# ------------------------------------------------------------ reading answers

def closed(blocks, answers):
    """The blocks whose last line (anything else) is answered."""
    return [b for b in blocks if (answers.get(b["lines"][-1]["line"]) or "").strip()]


def named(block, answers):
    """The ids named on a block's last line."""
    v = (answers.get(block["lines"][-1]["line"]) or "").strip().lower()
    return set() if v in ("", "none") else {x.strip() for x in v.split("|") if x.strip()}


def grade_of(v):
    """A 0-3 line's answer: an int, None for unsure, 0 when blank."""
    v = (v or "").strip().lower()
    if v == "unsure":
        return None
    if v == "":
        return 0
    return int(float(v))


def yes_no(v):
    v = (v or "").strip().lower()
    return v if v in ("yes", "unsure") else "no"


def block_answers(block, answers, asset_ids, lo=2):
    """What a closed block says: ({set: grade, None when unsure}, assets the
    person gave (graded lo or more, or named on the last line), assets left
    unsure, {set: (types yes, types unsure)}, measures named on the last
    line)."""
    measures, assets, unsure, types = {}, set(), set(), {}
    other = named(block, answers)
    for line in block["lines"]:
        v = answers.get(line["line"], "")
        if line["kind"] == "measure":
            measures[line["label_set"]] = grade_of(v)
        elif line["kind"] == "asset":
            g = grade_of(v)
            if g is None:
                unsure.add(line["asset_id"])
            elif g >= lo:
                assets.add(line["asset_id"])
        elif line["kind"] == "type":
            yes, uns = types.setdefault(line["label_set"], (set(), set()))
            a = yes_no(v)
            if a == "yes":
                yes.add(line["type_id"])
            elif a == "unsure":
                uns.add(line["type_id"])
    asset_ids = set(asset_ids)
    return measures, assets | (other & asset_ids), unsure, types, other - asset_ids


def unreadable(blocks, answers):
    """The lines whose answer is not on their list (a pasted or typed value
    the dropdown would refuse), in sheet order."""
    ok = {"measure": GRADES, "asset": GRADES}
    out = []
    for b in blocks:
        for ln in b["lines"]:
            v = (answers.get(ln["line"]) or "").strip().lower()
            if not v or ln["kind"] == "other":
                continue
            if v not in ok.get(ln["kind"], YES_NO):
                out.append(ln["line"])
    return out


def pair_answer(ga, gb, lo=2):
    """The pair check two measure grades answer."""
    if ga is None or gb is None:
        return "unsure"
    return {(True, True): "both", (True, False): "only A", (False, True): "only B",
            (False, False): "neither"}[(ga >= lo, gb >= lo)]


# ------------------------------------------------------------ the workbook

HOW_TO = [
    "Two tabs to grade: 'sentences' and 'packages'. Each block starts with the passage (the "
    "sentence marked »like this«) and its project, section and component; under it, one line "
    "per question.",
    "Sentences: grade each measure line 0-3 on the scale (tab 'scale'): does the sentence commit "
    "the project to the measure? Read the passage to understand the sentence, but grade the "
    "sentence. For a measure you grade 2 or 3, answer its type lines yes, no or unsure (tab "
    "'types'). Grade each asset line 0-3: is the sentence about this digital asset (reading the "
    "passage, so 'it' counts as the asset it refers to)?",
    "Only mark what applies: in a block you close, a blank line counts as 0 or no. Close a block "
    "by answering its last line, 'anything else that applies': none, or the ids (tab 'lists') "
    "of measures and assets that apply and are not listed. A block not closed is not read.",
    "Packages: answer each asset line yes, no or unsure (does the package buy, build, install, "
    "connect, operate or design the asset?) and each commitment line (would carrying out the "
    "package deliver or apply the commitment?); close the block the same way.",
    "unsure on any line leaves that question out. A grade already given to the same sentence "
    "under the same question is filled in; change it if you disagree.",
    "The model's answers are not shown. Do not reorder, insert or delete rows: the line id "
    "joins each answer to its key.",
]
COLS = [("line", 10), ("what", 10), ("label", 32), ("text", 110), ("answer", 10), ("note", 24)]


def mark(paragraph, start, end):
    """The paragraph with the sentence between start and end marked."""
    if 0 <= start < end <= len(paragraph):
        return paragraph[:start] + "»" + paragraph[start:end] + "«" + paragraph[end:]
    return paragraph


def write_workbook(path, sblocks, pblocks, scale, measures, assets):
    """The workbook: how to, sentences, packages, and the reference tabs
    (scale, types, lists). measures: [(id, name, definition, [(type, name,
    definition)])]; assets: [(id, name, definition)]."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    import check_units
    wb = Workbook()
    wb.active.title = "how to"
    for line in HOW_TO:
        wb.active.append([line])
    wb.active.column_dimensions["A"].width = 120
    wrap = Alignment(wrap_text=True, vertical="top")
    shade = PatternFill("solid", fgColor="EEF2F7")
    for title, blocks in ((SHEET, sblocks), (PACKAGES, pblocks)):
        ws = wb.create_sheet(title)
        ws.append([c for c, _ in COLS])
        for i, (_, w) in enumerate(COLS, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        rows_of = {"grade": [], "yes": []}
        for k, b in enumerate(blocks):
            fill = shade if k % 2 else None
            head = ws.max_row + 1
            ws.append([b["block"], b["what"], b["where"], check_units.cell_text(b["passage"]),
                       "", ""])
            ws.row_dimensions[head].height = min(400, 15 * (len(b["passage"]) // 95 + 1))
            if b.get("sentence"):
                ws.append(["", "sentence", "", check_units.cell_text(b["sentence"]), "", ""])
                ws.cell(ws.max_row, 4).font = Font(bold=True)
            for line in b["lines"]:
                ws.append([line["line"], line["kind"].replace("package_", ""), line["label"],
                           check_units.cell_text(line.get("text", "")), line.get("answer", ""),
                           ""])
                if line["kind"] in ("measure", "asset"):
                    rows_of["grade"].append(ws.max_row)
                elif line["kind"] != "other":
                    rows_of["yes"].append(ws.max_row)
            for row in ws.iter_rows(min_row=head, max_row=ws.max_row):
                for cell in row:
                    cell.alignment = wrap
                    if fill:
                        cell.fill = fill
        for kind, values in (("grade", GRADES), ("yes", YES_NO)):
            if rows_of[kind]:
                dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"',
                                    allow_blank=True)
                ws.add_data_validation(dv)
                for a, b in runs(rows_of[kind]):
                    dv.add(f"E{a}:E{b}")
    s = wb.create_sheet("scale")
    s.append(["grade", "meaning"])
    for i, line in enumerate(scale):
        s.append([i, line])
    s.column_dimensions["B"].width = 120
    t = wb.create_sheet("types")
    t.append(["measure", "type", "name", "definition"])
    for mid, mname, _, ts in measures:
        for tid, tname, tdef in ts:
            t.append([mname, tid, tname, check_units.cell_text(tdef)])
    for c, w in zip("ABCD", (26, 30, 30, 100)):
        t.column_dimensions[c].width = w
    lists = wb.create_sheet("lists")
    lists.append(["id", "kind", "name", "definition"])
    for mid, mname, mdef, _ in measures:
        lists.append([mid, "measure", mname, check_units.cell_text(mdef)])
    for aid, aname, adef in assets:
        lists.append([aid, "asset", aname, check_units.cell_text(adef)])
    for c, w in zip("ABCD", (28, 10, 30, 100)):
        lists.column_dimensions[c].width = w
    wb.save(path)


def runs(rows):
    """[(first, last)] of the consecutive stretches in a sorted list of rows."""
    out = []
    for r in rows:
        if out and r == out[-1][1] + 1:
            out[-1] = (out[-1][0], r)
        else:
            out.append((r, r))
    return out


def read_workbook(path):
    """{line id: answer} of both graded tabs (blank answers left out)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    out = {}
    for title in (SHEET, PACKAGES):
        if title not in wb.sheetnames:
            continue
        rows = wb[title].iter_rows(values_only=True)
        head = list(next(rows))
        li, ai = head.index("line"), head.index("answer")
        for r in rows:
            if r[li] and r[ai] is not None and str(r[ai]).strip() != "":
                v = r[ai]
                out[str(r[li])] = str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)
    return out


# ------------------------------------------------------------ drawing the sheets

SENTENCE_KEY = ["row_id", "block", "kind", "label_set", "label_id", "type_id", "asset_id",
                "unit_id", "text_sha256", "weight", "chance"]
PACKAGE_KEY = ["row_id", "block", "kind", "label_set", "asset_id", "unit_id", "package_id",
               "weight", "chance"]


def key_path(data, what, seed):
    return where(data, "review", f"check_{what}_s{seed}.key.csv")


def measure_stratum(m, u, design):
    p = m["posterior"](u)
    if design == "eval":
        return "eval:committed" if p >= m["floor"] else "eval:not_committed"
    return "train:near" if abs(p - m["floor"]) <= NEAR else "train:far"


def same_kind(r, design):
    """Whether a stored label comes from a sheet of the design's kind: a
    training sheet never shows an evaluation's grades (they would then be
    stored as training rows and fit), nor an evaluation sheet a training's
    (it is graded afresh, so its figures are out of sample)."""
    return (r.get("stratum") or "").startswith("eval:") == (design == "eval")


def prior_grades(data, frames, design):
    """{"measures": {(set, text SHA-256, question SHA-256): label row},
    "assets": {text SHA-256: (assets given, assets unsure)}} of what sheets
    of the same kind already graded under the current questions."""
    out = {"measures": {}, "assets": {}}
    for name, m in frames["measures"].items():
        path = where(data, "labels", f"{name}_units.csv")
        if not os.path.exists(path):
            continue
        with open(path, newline="", encoding="utf-8") as fh:
            for r in sorted(csv.DictReader(fh), key=lambda r: r.get("graded_on", "")):
                # an 'outside' row is a measure nobody looked at on that
                # sentence, not a grade to show again
                if r.get("grade", "") != "" and r.get("question_sha256") == m["qsha"] \
                        and same_kind(r, design) \
                        and not (r.get("stratum") or "").endswith(":outside"):
                    out["measures"][(name, r["text_sha256"], m["qsha"])] = r
    path = where(data, "labels", "assets_units.csv")
    if os.path.exists(path):
        with open(path, newline="", encoding="utf-8") as fh:
            for r in sorted(csv.DictReader(fh), key=lambda r: r.get("graded_on", "")):
                if r.get("question_sha256") == frames["assets"]["qsha"] and same_kind(r, design):
                    out["assets"][r["text_sha256"]] = (
                        set(r.get("asset", "").split("|")) - {"", "none"},
                        set(r.get("unsure", "").split("|")) - {""})
    return out


def context(frames, data, blocks):
    """Each sentence block's passage (its paragraph, the sentence marked) and
    where it sits (project, section, component)."""
    import measure
    import units as units_mod
    corpus = frames["corpus"]
    every = {u["unit_id"]: u for u in units_mod.search_units(data)}
    index = {r["unit_id"]: r for r in corpus.index}
    cache = {}
    for b in blocks:
        u = every[b["unit_id"]]
        para = corpus.paras[u["paragraph_id"]]
        doc = measure._text(cache, data, u["doc_id"])
        ps = int(para["char_start"])
        paragraph = doc[ps:int(para["char_end"])]
        b["passage"] = mark(paragraph, int(u["char_start"]) - ps, int(u["char_end"]) - ps)
        state = measure.state_for(index[b["unit_id"]], para, b["text"], corpus.names)
        b["what"] = "passage"
        b["where"] = " · ".join(x for x in (index[b["unit_id"]].get("project_ids", ""),
                                            state.get("section", ""),
                                            state.get("component", "")) if x)
        b["sentence"] = b["text"]
        b["paragraph_id"] = u["paragraph_id"]


def build(data, seed, design="train", near_gate=False, projects=None, dry_run=False):
    """Draw both sheets into data/review/check_units_all_s<seed>.xlsx and
    write every decision's key file. Returns (workbook, sentences, packages).
    dry_run: say how many sentences and packages each decision would bring,
    and write nothing."""
    import check_units
    import classify
    import commit_checks
    import judged
    import labels
    import link_packages
    xlsx = check_units.all_path(data, seed)
    if not dry_run and (os.path.exists(xlsx) or os.path.exists(key_path(data, "sentences", seed))):
        raise SystemExit(f"check: {xlsx} or its keys exist and may hold grades; use another --seed")
    frames = load_frames(data, design, near_gate, projects)
    writes = [key_path(data, w, seed) for w in ("sentences", "packages", "assets", "pairs",
                                                 "package_assets", "links")]
    for name in frames["measures"]:
        base = where(data, "review", f"check_units_{name}_s{seed}")
        writes += [base + ".key.csv", base + ".xlsx"]
    taken = [p for p in writes if os.path.exists(p)]
    if taken and not dry_run:
        raise SystemExit(f"check: seed {seed} is taken ({os.path.basename(taken[0])} exists "
                         f"and may hold grades); use another --seed")
    members, targets = sentence_members(frames, design, near_gate)
    chance = inclusion(members, targets)
    drawn = poisson(chance, seed)
    if dry_run:
        pm, pt = package_members(frames, design)
        pchance = inclusion(pm, pt)
        for line in size_lines(members, targets, chance, drawn, pchance, seed):
            print(line)
        return xlsx, len(drawn), len(poisson(pchance, seed))
    sblocks = sentence_blocks(frames, drawn, chance, design, prior_grades(data, frames, design))
    number_lines(sblocks)
    context(frames, data, sblocks)
    pm, pt = package_members(frames, design)
    pchance = inclusion(pm, pt)
    pblocks = package_blocks(frames, poisson(pchance, seed), pchance)
    number_lines(pblocks)
    pk = frames["packages"]
    for b in pblocks:
        p = pk["packages"][b["package_id"]]
        b.update(what="package", where=" · ".join(x for x in (p["project_id"], p["component"]) if x),
                 passage=p["text"], sentence="")
    pre = "eval:" if design == "eval" else "train:"
    fmt = lambda x: f"{x:.6g}"
    os.makedirs(where(data, "review"), exist_ok=True)
    # the sentences and packages as drawn: what a closed block is read against
    srows = [{"row_id": ln["line"], "block": b["block"], "kind": ln["kind"],
              "label_set": ln.get("label_set", ""), "label_id": ln.get("label_id", ""),
              "type_id": ln.get("type_id", ""), "asset_id": ln.get("asset_id", ""),
              "unit_id": b["unit_id"], "text_sha256": text_sha(b["text"]),
              "weight": fmt(b["weight"]), "chance": fmt(b["chance"])}
             for b in sblocks for ln in b["lines"]]
    check_units._write_key(key_path(data, "sentences", seed), SENTENCE_KEY, srows)
    prows = [{"row_id": ln["line"], "block": b["block"], "kind": ln["kind"],
              "label_set": ln.get("label_set", ""), "asset_id": ln.get("asset_id", ""),
              "unit_id": ln.get("row", {}).get("unit_id", ""), "package_id": b["package_id"],
              "weight": fmt(b["weight"]), "chance": fmt(b["chance"])}
             for b in pblocks for ln in b["lines"]]
    check_units._write_key(key_path(data, "packages", seed), PACKAGE_KEY, prows)
    # each measure's key, as a sheet of its own would have it: its lines, and
    # one 'outside' row a block the measure never read (named on the last
    # line or not), which is what the pool missed
    cols = check_units.KEY_COLS + ["weight", "frame"]
    for name, m in frames["measures"].items():
        rows = []
        for b in sblocks:
            line = next((ln for ln in b["lines"] if ln["kind"] == "measure"
                         and ln["label_set"] == name), None)
            u = m["units"].get(b["unit_id"])
            common = {"label_id": m["label_id"], "unit_id": b["unit_id"],
                      "paragraph_id": b["paragraph_id"], "question_sha256": m["qsha"],
                      "measure_dir": m["dir"], "weight": fmt(b["weight"])}
            if line:
                rows.append({**common, "row_id": line["line"], "text_sha256": text_sha(u["text"]),
                             "stratum": measure_stratum(m, u, design), "frame": "pool"})
            else:
                rows.append({**common, "row_id": f"{b['block']}.o", "frame": "outside",
                             "text_sha256": text_sha(b["text"]), "stratum": pre + "outside"})
        check_units._write_key(where(data, "review", f"check_units_{name}_s{seed}.key.csv"),
                               cols, rows)
    a = frames["assets"]
    check_units._write_key(key_path(data, "assets", seed), cols, [
        {"row_id": b["block"], "label_id": "assets", "unit_id": b["unit_id"],
         "paragraph_id": b["paragraph_id"], "text_sha256": text_sha(b["text"]),
         "question_sha256": a["qsha"], "stratum": pre + "assets", "measure_dir": "",
         "weight": fmt(b["weight"]), "frame": "pool"} for b in sblocks])
    # the pair checks the sentences answer
    block_of = {b["unit_id"]: b for b in sblocks}
    texts = {(n, uid): u["text"] for n, m in frames["measures"].items() for uid, u in m["units"].items()}
    prs = []
    for r in frames["pairs"]:
        b = block_of.get(r["unit_id"])
        if not b:
            continue
        prs.append({"row_id": f"P{len(prs) + 1:04d}", "unit_id": r["unit_id"], "set_a": r["set_a"],
                    "set_b": r["set_b"],
                    "text_sha256": text_sha(texts.get((r["set_a"], r["unit_id"]), b["text"])),
                    "question_sha256": judged.sha(commit_checks.pair_question(r["set_a"], r["set_b"])),
                    "p_keep_a": r.get("p_keep_a", ""), "p_keep_b": r.get("p_keep_b", ""),
                    "stratum": pre + "pairs", "stratum_size": "", "stratum_drawn": "",
                    "run": frames["pairs_run"], "weight": fmt(b["weight"]), "block": b["block"]})
    check_units._write_key(key_path(data, "pairs", seed),
                           judged.PAIR_KEY_COLS + ["weight", "block"], prs)
    # the packages' lines, in the shapes of the package-asset and link keys
    assets = [x for x in labels.load_set("assets") if x["active"] == "true"]
    aq = {k: judged.sha(q) for k, q in link_packages.asset_questions(assets).items()}
    pa, pl = [], []
    for b in pblocks:
        p = pk["packages"][b["package_id"]]
        for ln in b["lines"]:
            r = ln.get("row")
            if ln["kind"] == "package_asset":
                pa.append({"row_id": ln["line"], "package_id": b["package_id"],
                           "asset_id": ln["asset_id"], "package_sha256": p["sha"],
                           "question_sha256": aq.get(ln["asset_id"], ""),
                           "probability": f"{r['p']:.3f}",
                           "stratum": (pre if design == "eval" else "") +
                           (band(r["p"], judged.BANDS) or ""),
                           "stratum_size": "", "stratum_drawn": "", "run": pk["run"],
                           "weight": fmt(b["weight"])})
            elif ln["kind"] == "package_link":
                u = pk["units"][(r["label_set"], r["unit_id"])]
                pl.append({"row_id": ln["line"], "label_set": r["label_set"],
                           "unit_id": r["unit_id"], "package_id": b["package_id"],
                           "text_sha256": text_sha(u["text"]), "package_sha256": p["sha"],
                           "question_sha256": check_units.link_question_sha(),
                           "probability": f"{r['p']:.3f}",
                           "stratum": (pre if design == "eval" else "") +
                           (band(r["p"], check_units.LINK_BANDS) or ""),
                           "stratum_size": "", "stratum_drawn": "", "link_run": pk["run"],
                           "weight": fmt(b["weight"])})
    check_units._write_key(key_path(data, "package_assets", seed),
                           judged.ASSET_KEY_COLS + ["weight"], pa)
    check_units._write_key(key_path(data, "links", seed),
                           check_units.LINK_KEY_COLS + ["weight"], pl)
    run0 = next(iter(frames["measures"]))
    scale = classify.load_labels(run0, None)[1]
    measures = [(n, m["name"], labels.load_set(n)[0].get("definition", ""),
                 [(t, tn, td) for t, (tn, td) in m["types"].items()])
                for n, m in frames["measures"].items()]
    adefs = [(x["id"], x["name"], x.get("definition", "")) for x in assets]
    write_workbook(xlsx + ".part.xlsx", sblocks, pblocks, scale, measures, adefs)
    os.replace(xlsx + ".part.xlsx", xlsx)
    expected = sum(chance.values())
    print(f"check: {len(sblocks)} sentences drawn (expected {expected:.0f}), "
          f"{sum(len(b['lines']) for b in sblocks)} lines; {len(pblocks)} packages "
          f"(expected {sum(pchance.values()):.0f}); "
          f"{sum(1 for b in sblocks for ln in b['lines'] if ln.get('answer'))} lines filled "
          f"from earlier grades")
    return xlsx, len(sblocks), len(pblocks)


def size_lines(members, targets, chance, drawn, pchance, seed):
    """What a draw would bring: per decision, the rows it wants and the
    sentences it would see; then the sentences and packages in all."""
    sees = {}
    picked = set(drawn)
    for uid, strata in members.items():
        for s in set(strata):
            d = s.rsplit(":", 2)[0] if s.startswith(("measure:", "asset:")) else s.split(":")[0]
            if uid in picked:
                sees[d] = sees.get(d, 0) + 1
    want = {}
    for s, n in targets.items():
        d = s.rsplit(":", 2)[0] if s.startswith(("measure:", "asset:")) else s.split(":")[0]
        want[d] = want.get(d, 0) + n
    out = ["| decision | rows wanted | sentences drawn that it reads |", "|---|---|---|"]
    for d in sorted(want):
        out.append(f"| {d} | {want[d]} | {sees.get(d, 0)} |")
    out += ["", f"sentences: {len(drawn)} drawn at seed {seed} (expected {sum(chance.values()):.0f}, "
            f"against {sum(want.values())} rows wanted: the rest is overlap); packages: "
            f"{len(poisson(pchance, seed))} (expected {sum(pchance.values()):.0f})"]
    return out


# ------------------------------------------------------------ reading them back

def is_combined(data, seed):
    return os.path.exists(key_path(data, "sentences", seed))


def _read_key(path):
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def blocks_from_key(rows):
    """[block] rebuilt from a sentences or packages key, lines in order."""
    out, by = [], {}
    for r in rows:
        b = by.get(r["block"])
        if b is None:
            b = by[r["block"]] = {"block": r["block"], "unit_id": r.get("unit_id", ""),
                                  "package_id": r.get("package_id", ""),
                                  "weight": float(r["weight"]), "lines": []}
            out.append(b)
        b["lines"].append(dict(r, line=r["row_id"]))
    return out


class Graded:
    """The person's answers on both sheets, read once."""

    def __init__(self, data, seed):
        import check_units
        import labels
        import measure
        import panel
        self.data, self.seed = data, seed
        self.lo = min(panel.grades())
        self.answers = read_workbook(check_units.all_path(data, seed))
        sblocks = blocks_from_key(_read_key(key_path(data, "sentences", seed)))
        pblocks = blocks_from_key(_read_key(key_path(data, "packages", seed)))
        self.bad = unreadable(sblocks + pblocks, self.answers)
        for line in self.bad:
            self.answers[line] = "unsure"    # left out, and said so
        if self.bad:
            print(f"check: {len(self.bad)} answers are not on their line's list and are left "
                  f"out: {', '.join(self.bad[:20])}{' ...' if len(self.bad) > 20 else ''}")
        self.sblocks = closed(sblocks, self.answers)
        self.pblocks = closed(pblocks, self.answers)
        self.asset_ids = [a["id"] for a in labels.load_set(measure.ASSET_SET)]
        self.said = {b["block"]: block_answers(b, self.answers, self.asset_ids, self.lo)
                     for b in self.sblocks}

    def measure_grades(self, name):
        """{row_id: (grade, assets, types, unsure types)} of a measure's key:
        its lines in closed blocks (unsure left out), and its 'outside' row
        in every closed block that does not ask it (2 if the person named
        the measure on the last line, else 0)."""
        out = {}
        for b in self.sblocks:
            measures, assets, _, types, others = self.said[b["block"]]
            a = "|".join(sorted(assets)) or "none"
            line = next((ln for ln in b["lines"] if ln["kind"] == "measure"
                         and ln["label_set"] == name), None)
            if line is None:
                out[f"{b['block']}.o"] = (2 if name in others else 0, a, "", "")
                continue
            g = measures[name]
            if g is None:
                continue
            yes, uns = types.get(name, (set(), set()))
            has_types = any(ln["kind"] == "type" and ln["label_set"] == name for ln in b["lines"])
            t = ("|".join(sorted(yes)) or "none") if has_types and g >= self.lo else ""
            out[line["line"]] = (g, a, t, "|".join(sorted(uns)) if t else "")
        return out

    def asset_grades(self):
        """{block: (grade, assets, '', unsure assets)} for assets_units.csv:
        the assets the person gave the sentence ('none' for none)."""
        out = {}
        for b in self.sblocks:
            _, assets, unsure, _, _ = self.said[b["block"]]
            grades = [grade_of(self.answers.get(ln["line"], "")) for ln in b["lines"]
                      if ln["kind"] == "asset"]
            g = max([x for x in grades if x is not None] + [self.lo if assets else 0])
            out[b["block"]] = (g, "|".join(sorted(assets)) or "none", "", "|".join(sorted(unsure)))
        return out

    def pair_answers(self):
        """{row_id: only A | only B | both | neither | unsure} of the pairs key."""
        out = {}
        for r in _read_key(key_path(self.data, "pairs", self.seed)):
            said = self.said.get(r["block"])
            if said is None:
                continue
            m = said[0]
            if r["set_a"] in m and r["set_b"] in m:
                out[r["row_id"]] = pair_answer(m[r["set_a"]], m[r["set_b"]], self.lo)
        return out

    def package_answers(self, kind):
        """{line: yes | no | unsure} of the package lines of a kind in closed blocks."""
        return {ln["line"]: yes_no(self.answers.get(ln["line"], ""))
                for b in self.pblocks for ln in b["lines"] if ln["kind"] == kind}


def store_assets(data, seed, graded=None):
    """The assets the person gave each sentence, as labels of the set
    'assets' (data/labels/assets_units.csv), which the asset layers, the
    asset gates and the tuning report read."""
    import check_units
    graded = graded or Graded(data, seed)
    key = _read_key(key_path(data, "assets", seed))
    got = graded.asset_grades()
    if got:
        check_units.store_labels("assets", data, key, got,
                                 os.path.basename(key_path(data, "assets", seed))[:-8])
    return len(got)


# ------------------------------------------------------------ the headline

def weighted(items):
    """(precision, low, high, effective n) of [(correct, weight)]."""
    import check_units
    import posterior
    if not items:
        return None, None, None, 0
    w = sum(x for _, x in items)
    p = sum(x for c, x in items if c) / w
    n = posterior.effective_n([x for _, x in items])
    lo, hi = check_units.wilson(round(p * n), round(n))
    return p, lo, hi, n


def headline_items(graded, units_of, levels, lo=2):
    """The decisions the dashboards show, each scored against the person:
    {'measure': [(correct, weight)], 'joint': [...], 'recall': [...],
    'by_measure': {set: (measure items, joint items)}}. A shown commitment
    is a sentence the pipeline commits for a measure, with each asset its
    asset link gives at an allowed level; it is right in measure and asset
    together when the person graded the measure lo or more and gave the
    sentence that asset. units_of: {set: {unit_id: unit row}}."""
    import measure
    import panel
    out = {"measure": [], "joint": [], "recall": [], "by_measure": {}}
    for b in graded.sblocks:
        measures, assets, unsure, _, _ = graded.said[b["block"]]
        w = b["weight"]
        for name, g in measures.items():
            u = units_of.get(name, {}).get(b["unit_id"])
            if u is None or g is None:
                continue
            mine = out["by_measure"].setdefault(name, ([], []))
            shown = measure.is_committed(u, panel.floor(name))
            links = set((u.get("asset_links") or u.get("asset_link") or "").split("|")) - {""}
            if u.get("asset_link_level", "unit") not in levels:
                links = set()
            if shown:
                out["measure"].append((g >= lo, w))
                mine[0].append((g >= lo, w))
                for a in sorted(links - unsure):
                    item = (g >= lo and a in assets, w)
                    out["joint"].append(item)
                    mine[1].append(item)
            if g >= lo:
                for a in sorted(assets):
                    out["recall"].append((shown and a in links, w))
    return out


def report(data, seed, graded=None):
    """data/reports/evaluate_combined.md: the precision of what the
    dashboards show, measure and asset together, with its interval; the
    measure alone; and by measure. Returns its path."""
    import panel
    graded = graded or Graded(data, seed)
    keys = {}
    units_of = {}
    import glob
    for path in glob.glob(where(data, "review", f"check_units_*_s{seed}.key.csv")):
        name = os.path.basename(path)[len("check_units_"):-len(f"_s{seed}.key.csv")]
        if name == "all" or name.endswith("_elusion"):
            continue
        rows = _read_key(path)
        if not rows:
            continue
        keys[name] = rows
        want = {r["unit_id"] for r in rows if r.get("frame") != "outside"}
        mdir = where(data, "measures", rows[0]["measure_dir"])
        with open(os.path.join(mdir, "units.csv"), newline="", encoding="utf-8") as fh:
            units_of[name] = {u["unit_id"]: u for u in csv.DictReader(fh) if u["unit_id"] in want}
    lv = panel.get("dashboard.commitment_link_levels", "unit|paragraph")
    levels = set(lv if isinstance(lv, list) else str(lv).split("|"))
    got = headline_items(graded, units_of, levels, graded.lo)
    is_eval = all(r["stratum"].startswith("eval:") for rows in keys.values() for r in rows)
    f = lambda x: "-" if x is None else f"{100 * x:.0f}%"
    p, l, h, n = weighted(got["joint"])
    pm, lm, hm, nm = weighted(got["measure"])
    pr, lr, hr, nr = weighted(got["recall"])
    lines = [f"# What the dashboards show, against {len(graded.sblocks)} graded sentences", "",
             ("An evaluation sheet: its figures are the ones to quote."
              if is_eval else
              "A training sheet: drawn where grades teach most, so these figures describe the "
              "hard cases, not the dashboards; quote an evaluation sheet's."), "",
             f"**Measure and asset together: {f(p)} precision** (95 per cent interval {f(l)} to "
             f"{f(h)}, {len(got['joint'])} commitments shown, effective sample {n:.0f}).", "",
             f"- Measure alone: {f(pm)} ({f(lm)} to {f(hm)}, {len(got['measure'])} sentences).",
             f"- Recall, measure and asset together, among what the search reached: {f(pr)} "
             f"({f(lr)} to {f(hr)}).",
             f"- A commitment counts when the pipeline commits the sentence (posterior at the "
             f"measure's floor) and its asset link is at a level the dashboards allow "
             f"({'|'.join(sorted(levels))}); it is right when the person graded the measure "
             f"{graded.lo} or more and gave the sentence that asset. Weights are one over each "
             f"sentence's chance of being drawn; the interval is Wilson on the effective sample.",
             "- What a sheet of the same kind graded before was filled in; an evaluation sheet "
             "is graded afresh, never from a training sheet's grades.", "",
             "| measure | precision, measure | sentences | precision, measure and asset | "
             "commitments |", "|---|---|---|---|---|"]
    for name, (mi, ji) in sorted(got["by_measure"].items()):
        lines.append(f"| {name} | {f(weighted(mi)[0])} | {len(mi)} | {f(weighted(ji)[0])} | "
                     f"{len(ji)} |")
    outside = {}
    for name in keys:
        for rid, (g, *_rest) in graded.measure_grades(name).items():
            if rid.endswith(".o") and g >= graded.lo:
                outside[name] = outside.get(name, 0) + 1
    lines += ["", "## Named on a sentence the measure never read", "",
              "- " + (", ".join(f"{k}: {v}" for k, v in sorted(outside.items())) or "none") +
              " (each such sentence stands for about one over its chance of being drawn; the "
              "skip rule reads them as graded yes below the gate).", ""]
    out = where(data, "reports", "evaluate_combined.md")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    with open(out, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return out
