#!/usr/bin/env python3
"""The pipeline as one module: choose projects, prepare them, read the results.

    import wbg
    sel = wbg.select(country="KE")               # or region=, fy=, project=, all
    wbg.prepare(sel)                             # fetch what is missing, rebuild what changed
    t = wbg.tables(sel)                          # the cleaned tables, cut to the selection
    wbg.summary(sel)                             # one row per project: where it stands
    wbg.upcoming(sel), wbg.open_notices(sel)     # package and notice detail
    wbg.export(sel)                              # optional: data/runs/<name>/

The same from the command line (./wbg, see `./wbg --help`):

    ./wbg list    --region "Western and Central Africa"
    ./wbg prepare --project P171528 --project P176932
    ./wbg summary --country KE
    ./wbg export  --fy 2024 --tables
    ./wbg pipeline                               # everything, as run.sh does

Projects are selected from the cohort (inputs/config/cohort.csv, included rows
only); that file is kept by people and only read here.

Nothing is redone that has not changed. Fetching is per project: a project is
fetched once, and again only with update=True. Every later stage runs over the
whole store from its caches, and only when its code or its inputs changed since
it last ran: data/reports/stages.json records what each stage last ran on. The
store is shared, so a selection prepared once is there for every later
selection that includes it.
"""
import argparse, csv, datetime, glob, hashlib, json, os, re, subprocess, sys
from collections import Counter, defaultdict

SRC = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SRC)
sys.path.insert(0, SRC)
import candidates as _candidates  # noqa: E402  (defaults for the command line)
import classify as _classify  # noqa: E402
import measure as _measure  # noqa: E402
import link_packages as _link  # noqa: E402
from paths import data_root, where  # noqa: E402

COHORT = os.path.join(ROOT, "inputs", "config", "cohort.csv")
# Statuses that mean the package is still to come.
UPCOMING = ("Pending", "Pending Implementation", "Under Review", "Planned")


# ------------------------------------------------------------------ selection

class Selection:
    """A set of cohort projects, with a name for its outputs."""

    def __init__(self, rows, name):
        self.rows = rows
        self.name = name
        self.ids = [r["project_id"] for r in rows]
        self._ids = set(self.ids)

    def __len__(self):
        return len(self.ids)

    def __contains__(self, pid):
        return pid in self._ids

    def __repr__(self):
        return f"Selection({self.name!r}, {len(self)} projects)"


def cohort(path=COHORT):
    with open(path, newline="", encoding="utf-8") as fh:
        return [r for r in csv.DictReader(fh) if r["included"].strip().lower() == "true"]


def _as_list(v):
    if v is None:
        return []
    if isinstance(v, str):
        return [x.strip() for x in v.split(",") if x.strip()]
    return [str(x).strip() for x in v if str(x).strip()]


# Columns a selection row takes from data/projects/projects.csv. Region,
# country and approval year come from the cohort, and from the projects
# interface only where the cohort is blank.
PROJECT_COLS = ("project_name", "status", "closing_date", "practice", "practice_code",
                "managing_unit", "unit_code")
FALLBACK = ("region", "country_code", "approval_fy")


def _projects(data):
    path = where(data or data_root(), "projects")
    if not os.path.exists(path):
        return None
    with open(path, newline="", encoding="utf-8") as fh:
        return {r["project_id"]: r for r in csv.DictReader(fh)}


def _has(values, *fields):
    """Does any value occur, without regard to case, in any of the fields?"""
    text = " | ".join(f.lower() for f in fields)
    return any(v in text for v in values)


def select(project=None, country=None, region=None, fy=None, practice=None, unit=None,
           cohort_set=None, path=COHORT, data=None):
    """Cohort projects matching every filter given; all of them when none is.

    cohort_set  the cohort's own set ('digital': operations whose object is a
                digital asset; 'digital_contributor': operations where a
                digital asset contributes to another sector's object)
    project   project ids, a list or comma-separated ('P171528,P176932')
    country   country codes ('KE', '3W')
    region    region names, or part of one, without regard to case ('western')
    fy        approval fiscal years ('2024')
    practice  Global Practice name or code, or part of one ('digital', 'transport')
    unit      managing unit or its code, or part of one ('IDD04', 'AFR EAST')

    Region, country and year are the cohort's, or the projects interface's
    where the cohort records none. Practice and unit come from the projects
    table (src/projects.py), so they need a prepared store. A project with no
    value for a filter is not matched by it."""
    rows = [dict(r) for r in cohort(path)]
    info = _projects(data)
    for r in rows:
        p = (info or {}).get(r["project_id"], {})
        for col in FALLBACK:
            r[col] = r.get(col) or p.get(col, "")
        for col in PROJECT_COLS:
            r[col] = p.get(col, "")
    given = {"project": project, "country": country, "region": region, "fy": fy,
             "practice": practice, "unit": unit, "set": cohort_set}
    blank = [k for k, v in given.items() if v is not None and not _as_list(v)]
    if blank:
        raise SystemExit(f"empty filter: --{', --'.join(blank)}")
    projects = [p.upper() for p in _as_list(project)]
    countries = [c.upper() for c in _as_list(country)]
    regions = [r.lower() for r in _as_list(region)]
    years = [re.sub(r"(?i)^fy\s*", "", y) for y in _as_list(fy)]
    years = [f"20{y}" if len(y) == 2 and y.isdigit() else y for y in years]
    pracs, units = [x.lower() for x in _as_list(practice)], [x.lower() for x in _as_list(unit)]
    sets = [x.lower() for x in _as_list(cohort_set)]
    known_sets = {r.get("cohort_set", "").lower() for r in rows}
    if sets and not set(sets) <= known_sets:
        raise SystemExit(f"no such cohort set: {', '.join(sorted(set(sets) - known_sets))} "
                         f"(the cohort has {', '.join(sorted(known_sets))})")
    if (pracs or units) and info is None:
        raise SystemExit("no projects table yet: run ./wbg prepare first")
    unknown = sorted(set(projects) - {r["project_id"] for r in rows})
    if unknown:
        raise SystemExit(f"not in the cohort (or not included): {', '.join(unknown)}")
    out = [r for r in rows
           if (not projects or r["project_id"] in projects)
           and (not countries or any(c in countries for c in r["country_code"].upper().split("|")))
           and (not regions or _has(regions, r["region"]))
           and (not years or r["approval_fy"] in years)
           and (not sets or r.get("cohort_set", "").lower() in sets)
           and (not pracs or _has(pracs, r["practice"], r["practice_code"]))
           and (not units or _has(units, r["managing_unit"], r["unit_code"]))]
    # Each value carries its filter's name, so '--practice digital' and
    # '--unit digital' do not share a folder; a long name keeps a hash of
    # the whole.
    parts = [*[f"project-{x}" for x in projects], *[f"country-{x}" for x in countries],
             *[f"region-{x}" for x in regions], *[f"fy{y}" for y in years],
             *[f"practice-{x}" for x in pracs], *[f"unit-{x}" for x in units],
             *[f"set-{x}" for x in sets]]
    name = re.sub(r"[^A-Za-z0-9._-]+", "-", "-".join(parts)).strip("-").lower() or "all"
    if len(name) > 80:
        name = name[:71].rstrip("-") + "-" + hashlib.sha256(name.encode()).hexdigest()[:8]
    return Selection(out, name)


# ------------------------------------------------------------------ preparing

# Each stage after fetching: the script, the code it depends on, and the files
# it reads. It runs when any of those changed since it last ran.
STAGES = [
    ("projects", "projects.py", ["projects.py", "fetch.py"],
     [("projects_json", "*.json"), ("documents",), ("cohort",)],
     [("projects",)]),
    ("clean", "clean.py",
     ["clean.py", "pdf_layout.py", "text_rules.py", "components.py", "sentences.py",
      "keys.py"],
     [("documents",), ("pdf", "*.pdf")],
     [("appraisal", "paragraphs.csv"), ("appraisal", "sentences.csv"),
      ("appraisal", "document_parts.csv")]),
    ("components", "components.py", ["components.py", "pdf_layout.py", "text_rules.py",
                                     "keys.py"],
     [("documents",), ("appraisal", "paragraphs.csv")],
     [("appraisal", "components.csv")]),
    ("audit", "audit.py", ["audit.py"], [("appraisal", "paragraphs.csv")],
     [("reports", "audit.txt")]),
    ("clean_procurement", "clean_procurement.py",
     ["clean_procurement.py", "glue.py", "text_rules.py", "plan_table.py", "links.py",
      "keys.py"],
     [("procurement", "packages_raw.csv"), ("procurement", "notices_raw.csv"),
      ("procurement", "awards_raw.csv"), ("appraisal", "components.csv"),
      ("appraisal", "paragraphs.csv")],
     [("procurement", "packages.csv"), ("procurement", "notices.csv"),
      ("procurement", "awards.csv"), ("procurement", "superseded_packages.csv")]),
    ("audit_procurement", "audit_procurement.py",
     ["audit_procurement.py", "glue.py", "clean_procurement.py"],
     [("procurement", "packages.csv"), ("procurement", "notices.csv"),
      ("procurement", "awards.csv"), ("procurement", "packages_raw.csv"),
      ("procurement", "notices_raw.csv"), ("procurement", "awards_raw.csv"),
      ("text", "*.txt")],
     [("reports", "audit_procurement.txt")]),
]


def _library_versions():
    """The PDF reader's version: an upgrade changes how text is read, so
    every stage after it must run again (AGENTS.md rule 7)."""
    try:
        from importlib.metadata import version
        return f"pdfplumber={version('pdfplumber')}"
    except Exception:
        return "pdfplumber=unknown"


ADD_WORKERS = 10   # parallel decision-model requests per set under ./wbg add


def add(args, data):
    """A measure (or several) from its label files to Postgres in one
    command, timed. Per set, in its own process: search, the funnel, type
    proposals (and --assign). Sets run side by side, up to --jobs at once:
    the matrix is memory-mapped, so they share one copy in the page cache
    and each takes one core. Then one load and one release for all of
    them, not one per stage. The time is printed per set and overall; the
    target is a measure visible in Metabase within five minutes of its
    files landing."""
    import time
    from concurrent.futures import ThreadPoolExecutor
    t_all = time.time()
    scope = _scope(args)
    dry = ["--dry-run"] if args.dry_run else []
    env = dict(os.environ, WBG_DATA=data)
    # The sentence embeddings first, once: every set reads the same matrix.
    _run("embed.py", *dry, "--context", "none", data=data)

    def one(name):
        t0 = time.time()
        log = []
        for script, extra in (("search.py", [name, *dry]),
                              ("measure.py", [name, "--workers", str(args.workers), *dry, *scope]),
                              ("kinds.py", [name, *dry, *scope]),
                              *((("kinds.py", [name, "--assign", *dry, *scope]),)
                                if args.assign else ())):
            cmd = [_python(), os.path.join(SRC, script), *extra]
            done = subprocess.run(cmd, env=env, cwd=ROOT, capture_output=True, text=True)
            log.append(f"== {script} {' '.join(extra)}\n{done.stdout.rstrip()}")
            if done.returncode:
                log.append(done.stderr.rstrip())
                return name, time.time() - t0, False, "\n".join(log)
        return name, time.time() - t0, True, "\n".join(log)

    with ThreadPoolExecutor(max_workers=max(1, args.jobs)) as pool:
        results = list(pool.map(one, args.label_set))
    failed = []
    for name, secs, ok, log in results:
        print(log, flush=True)
        print(f"-- {name}: {'done' if ok else 'FAILED'} in {secs:.0f}s", flush=True)
        if not ok:
            failed.append(name)
    if failed:
        print(f"add: {', '.join(failed)} failed; nothing loaded. Run again to resume.")
        return 1
    if args.dry_run:
        print(f"add: --dry-run, nothing loaded ({time.time() - t_all:.0f}s)")
        return 0
    t0 = time.time()
    _run("load_pg.py", *scope_sets(args), data=data)
    _run("release.py", data=data)
    print(f"add: {', '.join(args.label_set)} in Postgres; load and release {time.time() - t0:.0f}s, "
          f"everything {time.time() - t_all:.0f}s", flush=True)
    return 0


def scope_sets(args):
    """The load's scope flags: the cohort set(s) given, else --projects as
    the stages got it."""
    if getattr(args, "cohort_set", None):
        return [x for st in args.cohort_set for x in ("--set", st)]
    return _scope(args)


def _python():
    box = "/opt/yggdrasil/venvs/work/bin/python"
    return os.environ.get("WBG_PY") or (box if os.path.exists(box) else sys.executable)


def _run(script, *args, data, tail=0):
    """Run one stage. tail=N shows only the last N lines of its output (the
    audits print their whole report; it is in data/reports/ anyway)."""
    env = dict(os.environ, WBG_DATA=data)
    cmd = [_python(), os.path.join(SRC, script), *args]
    print(f"== {script} {' '.join(args)}".rstrip(), flush=True)
    if not tail:
        subprocess.run(cmd, check=True, env=env, cwd=ROOT)
        return
    done = subprocess.run(cmd, env=env, cwd=ROOT, capture_output=True, text=True)
    print("\n".join(done.stdout.rstrip().splitlines()[-tail:]), flush=True)
    if done.returncode:
        sys.stderr.write(done.stderr)
        raise subprocess.CalledProcessError(done.returncode, cmd)


def _fingerprint(data, code, inputs):
    """SHA-256 over the stage's code and the bytes of its input files. A PDF
    matched by a pattern is fingerprinted by name and size, which is enough:
    its content checksum is in documents.csv, also an input. Any other file
    matched by a pattern is read whole."""
    h = hashlib.sha256()
    h.update(_library_versions().encode())
    for name in code:
        with open(os.path.join(SRC, name), "rb") as fh:
            h.update(name.encode() + fh.read())
    for spec in inputs:
        # ("cohort",) is the hand-kept cohort file, outside the data folder.
        path = COHORT if spec == ("cohort",) else where(data, *spec)
        label = "cohort" if spec == ("cohort",) else os.path.relpath(path, data)
        if "*" in path:
            for p in sorted(glob.glob(path)):
                h.update(f"{os.path.basename(p)}:{os.path.getsize(p)}".encode())
                if not p.endswith(".pdf"):
                    with open(p, "rb") as fh:
                        h.update(fh.read())
        elif os.path.exists(path):
            # The path relative to the data folder: moving the folder is not
            # a change to any stage's input.
            h.update(label.encode())
            with open(path, "rb") as fh:
                for block in iter(lambda: fh.read(1 << 20), b""):
                    h.update(block)
        else:
            h.update(f"{label}:missing".encode())
    return h.hexdigest()


def _outputs(data, outputs):
    """SHA-256 over a stage's output files, or None if any is missing: a
    stage whose output was removed or replaced is not up to date, whatever
    its inputs say."""
    h = hashlib.sha256()
    for spec in outputs:
        path = where(data, *spec)
        if not os.path.exists(path):
            return None
        with open(path, "rb") as fh:
            for block in iter(lambda: fh.read(1 << 20), b""):
                h.update(block)
    return h.hexdigest()


def _state_path(data):
    return where(data, "reports", "stages.json")


def _load_state(data):
    """What each stage last ran on. A store that predates this file starts from
    what it already holds: a project with documents is counted as fetched, so
    the first prepare does not fetch the whole cohort again."""
    try:
        with open(_state_path(data), encoding="utf-8") as fh:
            return json.load(fh)
    except (OSError, ValueError):
        pass
    seen = {"appraisal": set(), "procurement": set(), "projects": set()}
    for r in _read(where(data, "documents")):
        seen["appraisal"].update(p for p in r["project_ids"].split("|") if p)
    for name in ("packages_raw.csv", "notices_raw.csv", "awards_raw.csv"):
        seen["procurement"].update(r["project_id"] for r in
                                   _read(where(data, "procurement", name)))
    seen["projects"].update(os.path.basename(p)[:-5] for p in
                            glob.glob(where(data, "projects_json", "P*.json")))
    stamp = "before stages.json"
    return {"fetched": {k: {p: stamp for p in sorted(v)} for k, v in seen.items()}}


def _fetch_status(data, kind, asked):
    """(ok, failed) from the fetch's status file; everything asked counts as
    ok when the fetcher wrote none."""
    try:
        with open(where(data, "fetch_status", f"{kind}.json"), encoding="utf-8") as fh:
            s = json.load(fh)
    except (OSError, ValueError):
        return list(asked), []
    ok = set(s.get("ok", []))
    return [p for p in asked if p in ok], [p for p in asked if p not in ok]


def _save_state(data, state):
    path = _state_path(data)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path + ".part", "w", encoding="utf-8") as fh:
        json.dump(state, fh, indent=1, sort_keys=True)
    os.replace(path + ".part", path)


def stale_stages(data=None):
    """The stages whose code, inputs or outputs differ from when they last
    ran: what ./wbg prepare would run now."""
    data = data or data_root()
    recorded = _load_state(data).get("stages", {})
    out = []
    for name, _script, code, inputs, outputs in STAGES:
        rec = recorded.get(name)
        if isinstance(rec, str):
            rec = {"in": rec, "out": None}
        out_now = _outputs(data, outputs)
        if not rec or rec["in"] != _fingerprint(data, code, inputs) or out_now is None \
                or rec["out"] not in (None, out_now):
            out.append(name)
    return out


def prepare(sel=None, update=False, data=None, fetch=True):
    """Bring the store up to date for a selection.

    1. The data folder is put in its layout (src/paths.py).
    2. Fetching: each selected project not fetched before - or every selected
       project with update=True - is fetched (appraisal documents, then
       procurement plans, notices and awards).
    3. Every later stage runs over the whole store, from its caches, only if
       its code or inputs changed since it last ran.

    Returns the names of the stages that ran."""
    data = data or data_root()
    sel = sel if sel is not None else select()
    _run("paths.py", data=data)
    state = _load_state(data)
    ran = []
    if fetch:
        # UTC, like every other timestamp the pipeline records.
        today = datetime.datetime.now(datetime.timezone.utc).date().isoformat()
        for kind, script in (("appraisal", "fetch.py"), ("procurement", "fetch_procurement.py"),
                             ("projects", "fetch_projects.py")):
            done = state.setdefault("fetched", {}).setdefault(kind, {})
            todo = [p for p in sel.ids if update or p not in done]
            if not todo:
                continue
            args = ["--projects", ",".join(todo)]
            if update and kind == "procurement":
                args.append("--refresh-records")      # notices and awards change
            _run(script, *args, data=data)
            # Only what the fetch says it fully read counts as fetched; a
            # project whose request failed is tried again next time.
            ok, _failed = _fetch_status(data, kind, todo)
            for p in ok:
                done[p] = today
            _save_state(data, state)
            ran.append(script)
    for name, script, code, inputs, outputs in STAGES:
        fp = _fingerprint(data, code, inputs)
        rec = state.get("stages", {}).get(name)
        if isinstance(rec, str):                 # a state file from before outputs
            rec = {"in": rec, "out": None}
        out = _outputs(data, outputs)
        if rec and rec["in"] == fp and out is not None and rec["out"] in (None, out):
            if rec["out"] is None:
                state["stages"][name] = {"in": fp, "out": out}
                _save_state(data, state)
            continue
        _run(script, data=data, tail=2 if name.startswith("audit") else 0)
        state.setdefault("stages", {})[name] = {"in": _fingerprint(data, code, inputs),
                                                "out": _outputs(data, outputs)}
        _save_state(data, state)
        ran.append(name)
    if not ran:
        print("nothing to do: every stage is up to date", flush=True)
    return ran


# ------------------------------------------------------------------ reading

def _read(path):
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _in(sel, ids):
    want = set(sel.ids)
    return any(p in want for p in (ids or "").split("|"))


def tables(sel=None, data=None, text=False):
    """The cleaned tables, cut to the selection, as lists of rows (dicts).

    projects, documents, paragraphs, sentences, components, packages,
    superseded_packages, notices, awards. With text=True each paragraph also carries its text (the
    tables themselves store offsets, not text)."""
    data = data or data_root()
    sel = sel if sel is not None else select()
    # A document belongs to the projects its parts belong to
    # (document_parts.csv): its listing, or, for a combined appraisal
    # document, the operations its data sheets name.
    owners = defaultdict(list)
    for r in _read(where(data, "appraisal", "document_parts.csv")):
        for p in r["project_ids"].split("|"):
            if p not in owners[r["doc_id"]]:
                owners[r["doc_id"]].append(p)
    docs = [dict(r, listed_project_ids=r["project_ids"],
                 project_ids="|".join(sorted(owners[r["doc_id"]])) if owners.get(r["doc_id"])
                 else r["project_ids"])
            for r in _read(where(data, "documents"))]
    out = {
        "projects": [r for r in _read(where(data, "projects")) if r["project_id"] in sel],
        "documents": [r for r in docs if _in(sel, r["project_ids"])],
        "paragraphs": [r for r in _read(where(data, "appraisal", "paragraphs.csv"))
                       if _in(sel, r["project_ids"])],
        "components": [r for r in _read(where(data, "appraisal", "components.csv"))
                       if _in(sel, r["project_ids"])],
    }
    # By paragraph, not document: a shared document's other parts belong to
    # other projects.
    paras = {r["paragraph_id"] for r in out["paragraphs"]}
    out["sentences"] = [r for r in _read(where(data, "appraisal", "sentences.csv"))
                        if r["paragraph_id"] in paras]
    for name in ("packages", "superseded_packages", "notices", "awards"):
        out[name] = [r for r in _read(where(data, "procurement", f"{name}.csv"))
                     if r["project_id"] in sel]
    if text:
        cache = {}
        for r in out["paragraphs"]:
            d = r["doc_id"]
            if d not in cache:
                with open(where(data, "text", f"{d}.txt"), encoding="utf-8") as fh:
                    cache[d] = fh.read()
            r["text"] = cache[d][int(r["char_start"]):int(r["char_end"])]
    return out


def _amount(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def summary(sel=None, data=None, today=None, t=None):
    """One row per selected project: what the store holds and where it stands.

    Documents and the latest appraisal-side disclosure; components and their
    cost; the latest procurement plan; packages by STEP status; what is still
    to come (Pending, Pending Implementation, Under Review, Planned) with its
    estimated value, how many are planned for a date already past, and the
    next planned date; notices, and those still open
    (deadline today or later); signed contracts and their value. Dates are
    ISO; amounts are US dollars as the plans and the awards interface give
    them. A project with no procurement plan has blank plan columns."""
    data = data or data_root()
    sel = sel if sel is not None else select()
    today = today or datetime.date.today().isoformat()
    t = t or tables(sel, data)
    by = defaultdict(lambda: defaultdict(list))
    for name in ("documents", "components"):
        for r in t[name]:
            for p in r["project_ids"].split("|"):
                by[p][name].append(r)
    for name in ("packages", "notices", "awards"):
        for r in t[name]:
            by[r["project_id"]][name].append(r)
    rows = []
    for c in sel.rows:
        p = c["project_id"]
        g = by[p]
        docs = g["documents"]
        kinds = Counter(d["doc_kind"] for d in docs)
        comps = _latest_components(g["components"])
        pk = g["packages"]
        statuses = Counter(r["status"] for r in pk)
        up = [r for r in pk if r["status"] in UPCOMING]
        next_dates = sorted(r["planned_date"] for r in up if r["planned_date"] >= today)
        open_n = [n for n in g["notices"] if n["deadline_date"] and n["deadline_date"] >= today]
        rows.append({
            "project_id": p, "project_name": c.get("project_name", ""),
            "country_code": c["country_code"], "region": c["region"],
            "approval_fy": c["approval_fy"], "status": c.get("status", ""),
            "closing_date": c.get("closing_date", ""), "practice": c.get("practice", ""),
            "managing_unit": c.get("managing_unit", ""),
            "documents": len(docs),
            "documents_by_kind": "; ".join(f"{k} {v}" for k, v in sorted(kinds.items())),
            "latest_document_date": max((d["disclosure_date"] for d in docs), default=""),
            "components": "; ".join(f"{n}. {name}" + (f" (US${cost}m)" if cost else "")
                                    for n, name, cost in comps),
            "components_cost_usd_m": f"{sum(_amount(x[2]) for x in comps):.2f}" if comps else "",
            "latest_plan_date": max((r["plan_disclosure_date"] for r in pk), default=""),
            "packages": len(pk),
            "packages_by_status": "; ".join(f"{k} {v}" for k, v in statuses.most_common()),
            "upcoming_packages": len(up),
            # Of those, how many the project's latest plan no longer lists:
            # their status is from an older plan (packages.in_latest_plan).
            "upcoming_not_in_latest_plan": sum(1 for r in up
                                               if r.get("in_latest_plan") == "false"),
            "upcoming_estimated_usd": f"{sum(_amount(r['estimated_amount']) for r in up):.2f}",
            # Still to come, but planned for a date already past: the plan has
            # not been updated, or the package is late. Either is worth a question.
            "upcoming_planned_date_passed": sum(1 for r in up if r["planned_date"]
                                                and r["planned_date"] < today),
            "next_planned_date": next_dates[0] if next_dates else "",
            "notices": len(g["notices"]),
            "open_notices": len(open_n),
            "latest_notice_date": max((n["publication_date"] for n in g["notices"]), default=""),
            "contracts_signed": len(g["awards"]),
            "contracts_value_usd": f"{sum(_amount(a['total_amount']) for a in g['awards']):.2f}",
            "latest_contract_date": max((a["signed_date"] for a in g["awards"]), default=""),
        })
    return rows


def _latest_components(rows):
    """(number, name, cost) of top-level components, from the most recent
    document that lists them; the data sheet's list where it has one."""
    tops = [r for r in rows if r["level"] == "component" and r["number"]]
    if not tops:
        return []
    latest = max(r["disclosure_date"] for r in tops)
    pick = [r for r in tops if r["disclosure_date"] == latest]
    pref = [r for r in pick if r["source"] in ("datasheet", "restructuring")] or pick
    seen, out = set(), []
    for r in sorted(pref, key=lambda r: [int(x) for x in r["number"].split(".") if x.isdigit()]):
        if r["number"] not in seen:
            seen.add(r["number"])
            out.append((r["number"], r["name"], r["cost_usd_m"]))
    return out


def upcoming(sel=None, data=None, t=None):
    """Packages still to come, one row each, soonest planned date first, with
    the notices linked to each (links.py): how many, and the latest one's
    publication date and deadline."""
    t = t or tables(sel, data)
    cols = ("project_id", "package_id", "borrower_ref", "description_clean", "category",
            "method", "method_name", "market_approach", "component_number", "component",
            "status", "status_source", "status_as_of", "planned_date", "revised_date",
            "estimated_amount", "currency", "amount_source", "plan_disclosure_date",
            "in_latest_plan")
    notices = defaultdict(list)
    for n in t["notices"]:
        if n.get("package_id"):
            notices[n["package_id"]].append(n)
    rows = []
    for r in t["packages"]:
        if r["status"] not in UPCOMING:
            continue
        row = {k: r.get(k, "") for k in cols}
        mine = sorted(notices[r["package_id"]], key=lambda n: n["publication_date"])
        row["notices"] = len(mine)
        row["latest_notice_date"] = mine[-1]["publication_date"] if mine else ""
        row["latest_notice_deadline"] = mine[-1]["deadline_date"] if mine else ""
        rows.append(row)
    return sorted(rows, key=lambda r: (r["planned_date"] or "9999", r["project_id"],
                                       r["package_id"]))


def open_notices(sel=None, data=None, today=None, t=None):
    """Notices whose deadline is today or later, soonest deadline first."""
    today = today or datetime.date.today().isoformat()
    t = t or tables(sel, data)
    cols = ("project_id", "notice_id", "notice_type", "publication_date", "deadline_date",
            "borrower_ref", "description_clean", "category", "method", "method_name",
            "package_id", "package_link")
    pk = {r["package_id"]: r for r in t["packages"]}
    rows = []
    for n in t["notices"]:
        if not (n["deadline_date"] and n["deadline_date"] >= today):
            continue
        row = {k: n.get(k, "") for k in cols}
        p = pk.get(n.get("package_id") or "", {})
        row["package_status"] = p.get("status", "")
        row["package_estimated_amount"] = p.get("estimated_amount", "")
        rows.append(row)
    return sorted(rows, key=lambda r: (r["deadline_date"], r["project_id"], r["notice_id"]))


# ------------------------------------------------------------------ writing

def _write_csv(path, rows):
    if not rows:
        # An empty file rather than none, so a CSV from an earlier export into
        # the same folder is not left behind looking current.
        open(path, "w").close()
        return
    with open(path + ".part", "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    os.replace(path + ".part", path)


def export(sel=None, name=None, data=None, include_tables=False, today=None):
    """Write a selection's results to data/runs/<name>/ and return the folder.

    summary.xlsx  sheets: projects (summary), upcoming packages, open notices,
                  components
    *.csv         with include_tables=True, every table cut to the selection

    Only what is asked for is written; nothing else reads these folders."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    data = data or data_root()
    sel = sel if sel is not None else select()
    today = today or datetime.date.today().isoformat()
    out = where(data, "runs", f"{name or sel.name}-{today}")
    os.makedirs(out, exist_ok=True)
    t = tables(sel, data)
    sheets = [("projects", summary(sel, data, today, t)),
              ("upcoming packages", upcoming(sel, data, t)),
              ("open notices", open_notices(sel, data, today, t)),
              ("components", t["components"])]
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        if not rows:
            ws.append(["(none)"])
            continue
        ws.append(list(rows[0]))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(k, "") for k in rows[0]])
        ws.freeze_panes = "A2"
    path = os.path.join(out, "summary.xlsx")
    wb.save(path + ".part")
    os.replace(path + ".part", path)
    if include_tables:
        for tname, rows in t.items():
            _write_csv(os.path.join(out, f"{tname}.csv"), rows)
    return out


# ------------------------------------------------------------------ command line

def _filters(ap):
    ap.add_argument("--project", action="append", help="project id (repeat or comma-separate)")
    ap.add_argument("--set", action="append", dest="cohort_set",
                    help="the cohort's set: digital or digital_contributor")
    ap.add_argument("--country", action="append", help="country code as in the cohort")
    ap.add_argument("--region", action="append", help="region name, or part of one")
    ap.add_argument("--fy", action="append", help="approval fiscal year")
    ap.add_argument("--practice", action="append",
                    help="Global Practice name or code, or part of one ('digital')")
    ap.add_argument("--unit", action="append",
                    help="managing unit or its code, or part of one ('IDD04')")


def _scope(args):
    """['--projects', 'P1,P2'] when a selection filter was given, else []."""
    if not any(getattr(args, k, None) for k in
               ("project", "cohort_set", "country", "region", "fy", "practice", "unit")):
        return []
    return ["--projects", ",".join(_sel(args).ids)]


def _sel(args):
    return select(project=args.project and ",".join(args.project),
                  country=args.country and ",".join(args.country),
                  region=args.region and ",".join(args.region),
                  fy=args.fy and ",".join(args.fy),
                  cohort_set=args.cohort_set and ",".join(args.cohort_set),
                  practice=args.practice and ",".join(args.practice),
                  unit=args.unit and ",".join(args.unit))


def _print_rows(rows, cols):
    if not rows:
        print("(none)")
        return
    width = {c: min(40, max(len(c), *(len(str(r.get(c, ""))) for r in rows))) for c in cols}
    print("  ".join(c.ljust(width[c]) for c in cols))
    for r in rows:
        print("  ".join(str(r.get(c, ""))[:40].ljust(width[c]) for c in cols))


def main(argv=None):
    ap = argparse.ArgumentParser(prog="wbg", description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("list", help="show the projects a selection covers")
    _filters(p)
    p = sub.add_parser("prepare", help="fetch what is missing and rebuild what changed")
    _filters(p)
    p.add_argument("--update", action="store_true",
                   help="fetch the selected projects again, even if fetched before")
    p = sub.add_parser("summary", help="where each selected project stands")
    _filters(p)
    p = sub.add_parser("export", help="write a selection's results to data/runs/")
    _filters(p)
    p.add_argument("--name", help="folder name (default: from the filters)")
    p.add_argument("--tables", action="store_true", help="also write every table as CSV")
    p = sub.add_parser("search", help="embed the sentences (only what is not cached) and "
                                      "score them against a label set (inputs/taxonomy/ or "
                                      "analysis/labels/)")
    p.add_argument("label_set", help="assets, measures, or any set in analysis/labels/")
    p.add_argument("--dry-run", action="store_true", help="what it would embed, and cost")
    p.add_argument("--embed-context", choices=("none", "section"), default="none",
                   help="embed each sentence with its section title and component heading in "
                        "front (a re-embedding; the index, the search and every run keyed on "
                        "them are new)")
    p = sub.add_parser("classify", help="grade every model paragraph against a label set "
                                        "with the decision model (only what is not stored)")
    p.add_argument("label_set", help="assets, measures, or any set in analysis/labels/")
    _filters(p)
    p.add_argument("--dry-run", action="store_true", help="what it would ask, and cost")
    p.add_argument("--limit", type=int, default=0, help="first N paragraphs (smoke test)")
    p.add_argument("--min-cosine", type=float, default=_classify.MIN_COSINE,
                   help="ask a label of a paragraph only when a unit scores at least this")
    p = sub.add_parser("candidates", help="every unit a label's patterns or its embedding "
                                          "top N find, with p_commitment, firmness and types "
                                          "(after ./wbg search)")
    p.add_argument("label_set", help="any set in analysis/labels/ or inputs/taxonomy/")
    p.add_argument("--top", type=int, default=_measure.TOP, help="pool by count (0 = off)")
    p.add_argument("--min-cosine", type=float, default=_measure.MIN_COSINE,
                   help="pool: every unit whose best cosine reaches this (0 = off)")
    p = sub.add_parser("measure", help="grade a measure's candidates with the decision model, "
                                       "with the asset map as columns (after ./wbg search "
                                       "<set> and ./wbg classify assets)")
    p.add_argument("label_set", help="a set in analysis/labels/")
    _filters(p)
    p.add_argument("--context", default="sentence", choices=("sentence", "paragraph"))
    p.add_argument("--gate", default="none", choices=("none", "link", "unit", "paragraph", "component"))
    p.add_argument("--asset-min", type=float, default=0.0)
    p.add_argument("--top", type=int, default=_measure.TOP, help="round 0 pool by count (0 = off)")
    p.add_argument("--min-cosine", type=float, default=_measure.MIN_COSINE,
                   help="round 0 pool: units whose best cosine reaches this")
    p.add_argument("--batch", type=int, default=_measure.BATCH, help="units per feedback round")
    p.add_argument("--min-yield", type=int, default=_measure.MIN_YIELD,
                   help="a round yielding fewer committed units than this is the last")
    p.add_argument("--max-rounds", type=int, default=_measure.MAX_ROUNDS)
    p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("add", help="a measure from its files to Postgres in one go: search, "
                                   "the funnel, type proposals, one load and one release; "
                                   "several sets run side by side")
    p.add_argument("label_set", nargs="+", help="set(s) in analysis/labels/")
    _filters(p)
    p.add_argument("--jobs", type=int, default=3,
                   help="sets run at once (each is one process on one core)")
    p.add_argument("--workers", type=int, default=ADD_WORKERS,
                   help="parallel requests to the decision model per set")
    p.add_argument("--assign", action="store_true",
                   help="also assign approved types (<set>_types.csv) after the funnel")
    p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("link", help="which asset each procurement package buys and which "
                                    "commitments it would carry out (decision model, stored; "
                                    "after ./wbg measure)")
    _filters(p)
    p.add_argument("--asset-top", type=int, default=_link.ASSET_TOP,
                   help="assets asked per package, its best by cosine")
    p.add_argument("--label-set", action="append", dest="sets",
                   help="measure set(s) to pair; default: every set with a current run")
    p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("types", help="propose a measure's types from its committed units "
                                     "(generative model, stored), or --assign them with the "
                                     "decision model once a person has approved the list")
    p.add_argument("label_set")
    _filters(p)
    p.add_argument("--assign", action="store_true")
    p.add_argument("--runs", type=int, default=3)
    p.add_argument("--context", default="sentence", choices=("sentence", "paragraph"))
    p.add_argument("--gate", default="none", choices=("none", "link", "unit", "paragraph", "component"))
    p.add_argument("--asset-min", type=float, default=0.0)
    p.add_argument("--dry-run", action="store_true")
    p = sub.add_parser("check", help="draw a measure's units for a person to grade, or "
                                     "--evaluate the grades: precision and recall with intervals")
    p.add_argument("label_set")
    _filters(p)
    p.add_argument("--evaluate", action="store_true")
    p.add_argument("--elusion", action="store_true",
                   help="draw (or evaluate) units the funnel never read, below the pool")
    p.add_argument("--seed", type=int, default=1)
    p.add_argument("--context", default="sentence", choices=("sentence", "paragraph"))
    p.add_argument("--gate", default="none", choices=("none", "link", "unit", "paragraph", "component"))
    p.add_argument("--asset-min", type=float, default=0.0)
    p = sub.add_parser("evaluate", help="measure the funnel for a set against the "
                                        "hand-labelled activities workbook")
    p.add_argument("label_set")
    _filters(p)
    p.add_argument("--column", required=True)
    p.add_argument("--text-regex")
    p.add_argument("--context", default="sentence", choices=("sentence", "paragraph"))
    p.add_argument("--gate", default="none", choices=("none", "link", "unit", "paragraph", "component"))
    p.add_argument("--asset-min", type=float, default=0.0)
    p = sub.add_parser("words", help="the words that set a label's candidates apart from "
                                     "the corpus: what types it has")
    p.add_argument("label_set")
    p.add_argument("label_id")
    p.add_argument("--top", type=int, default=_measure.TOP, help="the measure's pool by count")
    p.add_argument("--min-cosine", type=float, default=_measure.MIN_COSINE,
                   help="the measure's pool by cosine")
    p.add_argument("--committed", action="store_true",
                   help="only candidates with p_commitment >= 0.5")
    p = sub.add_parser("sensitivity", help="how much the cleaned text depends on the "
                                           "paragraph, component-tag and sentence settings")
    p.add_argument("--docs", type=int, default=10,
                   help="documents to re-read from the PDF for the paragraph test")
    p.add_argument("--skip-paragraphs", action="store_true",
                   help="skip the slow paragraph test")
    sub.add_parser("status", help="which code and which data release the store is at, "
                                  "and what is out of date")
    sub.add_parser("release", help="record the store's current data as a release "
                                   "(data/releases/), if it changed")
    p = sub.add_parser("pipeline", help="everything run.sh does: prepare all, release, "
                                        "load, review sheets, run report")
    p.add_argument("--update", action="store_true", help="fetch every project again")
    args = ap.parse_args(argv)

    data = data_root()
    if args.cmd == "pipeline":
        prepare(select(), update=args.update, data=data)
        _run("release.py", data=data)
        _run("load_pg.py", data=data)
        _run("review_sheets.py", data=data)
        _run("run_report.py", "--topic", "Pipeline run", data=data)
        return 0
    if args.cmd == "status":
        import release
        print("\n".join(release.status(data)))
        return 0
    if args.cmd == "release":
        _run("release.py", data=data)
        return 0
    if args.cmd == "search":
        dry = ["--dry-run"] if args.dry_run else []
        _run("embed.py", *dry, "--context", args.embed_context, data=data)
        _run("search.py", args.label_set, *dry, data=data)
        if not dry:
            _run("release.py", data=data)
        return 0
    if args.cmd == "classify":
        extra = (["--dry-run"] if args.dry_run else []) + \
            (["--limit", str(args.limit)] if args.limit else []) + \
            (["--min-cosine", str(args.min_cosine)] if args.min_cosine else [])
        _run("classify.py", args.label_set, *extra, *_scope(args), data=data)
        if not (args.dry_run or args.limit):
            _run("release.py", data=data)
        return 0
    if args.cmd == "candidates":
        _run("candidates.py", args.label_set, "--top", str(args.top), "--min-cosine",
             str(args.min_cosine), data=data)
        return 0
    if args.cmd == "measure":
        _run("measure.py", args.label_set, "--context", args.context, "--gate", args.gate,
             "--asset-min", str(args.asset_min), "--top", str(args.top),
             "--min-cosine", str(args.min_cosine), "--batch",
             str(args.batch), "--min-yield", str(args.min_yield), "--max-rounds",
             str(args.max_rounds), *(["--dry-run"] if args.dry_run else []), *_scope(args),
             data=data)
        if not args.dry_run:
            _run("release.py", data=data)
        return 0
    if args.cmd == "add":
        return add(args, data)
    if args.cmd == "link":
        _run("link_packages.py", "--asset-top", str(args.asset_top),
             *[x for st in (args.sets or []) for x in ("--label-set", st)],
             *(["--dry-run"] if args.dry_run else []), *_scope(args), data=data)
        if not args.dry_run:
            _run("release.py", data=data)
        return 0
    if args.cmd == "types":
        _run("kinds.py", args.label_set, *(["--assign"] if args.assign else []),
             "--runs", str(args.runs), "--context", args.context, "--gate", args.gate,
             "--asset-min", str(args.asset_min), *(["--dry-run"] if args.dry_run else []),
             *_scope(args), data=data)
        return 0
    if args.cmd == "check":
        _run("check_units.py", args.label_set, *(["--evaluate"] if args.evaluate else []),
             *(["--elusion"] if args.elusion else []), "--seed", str(args.seed), "--context", args.context, "--gate", args.gate,
             "--asset-min", str(args.asset_min), data=data)
        return 0
    if args.cmd == "evaluate":
        _run("evaluate_activities.py", args.label_set, "--column", args.column,
             *(["--text-regex", args.text_regex] if args.text_regex else []),
             "--context", args.context, "--gate", args.gate, "--asset-min", str(args.asset_min),
             *_scope(args), data=data)
        return 0
    if args.cmd == "words":
        _run("candidates.py", args.label_set, args.label_id, "--words", "--top", str(args.top),
             "--min-cosine", str(args.min_cosine),
             *(["--committed"] if args.committed else []), data=data)
        return 0
    if args.cmd == "sensitivity":
        _run("sensitivity.py", "--docs", str(args.docs),
             *(["--skip-paragraphs"] if args.skip_paragraphs else []), data=data)
        return 0
    sel = _sel(args)
    if args.cmd == "list":
        _print_rows(sel.rows, ["project_id", "country_code", "region", "approval_fy",
                               "status", "closing_date", "practice", "unit_code"])
        print(f"\n{len(sel)} projects")
    elif args.cmd == "prepare":
        prepare(sel, update=args.update, data=data)
    elif args.cmd == "summary":
        _print_rows(summary(sel, data), ["project_id", "country_code", "status",
                                         "closing_date", "latest_plan_date",
                                         "packages", "upcoming_packages",
                                         "upcoming_not_in_latest_plan", "upcoming_estimated_usd",
                                         "upcoming_planned_date_passed", "next_planned_date",
                                         "open_notices", "contracts_signed"])
    elif args.cmd == "export":
        print(f"wrote {export(sel, args.name, data, include_tables=args.tables)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
