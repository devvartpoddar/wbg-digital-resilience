#!/usr/bin/env python3
"""Build the hand-annotation workbook for the gold set.

WHAT THIS PRODUCES. One xlsx with two sheets. `paragraphs` holds one row per
sampled paragraph with its full cleaned text, its id and a `read` tick.
`rows` holds the annotation table, pre-seeded with blank rows that already
carry the right `paragraph_id`, so the annotator never types an id and the
join back to the corpus cannot silently fail. Column layout is fixed by
analysis/gold-set-spec.md.

THE SAMPLING RULE THAT MATTERS. No slice may be selected by any method under
test. The comparison in gold-set-spec.md ranks five candidate methods against
each other, and a sample drawn by embedding retrieval or by clustering would
hand whichever method drew it a recall advantage it did not earn. So:

  flagged   selected by a human (the portfolio tracker), which biases toward
            climate language - disclosed, not hidden
  targeted  selected by section heading and project, which are facts about
            document structure
  random    selected by a seeded RNG over prose paragraphs

None of the three consults an embedding, a cluster or a classifier.

THE HASH CHECK IS THE POINT, as it is in analysis/discovery/embed_clauses.py.
paragraphs.csv stores offsets, not text. If a clean file moved, the offsets
still resolve and still yield something - just not what was measured. Checking
`text_sha256` turns that into an abort rather than a workbook of text the
annotator reads and the scorer cannot find.

THE TRACKER IS OPTIONAL AND NEVER COMMITTED. Without --tracker the flagged
slice is skipped and the run says so. The file itself is hand-labelling owned
by the task team; .gitignore keeps it out of the repository.
"""
import argparse, csv, hashlib, os, random, re, sys
from collections import Counter, defaultdict

PROSE_BLOCKS = ("narrative", "annex")
NGRAM = 8               # words; ~3 distinct hits is far past coincidence
MIN_NGRAM_HITS = 3
PREVIEW_CHARS = 120

# Her adaptation categories -> a suggested asset and direction. Suggestions the
# annotator overwrites, never a label. Keyed on a lowercase substring of the
# category heading so a reworded heading still lands.
CATEGORY_HINTS = [
    ("resilient telecom",            "telecom network",      "resilience_of_asset"),
    ("resilient data center",        "data centre",          "resilience_of_asset"),
    ("site-specific climate risk",   "data centre",          "resilience_of_asset"),
    ("resilient it equipment",       "IT equipment",         "resilience_of_asset"),
    ("business continuity",          "",                     "resilience_of_asset"),
    ("buildings/labs/tech hubs",     "buildings",            "resilience_of_asset"),
    ("digitization of dpi",          "DPI",                  "digital_for_resilience"),
    ("climate applications",         "climate application",  "digital_for_resilience"),
    ("digital skills",               "",                     "digital_for_resilience"),
    ("devices and internet",         "access/connectivity",  "digital_for_resilience"),
    ("enabling environment",         "",                     ""),
]

ROW_COLUMNS = ["paragraph_id", "where", "kind", "quote", "label", "direction",
               "asset", "asset_in_para", "hazard", "stem"]
KINDS = ["measure", "asset", "activity", "context"]
DIRECTIONS = ["resilience_of_asset", "digital_for_resilience"]


def norm_words(text):
    """Lowercase word list with punctuation flattened, for n-gram matching."""
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).split()


def ngrams(words, n=NGRAM):
    return {" ".join(words[i:i + n]) for i in range(max(0, len(words) - n + 1))}


def read_csv(path):
    with open(path, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def prepared(data_dir, name):
    """prepared/ if the pipeline wrote there, else the flat data dir."""
    p = os.path.join(data_dir, "intermediate", "prepared", name)
    return p if os.path.exists(p) else os.path.join(data_dir, name)


def truthy(v):
    return str(v).strip().lower() in ("1", "true", "yes", "y")


def load_corpus(data_dir, log):
    """paragraph rows (prose, in scope) with text resolved and hash-checked."""
    paras = read_csv(prepared(data_dir, "paragraphs.csv"))
    links = read_csv(prepared(data_dir, "span_project.csv"))
    doc_project = {}
    for r in links:
        doc_project.setdefault(r["doc_id"], r["project_id"])

    bodies, out, skipped, bad_hash = {}, [], Counter(), []
    for r in paras:
        if r.get("block_type") not in PROSE_BLOCKS:
            skipped["not prose"] += 1
            continue
        if "in_scope" in r and not truthy(r["in_scope"]):
            skipped["out of scope"] += 1
            continue
        doc = r["doc_id"]
        if doc not in bodies:
            path = os.path.join(data_dir, "clean", f"{doc}.txt")
            if not os.path.exists(path):
                skipped["clean file missing"] += 1
                bodies[doc] = None
            else:
                with open(path, encoding="utf-8") as fh:
                    bodies[doc] = fh.read()
        if bodies[doc] is None:
            continue
        text = bodies[doc][int(r["char_start"]):int(r["char_end"])]
        want = r.get("text_sha256") or ""
        if want and hashlib.sha256(text.encode("utf-8")).hexdigest() != want:
            bad_hash.append(r["paragraph_id"])
            continue
        if not text.strip():
            skipped["resolved empty"] += 1
            continue
        out.append({
            "paragraph_id": r["paragraph_id"],
            "doc_id": doc,
            "project_id": doc_project.get(doc, ""),
            "section_label": r.get("section_label", ""),
            "block_type": r["block_type"],
            "token_count": int(r.get("token_count") or 0),
            "text": text,
        })
    if bad_hash:
        raise SystemExit(
            f"build_sample: {len(bad_hash)} paragraphs do not match their "
            f"recorded text_sha256 (first: {', '.join(bad_hash[:3])}).\n"
            f"  paragraphs.csv and data/clean/ have diverged - re-run "
            f"src/clean.py and src/segment.py before building the sample.")
    for why, n in sorted(skipped.items()):
        log(f"  skipped {n:,} paragraphs: {why}")
    log(f"  usable prose paragraphs: {len(out):,} "
        f"across {len({p['project_id'] for p in out if p['project_id']})} projects")
    return out


def read_tracker(path, log):
    """[(project_id, excerpt, [category headings])] from '2. Activities Review'.

    Column letters are read from the header row rather than hard-coded, so a
    reordered or newly inserted column does not silently shift the read.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((s for s in wb.sheetnames if "activities review" in s.lower()), None)
    if sheet is None:
        raise SystemExit(f"build_sample: no 'Activities Review' sheet in {path}")
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows)
                 if r and any(str(c).strip() == "Project ID" for c in r if c))
    hdr = [str(c).strip() if c is not None else "" for c in rows[hdr_i]]
    idx = {name: i for i, name in enumerate(hdr) if name}

    def need(name):
        if name not in idx:
            raise SystemExit(f"build_sample: tracker has no '{name}' column")
        return idx[name]

    c_proj, c_text = need("Project ID"), need("Measures / Activities")
    # every column between the Adaptation flag and the Mitigation flag
    lo, hi = need("Adaptation"), need("Mitigation")
    cat_cols = [(i, hdr[i]) for i in range(lo + 1, hi) if hdr[i]]

    out = []
    for r in rows[hdr_i + 1:]:
        if not r or c_text >= len(r) or not r[c_text]:
            continue
        text = str(r[c_text]).strip()
        if not text:
            continue
        m = re.search(r"P\d{5,}", str(r[c_proj] or ""))
        cats = [name for i, name in cat_cols if i < len(r) and r[i]]
        out.append((m.group(0) if m else "", text, cats))
    log(f"  tracker sheet {sheet!r}: {len(out)} excerpt rows, "
        f"{len({p for p, _, _ in out if p})} projects, "
        f"{len(cat_cols)} category columns")
    log(f"  excerpt rows with no category at all: "
        f"{sum(1 for _, _, c in out if not c)}")
    return out


def hint_for(categories):
    for cat in categories:
        low = cat.lower()
        for key, asset, direction in CATEGORY_HINTS:
            if key in low:
                return asset, direction
    return "", ""


def match_excerpts(excerpts, paras, log):
    """paragraph_id -> (categories, had_no_category) for matched paragraphs."""
    by_project = defaultdict(list)
    for p in paras:
        by_project[p["project_id"]].append(p)

    index = {}
    for proj, plist in by_project.items():
        inv = defaultdict(set)
        for p in plist:
            for g in ngrams(norm_words(p["text"])):
                inv[g].add(p["paragraph_id"])
        index[proj] = inv

    hits, matched_excerpts = {}, 0
    for proj, text, cats in excerpts:
        inv = index.get(proj)
        if inv is None:
            continue
        counts = Counter()
        for g in ngrams(norm_words(text)):
            for pid in inv.get(g, ()):
                counts[pid] += 1
        found = [pid for pid, n in counts.items() if n >= MIN_NGRAM_HITS]
        if found:
            matched_excerpts += 1
        for pid in found:
            prev_cats, prev_none = hits.get(pid, ([], True))
            hits[pid] = (sorted(set(prev_cats) | set(cats)),
                         prev_none and not cats)
    log(f"  excerpts matched to >=1 paragraph: {matched_excerpts}/{len(excerpts)}"
        f"  ({matched_excerpts / max(1, len(excerpts)):.0%})")
    log(f"  distinct paragraphs matched: {len(hits):,}")
    log(f"  of those, from excerpts with no category: "
        f"{sum(1 for _, none in hits.values() if none):,}")
    return hits


def pick(pool, n, rng, exclude):
    """Deterministic sample of up to n, skipping already-taken ids."""
    avail = sorted((p for p in pool if p["paragraph_id"] not in exclude),
                   key=lambda p: p["paragraph_id"])
    rng.shuffle(avail)
    return avail[:n]


def build(args, log):
    log("corpus")
    paras = load_corpus(args.data, log)
    by_id = {p["paragraph_id"]: p for p in paras}
    rng = random.Random(args.seed)
    chosen, slice_of, hint_of = [], {}, {}

    hits = {}
    if args.tracker:
        log("tracker")
        excerpts = read_tracker(args.tracker, log)
        log("matching excerpts to paragraphs")
        hits = match_excerpts(excerpts, paras, log)
        # paragraphs from uncategorised excerpts first: candidate climate
        # content her taxonomy had no bucket for, and so the likeliest place
        # for a measure type with no name yet.
        priority = sorted(pid for pid, (_, none) in hits.items() if none)
        rest = sorted(pid for pid, (_, none) in hits.items() if not none)
        rng.shuffle(priority); rng.shuffle(rest)
        take = [pid for pid in priority + rest if pid in by_id][:args.flagged]
        for pid in take:
            chosen.append(by_id[pid]); slice_of[pid] = "flagged"
            hint_of[pid] = hint_for(hits[pid][0])
        log(f"  flagged slice: {len(take)} paragraphs "
            f"({sum(1 for pid in take if hits[pid][1])} from uncategorised excerpts)")
    else:
        log("tracker: not supplied - the flagged slice is SKIPPED")

    taken = {p["paragraph_id"] for p in chosen}
    sect = re.compile(args.targeted_sections, re.I) if args.targeted_sections else None
    if sect:
        pool = [p for p in paras if sect.search(p["section_label"] or "")]
        got = pick(pool, args.targeted, rng, taken)
        for p in got:
            chosen.append(p); slice_of[p["paragraph_id"]] = "targeted"
        taken |= {p["paragraph_id"] for p in got}
        log(f"targeted slice: {len(got)} of {len(pool):,} paragraphs whose "
            f"section heading matches /{args.targeted_sections}/")
        if len(got) < args.targeted:
            log(f"  SHORT by {args.targeted - len(got)} - widen "
                f"--targeted-sections or lower --targeted")

    got = pick(paras, args.random_n, rng, taken)
    for p in got:
        chosen.append(p); slice_of[p["paragraph_id"]] = "random"
    log(f"random slice: {len(got)} paragraphs, uniform over prose, seed {args.seed}")

    write_workbook(args.out, chosen, slice_of, hint_of, args, log)
    log(f"\ntotal paragraphs in the workbook: {len(chosen)}")
    counts = Counter(slice_of.values())
    log(f"  by slice: " + ", ".join(f"{k} {v}" for k, v in sorted(counts.items())))
    log(f"blank annotation rows: {sum(blank_rows(slice_of[p['paragraph_id']], args) for p in chosen)}")
    return 0


def blank_rows(slice_name, args):
    return args.rows_random if slice_name == "random" else args.rows_rich


def write_workbook(path, chosen, slice_of, hint_of, args, log):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as gcl
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    head = Font(bold=True)
    grey = PatternFill("solid", fgColor="EEEEEE")
    wrap = Alignment(wrap_text=True, vertical="top")

    ws = wb.active
    ws.title = "paragraphs"
    cols = ["paragraph_id", "project_id", "doc_id", "section_label", "slice",
            "tokens", "read", "text"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        ws.cell(1, c).font = head
        ws.cell(1, c).fill = grey
    for p in chosen:
        ws.append([p["paragraph_id"], p["project_id"], p["doc_id"],
                   p["section_label"], slice_of[p["paragraph_id"]],
                   p["token_count"], "", p["text"]])
    for w, c in zip((26, 12, 12, 30, 10, 8, 7, 150), range(1, len(cols) + 1)):
        ws.column_dimensions[gcl(c)].width = w
    for r in range(2, len(chosen) + 2):
        ws.cell(r, 8).alignment = wrap
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{gcl(len(cols))}{len(chosen) + 1}"
    yn = DataValidation(type="list", formula1='"Y"', allow_blank=True)
    ws.add_data_validation(yn)
    yn.add(f"G2:G{len(chosen) + 1}")

    rs = wb.create_sheet("rows")
    rs.append(ROW_COLUMNS)
    for c in range(1, len(ROW_COLUMNS) + 1):
        rs.cell(1, c).font = head
        rs.cell(1, c).fill = grey
    n = 0
    for p in chosen:
        pid = p["paragraph_id"]
        asset, direction = hint_of.get(pid, ("", ""))
        preview = re.sub(r"\s+", " ", p["text"])[:PREVIEW_CHARS]
        for _ in range(blank_rows(slice_of[pid], args)):
            rs.append([pid, preview, "", "", "", direction, asset, "", "", ""])
            n += 1
    for w, c in zip((26, 46, 12, 64, 30, 22, 22, 13, 18, 40),
                    range(1, len(ROW_COLUMNS) + 1)):
        rs.column_dimensions[gcl(c)].width = w
    rs.freeze_panes = "C2"
    rs.auto_filter.ref = f"A1:{gcl(len(ROW_COLUMNS))}{n + 1}"
    dv_kind = DataValidation(type="list", formula1=f'"{",".join(KINDS)}"',
                             allow_blank=True)
    dv_dir = DataValidation(type="list", formula1=f'"{",".join(DIRECTIONS)}"',
                            allow_blank=True)
    dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    for dv, col in ((dv_kind, "C"), (dv_dir, "F"), (dv_yn, "H")):
        rs.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n + 1}")
    for r in range(2, n + 2):
        rs.cell(r, 4).alignment = wrap
        rs.cell(r, 10).alignment = wrap

    tmp = path + ".part"
    wb.save(tmp)
    os.replace(tmp, path)
    log(f"\nwrote {path}")


def inspect(args, log):
    """Report what there is to sample from, before any slice is chosen."""
    paras = load_corpus(args.data, log)
    secs = Counter((p["section_label"] or "(blank)").strip() for p in paras)
    log(f"\ndistinct section labels: {len(secs)}")
    log("the 40 largest:")
    for name, n in secs.most_common(40):
        log(f"  {n:6,}  {name[:88]}")
    probe = ["cyber", "security", "data cent", "datacent", "identif", "dpi",
             "digital public", "connectiv", "resilien", "climate", "annex",
             "component"]
    log("\nparagraphs whose section heading contains:")
    for k in probe:
        rx = re.compile(k, re.I)
        log(f"  {sum(1 for p in paras if rx.search(p['section_label'] or '')):6,}  {k}")
    if args.tracker:
        read_tracker(args.tracker, log)
    return 0


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default="data")
    ap.add_argument("--tracker", default="",
                    help="Portfolio_tracker.xlsx; omit to skip the flagged slice")
    ap.add_argument("--out", default="goldset_sample.xlsx")
    ap.add_argument("--seed", type=int, default=20260918)
    ap.add_argument("--flagged", type=int, default=80)
    ap.add_argument("--targeted", type=int, default=40)
    ap.add_argument("--random-n", type=int, default=60)
    ap.add_argument("--targeted-sections", default="",
                    help="regex over section_label; set it after --inspect")
    ap.add_argument("--rows-rich", type=int, default=4,
                    help="blank annotation rows per flagged/targeted paragraph")
    ap.add_argument("--rows-random", type=int, default=2)
    ap.add_argument("--inspect", action="store_true",
                    help="report section labels and tracker shape; writes nothing")
    args = ap.parse_args()

    def log(m):
        print(m, flush=True)

    return inspect(args, log) if args.inspect else build(args, log)


if __name__ == "__main__":
    raise SystemExit(main())
