#!/usr/bin/env python3
"""Fill the asset and direction suggestion columns from the portfolio tracker.

WHY THIS IS A SEPARATE SCRIPT AND NOT PART OF build_sample.py. The tracker is
hand-labelling owned by the task team and stays on the annotator's own machine:
neither this repository nor the server holds anything but publicly disclosed
PAD text. build_sample.py therefore runs on the server, where the corpus is,
and never sees the tracker.

This script closes the gap without moving either file. By the time the workbook
exists it already carries each sampled paragraph's full text, so the excerpts
can be matched against the workbook itself rather than against the 20 MB of
cleaned documents. Run it on the machine that holds the tracker:

    python3 prefill.py --workbook goldset_sample.xlsx \
                       --tracker Portfolio_tracker.xlsx

It edits the workbook in place (after writing a .bak) and reports what it
matched. Nothing leaves the machine.

WHAT IT WRITES, AND WHAT IT WILL NOT OVERWRITE. Only the `asset` and
`direction` cells of blank annotation rows, and only where they are empty. A
cell the annotator has already typed into is never touched, so the script is
safe to re-run midway through annotation. It never writes `kind`, `quote`,
`label` or anything else - those are the annotator's alone, and a suggestion
there would be a label this project did not earn.

IT ALSO MEASURES SOMETHING. The match rate is the share of the workbook's
climate-section paragraphs the tracker had already caught, which is a coverage
statement about the existing hand-labelling that nothing else produces.
"""
import argparse, importlib.util, os, re, shutil, sys
from collections import Counter, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location(
    "build_sample", os.path.join(HERE, "build_sample.py"))
BS = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(BS)


def workbook_paragraphs(ws):
    """paragraph_id -> (project_id, text) from the `paragraphs` sheet."""
    hdr = [c.value for c in ws[1]]
    need = ("paragraph_id", "project_id", "text")
    for name in need:
        if name not in hdr:
            raise SystemExit(f"prefill: the paragraphs sheet has no {name!r} column")
    i_pid, i_proj, i_text = (hdr.index(n) for n in need)
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[i_pid]:
            continue
        out[str(row[i_pid])] = (str(row[i_proj] or ""), str(row[i_text] or ""))
    return out


def match(paras, excerpts, log):
    """paragraph_id -> [category headings], matched within the same project."""
    by_project = defaultdict(list)
    for pid, (proj, text) in paras.items():
        by_project[proj].append((pid, BS.ngrams(BS.norm_words(text))))

    hits = defaultdict(set)
    matched_excerpts = 0
    for proj, text, cats in excerpts:
        grams = BS.ngrams(BS.norm_words(text))
        found = [pid for pid, pg in by_project.get(proj, ())
                 if len(grams & pg) >= BS.MIN_NGRAM_HITS]
        if found:
            matched_excerpts += 1
        for pid in found:
            hits[pid].update(cats)
    log(f"tracker excerpts matching a paragraph in the workbook: "
        f"{matched_excerpts}/{len(excerpts)}")
    log(f"workbook paragraphs the tracker had already caught: "
        f"{len(hits)}/{len(paras)}  ({len(hits) / max(1, len(paras)):.0%})")
    return {pid: sorted(c) for pid, c in hits.items()}


def apply(ws, hits, log):
    hdr = [c.value for c in ws[1]]
    for name in ("paragraph_id", "asset", "direction"):
        if name not in hdr:
            raise SystemExit(f"prefill: the rows sheet has no {name!r} column")
    c_pid = hdr.index("paragraph_id") + 1
    c_asset = hdr.index("asset") + 1
    c_dir = hdr.index("direction") + 1

    filled, kept, no_hint = 0, 0, Counter()
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(r, c_pid).value
        if not pid or str(pid) not in hits:
            continue
        cats = hits[str(pid)]
        asset, direction = BS.hint_for(cats)
        if not asset and not direction:
            for cat in cats:
                no_hint[cat] += 1
            continue
        for col, value in ((c_asset, asset), (c_dir, direction)):
            cell = ws.cell(r, col)
            if not value:
                continue
            if cell.value not in (None, ""):
                kept += 1
                continue
            cell.value = value
            filled += 1
    log(f"suggestion cells filled: {filled}")
    log(f"cells left alone because they already had a value: {kept}")
    if no_hint:
        log("tracker categories with no asset/direction hint "
            "(add them to CATEGORY_HINTS in build_sample.py if they matter):")
        for cat, n in no_hint.most_common():
            log(f"  {n:4d}  {cat}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--tracker", required=True)
    ap.add_argument("--dry-run", action="store_true",
                    help="report the match; write nothing")
    args = ap.parse_args()

    def log(m):
        print(m, flush=True)

    import openpyxl
    for p in (args.workbook, args.tracker):
        if not os.path.exists(p):
            raise SystemExit(f"prefill: no such file: {p}")

    excerpts = BS.read_tracker(args.tracker, log)
    wb = openpyxl.load_workbook(args.workbook)
    for sheet in ("paragraphs", "rows"):
        if sheet not in wb.sheetnames:
            raise SystemExit(
                f"prefill: {args.workbook} has no {sheet!r} sheet - is this the "
                f"workbook build_sample.py produced?")
    paras = workbook_paragraphs(wb["paragraphs"])
    log(f"workbook paragraphs: {len(paras)}")
    hits = match(paras, excerpts, log)

    if args.dry_run:
        log("\n--dry-run: nothing written")
        return 0
    apply(wb["rows"], hits, log)
    backup = args.workbook + ".bak"
    if not os.path.exists(backup):
        shutil.copy2(args.workbook, backup)
        log(f"kept a copy at {backup}")
    tmp = args.workbook + ".part"
    wb.save(tmp)
    os.replace(tmp, args.workbook)
    log(f"wrote {args.workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
