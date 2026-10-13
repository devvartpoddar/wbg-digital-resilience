"""Classification with a decision model (src/classify.py), on an invented store
and a fake endpoint. No document text is committed."""
import csv
import hashlib
import json
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                "src"))
import classify  # noqa: E402
from clean import PARA_COLS  # noqa: E402

TEXTS = ["The project will finance towers on raised plinths.",
         "Background on the economy.",
         "The project will finance towers on raised plinths."]   # a repeat


def _csv(path, cols, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})


@pytest.fixture
def store(tmp_path, monkeypatch):
    data = tmp_path / "data"
    text = "\n\n".join(TEXTS)
    (data / "appraisal" / "text").mkdir(parents=True)
    (data / "appraisal" / "text" / "D.txt").write_text(text, encoding="utf-8")
    rows, at = [], 0
    for i, t in enumerate(TEXTS, 1):
        rows.append({"paragraph_id": f"D:p{i:05d}", "doc_id": "D", "project_ids": "P1",
                     "block": "narrative", "for_model": "true", "char_start": at,
                     "char_end": at + len(t), "section_title": "Components",
                     "component_number": "1" if i != 2 else "",
                     "text_sha256": hashlib.sha256(t.encode()).hexdigest()})
        at += len(t) + 2
    rows.append({"paragraph_id": "D:p00004", "doc_id": "D", "block": "footnote",
                 "for_model": "true", "char_start": 0, "char_end": 3})
    _csv(str(data / "appraisal" / "paragraphs.csv"), PARA_COLS, rows)
    _csv(str(data / "appraisal" / "components.csv"), ["doc_id", "source", "number", "name"],
         [{"doc_id": "D", "source": "datasheet", "number": "1", "name": "Connectivity"}])
    tax = tmp_path / "taxonomy"
    _csv(str(tax / "assets.csv"), ["asset_id", "asset_name", "definition", "active"], [
        {"asset_id": "towers", "asset_name": "Mobile networks", "definition": "masts",
         "active": "true"},
        {"asset_id": "old", "asset_name": "Old", "definition": "x", "active": "false"}])
    _csv(str(tax / "assets_scale.csv"), ["grade", "description"], [
        {"grade": 0, "description": "Not about {name}."},
        {"grade": 1, "description": "{name} mentioned."},
        {"grade": 2, "description": "{name} financed."}])
    calls = []

    def fake(body, **kw):
        calls.append(body)
        hit = "towers" in body["state"]["passage"]
        probs = {"0": 0.1, "1": 0.1, "2": 0.8} if hit else {"0": 0.9, "1": 0.1, "2": 0.0}
        return {"model": "typesafe/jev-1.13-20260917", "usage": {"cost": 0.00001},
                "answers": {"towers": {"type": "score", "score": 1.7, "confidence": 0.8,
                                       "probabilities": probs}}}
    monkeypatch.setattr(classify, "decide", fake)
    return str(data), str(tax), calls


def _run(store, **kw):
    data, tax, _ = store
    return classify.classify("assets", data, api_key="k", taxonomy=tax, log=lambda *_: None,
                             workers=1, **kw)


def test_each_label_is_one_graded_question_naming_it():
    labels = [{"id": "towers", "name": "Mobile networks", "definition": "masts"}]
    q = classify.questions(labels, ["Not about {name}.", "{name} financed."])
    assert q["towers"]["type"] == "score"
    assert "Mobile networks means: masts" in q["towers"]["instructions"]
    assert q["towers"]["criteria"] == ["Not about Mobile networks.", "Mobile networks financed."]


def test_every_narrative_paragraph_is_graded_and_a_repeat_asked_once(store):
    out = _run(store)
    _, _, calls = store
    assert len(calls) == 2                                  # three paragraphs, one repeat
    assert calls[0]["state"]["component"] == "1 Connectivity"
    rows = list(csv.DictReader(open(os.path.join(out, "scores.csv"))))
    assert [(r["paragraph_id"], r["grade"]) for r in rows] == [
        ("D:p00001", "2"), ("D:p00002", "0"), ("D:p00003", "2")]


def test_a_rerun_reads_the_store_and_asks_nothing(store):
    first = _run(store)
    _, _, calls = store
    n = len(calls)
    assert _run(store) == first and len(calls) == n


def test_an_incomplete_answer_is_not_stored(store, monkeypatch):
    monkeypatch.setattr(classify, "decide", lambda body, **kw: {"answers": {}})
    assert _run(store) is None
    data, _, _ = store
    assert not os.path.exists(classify.cache_root(data, classify.MODEL))


def test_the_grades_must_be_given(store):
    _, tax, _ = store
    os.remove(os.path.join(tax, "assets_scale.csv"))
    with pytest.raises(SystemExit):
        _run(store)


def test_a_dry_run_asks_nothing(store):
    _, _, calls = store
    assert _run(store, dry_run=True) is None and calls == []


def test_a_graded_sheet_becomes_labels_and_a_measurement(store):
    import check_sheet
    import evaluate
    from openpyxl import load_workbook
    data, tax, _ = store
    _run(store)
    path, n = check_sheet.build("assets", data, per=5, seed=1, taxonomy=tax)
    assert n == 3                                 # two strata: grade 2 (two), grade 0 (one)
    wb = load_workbook(path)
    ws = wb["grade these"]
    head = [c.value for c in ws[1]]
    assert "grade" in head and "p0" not in head   # the model's answer is not shown
    for row in ws.iter_rows(min_row=2):
        passage = row[head.index("passage")].value
        row[head.index("grade")].value = 2 if "towers" in passage else 0
    wb.save(path)
    out, graded = evaluate.evaluate("assets", data, seed=1, taxonomy=tax)
    assert graded == 3
    labels = list(csv.DictReader(open(os.path.join(data, "labels", "assets.csv"))))
    assert {r["grade"] for r in labels} == {"0", "2"}
    assert all(len(r["text_sha256"]) == 64 for r in labels)
    report = open(out).read()
    assert "| 0.5 | 1.00 | 1.00 |" in report
    with pytest.raises(SystemExit):              # a graded sheet is never overwritten
        check_sheet.build("assets", data, per=5, seed=1, taxonomy=tax)
