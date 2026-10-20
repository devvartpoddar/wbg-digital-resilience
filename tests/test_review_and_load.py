"""Unit tests for the data root, the review sheets and the Postgres loader.

Invented text throughout; no document text is ever committed. The loader test
needs a real Postgres and runs only when $WBG_PG_TEST names one.
"""
import csv
import os
import subprocess
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))

import paths  # noqa: E402
import load_pg as L  # noqa: E402


# --- where data goes ---------------------------------------------------------

def test_env_overrides_data_root(monkeypatch, tmp_path):
    monkeypatch.setenv("WBG_DATA", str(tmp_path))
    assert paths.data_root() == str(tmp_path)


def test_worktree_resolves_to_main_checkout(tmp_path):
    """The reason paths.py exists: a board card's worktree is deleted when the
    card completes, so data must land in the main checkout, not the worktree."""
    main = tmp_path / "repo"
    main.mkdir()
    run = lambda *a, cwd=main: subprocess.run(["git", *a], cwd=cwd, check=True,  # noqa: E731
                                              capture_output=True)
    run("init", "-q")
    run("-c", "user.email=t@t", "-c", "user.name=t", "commit", "-q", "--allow-empty", "-m", "x")
    wt = main / ".worktrees" / "t_1"
    run("worktree", "add", "-q", str(wt))
    assert paths.main_checkout(str(wt)) == str(main)


def test_not_a_repo_falls_back_to_itself(tmp_path):
    assert paths.main_checkout(str(tmp_path)) == str(tmp_path)


def test_convert_types_and_nulls():
    bad = {}
    assert L.convert("n_tokens", "12", bad) == 12
    assert L.convert("estimated_amount", "1,250.5", bad) == 1250.5
    assert L.convert("description", "", bad) is None
    assert L.convert("estimated_amount", "n/a", bad) is None
    assert bad == {"estimated_amount": 1}


@pytest.mark.skipif(not os.environ.get("WBG_PG_TEST"), reason="set WBG_PG_TEST to a scratch database")
def test_load_is_idempotent(tmp_path, monkeypatch, capsys):
    import psycopg
    with psycopg.connect(os.environ["WBG_PG_TEST"], autocommit=True) as conn:
        conn.execute("DROP SCHEMA IF EXISTS wbg CASCADE")
    data = tmp_path
    (data / "appraisal" / "text").mkdir(parents=True)
    (data / "raw").mkdir()
    (data / "appraisal" / "text" / "D1.txt").write_text("Alpha beta gamma. Delta epsilon.",
                                                       encoding="utf-8")
    with open(data / "raw" / "documents.csv", "w", newline="") as fh:
        fh.write("doc_id,title\nD1,Test\n")
    from clean import PARA_COLS
    with open(data / "appraisal" / "paragraphs.csv", "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=PARA_COLS)
        w.writeheader()
        for i, (a, z, n) in enumerate([(0, 17, 3), (18, 32, 2)], 1):
            w.writerow({"paragraph_id": f"D1:p{i:05d}", "doc_id": "D1", "project_ids": "P1",
                        "ordinal": i, "block": "narrative", "for_model": "true",
                        "char_start": a, "char_end": z, "n_tokens": n, "section_title": "x"})
    argv = ["load_pg.py", "--data", str(data), "--dsn", os.environ["WBG_PG_TEST"]]
    monkeypatch.setattr(sys, "argv", argv)
    L.main()
    monkeypatch.setattr(sys, "argv", argv)
    L.main()
    out = capsys.readouterr().out
    assert "paragraphs                  2 rows" in out and "paragraphs           unchanged" in out
    with psycopg.connect(os.environ["WBG_PG_TEST"]) as conn:
        texts = [r[0] for r in conn.execute("SELECT text FROM wbg.paragraphs ORDER BY 1")]
    assert texts == ["Alpha beta gamma.", "Delta epsilon."]


def test_dates_flags_and_stamps_are_typed_for_metabase():
    bad = {}
    assert L.col_type("planned_date") == "date" and L.col_type("status_as_of") == "date"
    assert L.col_type("fetched_at") == "timestamptz" and L.col_type("for_model") == "boolean"
    assert L.convert("planned_date", "2024-06-30", bad) == "2024-06-30"
    assert L.convert("planned_date", "30/06/2024", bad) is None
    assert L.convert("for_model", "true", bad) is True
    assert L.convert("for_model", "yes", bad) is None
    assert L.convert("fetched_at", "2026-10-13T14:31:08Z", bad) == "2026-10-13T14:31:08Z"
    assert L.convert("fetched_at", "yesterday", bad) is None
    assert bad == {"planned_date": 1, "for_model": 1, "fetched_at": 1}


def test_found_sets_reads_the_run_names(monkeypatch):
    monkeypatch.setattr(L.labels, "label_sets", lambda: ["assets", "cra", "ews"])
    sha = "cra-abc123def456|ews-0123456789ab|cra-assign-aaaaaaaaaaaa|9f8e7d6c5b4a"
    assert L.found_sets(sha) == ["cra", "ews"]
    assert L.found_sets(None) == []


@pytest.mark.skipif(not os.environ.get("WBG_PG_TEST"), reason="set WBG_PG_TEST to a scratch database")
def test_a_set_without_a_current_run_keeps_its_rows():
    """Replacing one set's rows, even when the table's columns changed,
    leaves every other set's rows in place (a measure being re-run must not
    vanish from the dashboards)."""
    import psycopg
    with psycopg.connect(os.environ["WBG_PG_TEST"], autocommit=True) as conn:
        conn.execute("DROP SCHEMA IF EXISTS wbg CASCADE")
        L.prepare(conn)
        cols = ["label_set", "unit_id", "text"]
        L.replace_rows(conn, "measure_findings", cols,
                       [{"label_set": "a", "unit_id": "u1", "text": "x"},
                        {"label_set": "b", "unit_id": "u2", "text": "y"}], "label_set", ["a", "b"], "s1")
        # New code adds a column: set a is replaced, set b (no current run) is carried over.
        L.replace_rows(conn, "measure_findings", cols + ["firmness"],
                       [{"label_set": "a", "unit_id": "u3", "text": "z", "firmness": "activity"}],
                       "label_set", ["a"], "s2")
        got = conn.execute("SELECT label_set, unit_id, firmness FROM wbg.measure_findings "
                           "ORDER BY 1").fetchall()
        assert got == [("a", "u3", "activity"), ("b", "u2", None)]


def test_cell_text_marks_what_xml_refuses_and_keeps_the_page_break():
    """A control character is marked, not deleted: two words either side of it
    must not run together into a sentence a person grades. A form feed is a
    page break the reader wants to see, so it keeps its own visible marker."""
    import check_units
    assert check_units.cell_text("are: \uffff technical") == "are: \ufffd technical"
    assert check_units.cell_text("law\uffff)") == "law\ufffd)"
    assert check_units.cell_text("a\x00b\x1fc") == "a\ufffdb\ufffdc"
    assert check_units.cell_text("page\fbreak") == "page\n[page break]\nbreak"
    assert check_units.cell_text("tabs\tand\nnewlines\rstay") == "tabs\tand\nnewlines\rstay"
    assert check_units.cell_text(None) == ""


def test_a_noncharacter_from_the_parse_does_not_stop_the_sheet(tmp_path):
    """The parse carries U+FFFF where the PDF's extractor met a glyph it could
    not map (nine files of the corpus do), and lxml refuses it at wb.save - so
    the sheet could not be written at all and the truncated file blocked the
    seed. The sheet is now written, with the character marked."""
    import check_units
    from openpyxl import load_workbook
    path = str(tmp_path / "sheet.xlsx")
    check_units._write_xlsx(path, [{"row_id": 1, "measure": "renewable energy",
                                    "measure means": "a measure", "project": "P1",
                                    "section": "A", "component": "C1",
                                    "sentence": "The activities are: \uffff technical assistance.",
                                    "paragraph": "p\uffffq", "grade": "", "asset": "", "note": ""}],
                            ["no"], ["datacentre"])
    got = [c.value for c in load_workbook(path)["grade these"][2]]
    assert got[6] == "The activities are: \ufffd technical assistance."
    assert got[7] == "p\ufffdq"


def test_an_evaluation_writes_its_headline_numbers(tmp_path):
    """check --evaluate keeps one row per measure and question in
    data/labels/measure_quality.csv; a later evaluation replaces it."""
    import check_units
    key = {"question_sha256": "q1", "measure_dir": "cra-abc"}
    rows = [{"grade": 3, "model_grade": 3}, {"grade": 3, "model_grade": 2},
            {"grade": 0, "model_grade": 3}, {"grade": 3, "model_grade": 3}]
    check_units.write_quality(str(tmp_path), "cra", key, rows, 3)
    check_units.write_quality(str(tmp_path), "cra", key, rows[:2], 3)
    got = list(csv.DictReader(open(tmp_path / "labels" / "measure_quality.csv")))
    assert len(got) == 1 and got[0]["graded_units"] == "2"
    assert got[0]["precision"] == "1.000" and got[0]["recall"] == "0.500"
    assert got[0]["version"]                     # the measure file's version column
