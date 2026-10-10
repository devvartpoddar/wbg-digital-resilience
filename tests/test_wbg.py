"""The packaging module (src/wbg.py): selecting, preparing, summarising.
Invented cohort and tables throughout."""
import csv
import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))

import wbg  # noqa: E402

COHORT = [
    {"project_id": "P000001", "country_code": "AA", "region": "Western and Central Africa",
     "approval_fy": "2024", "practice": "", "included": "true", "exclusion_reason": "",
     "cohort_version": "x"},
    {"project_id": "P000002", "country_code": "BB", "region": "South Asia",
     "approval_fy": "2022", "practice": "", "included": "true", "exclusion_reason": "",
     "cohort_version": "x"},
    {"project_id": "P000003", "country_code": "", "region": "", "approval_fy": "",
     "practice": "", "included": "true", "exclusion_reason": "", "cohort_version": "x"},
    {"project_id": "P000004", "country_code": "AA", "region": "Western and Central Africa",
     "approval_fy": "2024", "practice": "", "included": "false",
     "exclusion_reason": "no_documents_disclosed", "cohort_version": "x"},
]


def _write(path, rows):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


@pytest.fixture
def cohort(tmp_path):
    path = str(tmp_path / "cohort.csv")
    _write(path, COHORT)
    return path


def test_no_filter_selects_every_included_project(cohort):
    assert wbg.select(path=cohort).ids == ["P000001", "P000002", "P000003"]


@pytest.mark.parametrize("kw,ids", [
    ({"project": "P000002"}, ["P000002"]),
    ({"project": "P000001,P000002"}, ["P000001", "P000002"]),
    ({"country": "aa"}, ["P000001"]),
    ({"region": "western"}, ["P000001"]),
    ({"fy": "2022"}, ["P000002"]),
    ({"region": "western", "fy": "2022"}, []),
])
def test_filters_combine(cohort, kw, ids):
    assert wbg.select(path=cohort, **kw).ids == ids


def test_a_project_outside_the_cohort_is_refused(cohort):
    with pytest.raises(SystemExit):
        wbg.select(project="P000004", path=cohort)


def test_a_selection_is_named_for_its_filters(cohort):
    assert wbg.select(region="Western and Central", fy="2024", path=cohort).name == \
        "western-and-central-fy2024"
    assert wbg.select(path=cohort).name == "all"


@pytest.fixture
def store(tmp_path, cohort):
    data = str(tmp_path / "data")
    _write(os.path.join(data, "raw", "documents.csv"), [
        {"doc_id": "D1", "project_ids": "P000001", "doc_kind": "pad",
         "disclosure_date": "2023-01-01"},
        {"doc_id": "D2", "project_ids": "P000001|P000002", "doc_kind": "restructuring",
         "disclosure_date": "2025-01-01"}])
    _write(os.path.join(data, "appraisal", "components.csv"), [
        {"project_ids": "P000001", "doc_id": "D1", "doc_kind": "pad",
         "disclosure_date": "2023-01-01", "source": "datasheet", "level": "component",
         "number": "1", "name": "Rural radio masts", "cost_usd_m": "10.00", "action": "",
         "page": "5"},
        {"project_ids": "P000001", "doc_id": "D2", "doc_kind": "restructuring",
         "disclosure_date": "2025-01-01", "source": "restructuring", "level": "component",
         "number": "1", "name": "Rural radio masts", "cost_usd_m": "12.00", "action": "",
         "page": "3"}])
    pk = {"package_id": "", "project_id": "P000001", "borrower_ref": "", "status": "",
          "planned_date": "", "estimated_amount": "", "plan_disclosure_date": "2026-01-01",
          "description_clean": "x", "category": "goods", "method": "RFB"}
    _write(os.path.join(data, "procurement", "packages.csv"), [
        dict(pk, package_id="a", status="Pending Implementation", planned_date="2030-01-01",
             estimated_amount="100.00"),
        dict(pk, package_id="b", status="Pending", planned_date="2020-01-01",
             estimated_amount="50.00"),
        dict(pk, package_id="c", status="Signed", planned_date="2021-01-01")])
    _write(os.path.join(data, "procurement", "packages_raw.csv"), [
        {"package_version_id": "R1:00000", "project_id": "P000001"}])
    _write(os.path.join(data, "procurement", "notices.csv"), [
        {"notice_id": "N1", "project_id": "P000001", "publication_date": "2026-01-01",
         "deadline_date": "2030-02-01", "notice_type": "x", "borrower_ref": "",
         "description_clean": "x", "category": "goods", "method": "RFB"}])
    _write(os.path.join(data, "procurement", "awards.csv"), [
        {"contract_id": "C1", "project_id": "P000001", "signed_date": "2025-06-01",
         "total_amount": "75.50"}])
    return data


def test_summary_says_where_each_project_stands(store, cohort):
    sel = wbg.select(project="P000001", path=cohort)
    row = wbg.summary(sel, data=store, today="2026-10-10")[0]
    assert row["documents"] == 2 and row["latest_document_date"] == "2025-01-01"
    # The most recent document's list, with its cost.
    assert row["components"] == "1. Rural radio masts (US$12.00m)"
    assert row["packages"] == 3 and row["upcoming_packages"] == 2
    assert row["upcoming_estimated_usd"] == "150.00"
    assert row["upcoming_planned_date_passed"] == 1
    assert row["next_planned_date"] == "2030-01-01"
    assert row["open_notices"] == 1
    assert row["contracts_value_usd"] == "75.50"


def test_upcoming_lists_packages_still_to_come_soonest_first(store, cohort):
    sel = wbg.select(project="P000001", path=cohort)
    assert [r["package_id"] for r in wbg.upcoming(sel, data=store)] == ["b", "a"]


def test_a_shared_document_belongs_to_each_of_its_projects(store, cohort):
    sel = wbg.select(project="P000002", path=cohort)
    assert [d["doc_id"] for d in wbg.tables(sel, data=store)["documents"]] == ["D2"]


def test_export_writes_the_summary_workbook(store, cohort):
    from openpyxl import load_workbook
    sel = wbg.select(project="P000001", path=cohort)
    out = wbg.export(sel, data=store, today="2026-10-10")
    assert out.endswith(os.path.join("runs", "p000001-2026-10-10"))
    wb = load_workbook(os.path.join(out, "summary.xlsx"))
    assert wb.sheetnames == ["projects", "upcoming packages", "open notices", "components"]


def test_prepare_fetches_only_what_is_new_and_skips_what_did_not_change(
        store, cohort, monkeypatch):
    calls = []
    monkeypatch.setattr(wbg, "_run", lambda script, *a, data, tail=0: calls.append(
        (script, a)))
    sel = wbg.select(project="P000001", path=cohort)
    wbg.prepare(sel, data=store)
    # The store already holds P000001's documents and packages: nothing to fetch.
    assert not any(s.startswith("fetch") for s, _ in calls)
    assert {s for s, _ in calls} >= {"clean.py", "components.py"}
    calls.clear()
    wbg.prepare(sel, data=store)
    assert [s for s, _ in calls] == ["paths.py"], "nothing changed, nothing reruns"
    calls.clear()
    wbg.prepare(wbg.select(project="P000002", path=cohort), data=store)
    assert ("fetch_procurement.py", ("--projects", "P000002")) in calls
